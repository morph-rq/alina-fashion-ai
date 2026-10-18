from __future__ import annotations

import argparse
import base64
import json
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


PLACEHOLDER_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGNgYGD4DwABBAEA"
    "HnOcQAAAAABJRU5ErkJggg=="
)


def scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def sanitized_workbook(path: Path, output: Path) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            if info.filename.startswith("xl/media/"):
                payload = PLACEHOLDER_PNG
            else:
                payload = source.read(info)
            target.writestr(info, payload)
    return output


def inspect_workbook(path: Path) -> dict[str, Any]:
    clean_path = sanitized_workbook(path, Path("tmp/sanitized-source.xlsx"))
    workbook = openpyxl.load_workbook(clean_path, data_only=False)
    cached = openpyxl.load_workbook(clean_path, data_only=True)
    result: dict[str, Any] = {
        "path": str(path),
        "sheets": [],
        "defined_names": [str(item) for item in workbook.defined_names.values()],
    }
    for sheet in workbook.worksheets:
        cached_sheet = cached[sheet.title]
        formulas = []
        nonempty_rows = []
        styles = Counter()
        percentages = []
        for row in sheet.iter_rows():
            entries = []
            for cell in row:
                if cell.value is None:
                    continue
                styles[(cell.style_id, cell.number_format)] += 1
                value = scalar(cell.value)
                cached_value = scalar(cached_sheet[cell.coordinate].value)
                entry = {"cell": cell.coordinate, "value": value}
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    entry["cached"] = cached_value
                    formulas.append(entry)
                if "%" in str(cell.number_format) or (
                    isinstance(cell.value, (float, int)) and 0 < abs(cell.value) < 1
                ):
                    percentages.append(entry | {"number_format": cell.number_format})
                entries.append(entry)
            if entries:
                nonempty_rows.append({"row": row[0].row, "cells": entries})
        result["sheets"].append(
            {
                "name": sheet.title,
                "dimensions": sheet.calculate_dimension(),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "image_count": len(sheet._images),
                "formula_count": len(formulas),
                "formulas": formulas,
                "percentages": percentages,
                "nonempty_rows": nonempty_rows,
                "top_styles": [
                    {"style_id": style_id, "number_format": fmt, "count": count}
                    for (style_id, fmt), count in styles.most_common(20)
                ],
                "column_widths": {
                    key: value.width
                    for key, value in sheet.column_dimensions.items()
                    if value.width
                },
                "hidden_rows": [
                    idx for idx, dim in sheet.row_dimensions.items() if dim.hidden
                ],
                "hidden_columns": [
                    key for key, dim in sheet.column_dimensions.items() if dim.hidden
                ],
            }
        )
    return result


def inspect_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(path)
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(
            {
                "page": index,
                "width": float(page.mediabox.width),
                "height": float(page.mediabox.height),
                "text": page.extract_text() or "",
            }
        )
    return {"path": str(path), "page_count": len(reader.pages), "pages": pages}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, action="append", default=[])
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "workbook": inspect_workbook(args.workbook),
        "pdfs": [inspect_pdf(path) for path in args.pdf],
    }
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
