from __future__ import annotations

import hashlib
import json
import re
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.audit import record_change
from app.calculations import money
from app.config import ROOT_DIR, settings
from app.importer import normalized, parse_customer_allocations, parse_int
from app.models import (
    Brand,
    Customer,
    CustomerAlias,
    ImportBatch,
    ImportChange,
    ImportReview,
    Order,
    OrderItem,
    ReadyBatch,
    Season,
)


HEADER_ALIASES = {
    "article": {"article", "sku", "артикул", "art", "codice", "code"},
    "brand": {"brand", "бренд", "marchio"},
    "customer": {"customer", "client", "клиент", "initials", "allocation", "allocazione"},
    "size": {"size", "размер", "taglia"},
    "quantity": {"quantity", "qty", "количество", "qta", "pcs"},
    "ready": {"ready", "готово", "pronto"},
}


def _header_map(sheet) -> tuple[int, dict[str, int]]:
    for row_number in range(1, min(sheet.max_row, 20) + 1):
        values = {column: normalized(sheet.cell(row_number, column).value) for column in range(1, sheet.max_column + 1)}
        result: dict[str, int] = {}
        for field, aliases in HEADER_ALIASES.items():
            match = next((column for column, value in values.items() if value in aliases), None)
            if match:
                result[field] = match
        if "article" in result:
            return row_number, result
    raise ValueError("Ready workbook needs an Article/SKU column in its first 20 rows")


def _text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return str(value).strip() or None


def _truthy_ready(value: Any) -> bool:
    if value in (None, ""):
        return True
    return normalized(value) in {"1", "yes", "y", "true", "ready", "готово", "pronto", "x"}


def _clone_remainder(item: OrderItem, quantity: int) -> OrderItem:
    return OrderItem(
        order_id=item.order_id,
        customer_id=item.customer_id,
        ready_batch_id=None,
        article=item.article,
        name=item.name,
        color=item.color,
        size=item.size,
        quantity=quantity,
        wholesale_unit_price=item.wholesale_unit_price,
        selling_unit_price=item.selling_unit_price,
        markup_percentage=item.markup_percentage,
        status="Ordered" if item.status == "Ready" else item.status,
        image_path=item.image_path,
        source_file=item.source_file,
        source_sheet=item.source_sheet,
        source_row=item.source_row,
        source_total=money((item.wholesale_unit_price or Decimal("0")) * quantity),
        source_note=item.source_note,
        source_allocation_description=item.source_allocation_description,
        allocation_status=item.allocation_status,
        allocation_group_id=f"{item.allocation_group_id or item.id}:split",
    )


def process_ready_workbook(
    session: Session,
    source_path: str | Path,
    batch_name: str,
    ready_date: date | None = None,
    season_name: str | None = None,
) -> dict[str, Any]:
    path = Path(source_path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    existing = session.scalar(select(ImportBatch).where(ImportBatch.sha256 == digest))
    if existing:
        return json.loads(existing.summary_json or "{}") | {"already_imported": True}
    season = (
        session.scalar(select(Season).where(Season.name == season_name))
        if season_name
        else session.scalar(select(Season).order_by(Season.id.desc()).limit(1))
    )
    if season is None:
        raise ValueError("Import a master season workbook before processing a Ready list")
    if session.scalar(select(ReadyBatch).where(ReadyBatch.season_id == season.id, ReadyBatch.name == batch_name.strip())):
        raise ValueError(f"Ready batch {batch_name!r} already exists")

    try:
        stored_path = path.resolve().relative_to(ROOT_DIR.resolve()).as_posix()
    except ValueError:
        stored_path = None
    import_batch = ImportBatch(filename=path.name, sha256=digest, stored_path=stored_path)
    session.add(import_batch)
    session.flush()
    ready_batch = ReadyBatch(
        season_id=season.id,
        name=batch_name.strip(),
        ready_date=ready_date,
        status="Open",
        notes=f"Created automatically from Ready source {path.name}",
    )
    session.add(ready_batch)
    session.flush()

    aliases = {
        alias.upper(): code.upper()
        for alias, code in session.execute(select(CustomerAlias.alias, Customer.code).join(Customer)).all()
    }
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    header_row, columns = _header_map(sheet)
    candidate_items = list(session.scalars(
        select(OrderItem).options(
            selectinload(OrderItem.customer),
            selectinload(OrderItem.order).selectinload(Order.brand),
        ).where(
            OrderItem.status != "Cancelled",
            OrderItem.ready_batch_id.is_(None),
            OrderItem.order.has(season_id=season.id),
        )
    ).all())

    assigned_items = 0
    affected_customers: set[int] = set()
    reviews = 0
    processed_rows = 0
    for row_number in range(header_row + 1, sheet.max_row + 1):
        article = _text(sheet.cell(row_number, columns["article"]).value)
        if not article:
            continue
        if "ready" in columns and not _truthy_ready(sheet.cell(row_number, columns["ready"]).value):
            continue
        processed_rows += 1
        brand = _text(sheet.cell(row_number, columns["brand"]).value) if "brand" in columns else None
        size = _text(sheet.cell(row_number, columns["size"]).value) if "size" in columns else None
        quantity = parse_int(sheet.cell(row_number, columns["quantity"]).value) if "quantity" in columns else 1
        quantity = quantity or 1
        allocation = _text(sheet.cell(row_number, columns["customer"]).value) if "customer" in columns else None
        if "customer" in columns:
            customer_codes, allocation_reason = parse_customer_allocations(allocation, quantity, settings.default_customer_code, aliases)
        else:
            customer_codes, allocation_reason = [], None
        codes = {code for code in customer_codes if code}
        raw = {sheet.cell(header_row, column).value: sheet.cell(row_number, column).value for column in range(1, sheet.max_column + 1) if sheet.cell(row_number, column).value not in (None, "")}
        reason = allocation_reason
        if len(codes) > 1:
            reason = reason or "A Ready row spanning several customers must be split into one row per customer"
        customer_code = next(iter(codes), None) if len(codes) == 1 else None
        candidates = [
            item for item in candidate_items
            if normalized(item.article) == normalized(article)
            and (not brand or normalized(item.order.brand.name) == normalized(brand))
            and (not size or normalized(item.size) == normalized(size))
            and (not customer_code or (item.customer and item.customer.code.upper() == customer_code))
        ]
        if customer_code is None and len({item.customer_id for item in candidates}) > 1:
            reason = reason or "Article matches more than one customer; add a Customer/Allocation column"
        if reason or not candidates or sum(item.quantity for item in candidates) < quantity:
            if not reason:
                reason = "No unique set of unfulfilled master items reconciles to the requested article/customer/size/quantity"
            reviews += 1
            session.add(ImportReview(import_batch_id=import_batch.id, sheet=sheet.title, row=row_number, raw_values=json.dumps(raw, ensure_ascii=False, default=str), reason=reason))
            session.add(ImportChange(import_batch_id=import_batch.id, source_file=path.name, sheet=sheet.title, row=row_number, action="review", status="Review", previous_values=None, new_values=json.dumps(raw, ensure_ascii=False, default=str), reason=reason))
            continue

        remaining = quantity
        used: list[OrderItem] = []
        for item in sorted(candidates, key=lambda value: value.id):
            if remaining <= 0:
                break
            take = min(item.quantity, remaining)
            if take < item.quantity:
                session.add(_clone_remainder(item, item.quantity - take))
                item.quantity = take
                item.source_total = money((item.wholesale_unit_price or Decimal("0")) * take)
            item.ready_batch_id = ready_batch.id
            item.status = "Ready"
            remaining -= take
            used.append(item)
            assigned_items += take
            if item.customer_id:
                affected_customers.add(item.customer_id)
            if item in candidate_items:
                candidate_items.remove(item)
        session.add(ImportChange(import_batch_id=import_batch.id, source_file=path.name, sheet=sheet.title, row=row_number, action="ready_assigned", status="Applied", previous_values=None, new_values=json.dumps({"source": raw, "item_ids": [item.id for item in used], "ready_batch": ready_batch.name}, ensure_ascii=False, default=str), reason=None))
        record_change(session, entity_type="ReadyBatch", entity_id=ready_batch.id, field="items", previous_value=None, new_value=[item.id for item in used], source_user="Ready importer", description=f"{ready_batch.name}: assigned {quantity} unit(s) from {path.name} row {row_number}")

    summary = {
        "filename": path.name,
        "season": season.name,
        "ready_batch_id": ready_batch.id,
        "ready_batch": ready_batch.name,
        "rows": processed_rows,
        "assigned_units": assigned_items,
        "affected_customer_ids": sorted(affected_customers),
        "reviews": reviews,
        "already_imported": False,
    }
    import_batch.summary_json = json.dumps(summary, ensure_ascii=False)
    return summary
