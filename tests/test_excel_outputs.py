"""Excel output correctness: formulas follow the Decimal rule, private exports are allowlisted."""
from datetime import date
from decimal import Decimal as D
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from test_workspace_features import workspace_session, workspace_client  # noqa: F401
from app import excel_exchange as ex
from app.calculations import item_amounts
from app.excel_export import build_business_workbook, verify_business_workbook
from app.models import Customer, OrderItem, Payment


def _mixed_rows(db):
    item = db.scalar(select(OrderItem))
    other = Customer(code='ZZ', name='Zoe Secret', country_code='IT')
    db.add(other)
    db.flush()
    db.add_all([
        OrderItem(order_id=item.order_id, customer_id=item.customer_id, article='007', name='Half cents', size='M', quantity=3,
                  wholesale_unit_price=D('10.05'), markup_percentage=D('0.35')),
        OrderItem(order_id=item.order_id, customer_id=item.customer_id, article='=HYPERLINK("x")', name='Override', size='S', quantity=2,
                  wholesale_unit_price=D('40'), selling_unit_price=D('57.99'), markup_percentage=D('0.30')),
        OrderItem(order_id=item.order_id, customer_id=other.id, article='SECRET-1', name='Other customer dress', size='38', quantity=1,
                  wholesale_unit_price=D('999'), markup_percentage=D('0.30')),
        OrderItem(order_id=item.order_id, customer_id=item.customer_id, name='Cancelled row', size='L', quantity=1,
                  wholesale_unit_price=D('500'), markup_percentage=D('0.30'), status='Cancelled'),
    ])
    db.add(Payment(season_id=item.order.season_id, customer_id=item.customer_id, date=date(2026, 9, 1), amount=D('50'), type='payment'))
    db.add(Payment(season_id=item.order.season_id, customer_id=item.customer_id, date=date(2026, 9, 2), amount=D('5'), type='refund'))
    db.commit()
    return item, other


def test_business_workbook_formulas_follow_the_decimal_rule(workspace_session):
    db = workspace_session
    _mixed_rows(db)
    content = build_business_workbook(db)
    assert verify_business_workbook(content) == []
    book = load_workbook(BytesIO(content))
    assert {'Master summary', 'Checks', 'Products'} <= set(book.sheetnames)
    assert any(name.startswith('Test Atelier') for name in book.sheetnames)  # brand view
    products = book['Products']
    by_article = {products.cell(row, 6).value: row for row in range(5, products.max_row + 1)}
    assert products.cell(by_article['007'], 6).data_type == 's'
    assert products.cell(by_article['=HYPERLINK("x")'], 6).data_type == 's'
    half = db.scalar(select(OrderItem).where(OrderItem.article == '007'))
    assert D(str(products.cell(by_article['007'], 33).value)) == item_amounts(half).final == D('40.70')
    checks = book['Checks']
    labels = [checks.cell(row, 1).value for row in range(1, checks.max_row + 1)]
    assert 'Workbook schema' in labels and 'Products: final total' in labels


def test_private_customer_workbook_is_allowlisted_and_refunds_are_signed(workspace_session, workspace_client):
    db = workspace_session
    item, _other = _mixed_rows(db)
    response = workspace_client.get(f'/exports/customers/{item.customer_id}.xlsx?season=AW26')
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    text = ' '.join(str(cell.value) for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
    for forbidden in ('Wholesale', 'Markup', 'Zoe', 'SECRET-1', 'Other customer', 'Source', 'Cancelled row'):
        assert forbidden not in text, forbidden
    payments = book['Payments']
    rows = [[payments.cell(row, col).value for col in range(1, 5)] for row in range(5, payments.max_row + 1)]
    assert ['AW26', rows[1][1], 'Refund', -5.0] in rows
    products = book['Products']
    override_row = next(row for row in range(5, products.max_row + 1) if products.cell(row, 4).value == 'Override')
    assert products.cell(override_row, 8).value == 57.99 and products.cell(override_row, 9).value == 115.98
    internal = workspace_client.get(f'/exports/customers/{item.customer_id}/internal.xlsx?season=AW26')
    assert internal.status_code == 200 and 'Products' in load_workbook(BytesIO(internal.content)).sheetnames


def test_editable_workbook_line_formula_matches_app_rounding(workspace_session, tmp_path, monkeypatch):
    db = workspace_session
    monkeypatch.setattr(ex, 'EXCHANGE_DIR', tmp_path / 'exchange')
    _mixed_rows(db)
    sheet = load_workbook(BytesIO(ex.export_editable(db)))['Orders']
    row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 6).value == '007')
    assert sheet.cell(row, 22).value == (f'=IF(U{row}="","",IF(M{row}<>"",ROUND(M{row}*J{row},2),'
                                         f'ROUND(K{row}*J{row},2)+ROUND(ROUND(K{row}*J{row},2)*L{row}/100,2)))')


def test_libreoffice_recalculation_matches_app_amounts(workspace_session):
    """A controlled recalculation (when LibreOffice is installed) reproduces every app amount."""
    import pytest
    from app.recalculation import recalculate, soffice_path
    if soffice_path() is None:
        pytest.skip('LibreOffice is not installed')
    db = workspace_session
    _mixed_rows(db)
    book = load_workbook(BytesIO(recalculate(build_business_workbook(db))), data_only=True)
    products = book['Products']
    for row in range(5, products.max_row + 1):
        if products.cell(row, 1).value is None:
            continue
        assert round(products.cell(row, 14).value, 2) == round(products.cell(row, 33).value, 2), row
    checks = book['Checks']
    statuses = [checks.cell(row, 4).value for row in range(1, checks.max_row + 1) if checks.cell(row, 4).value in ('OK', 'CHECK')]
    assert statuses and set(statuses) == {'OK'}
    master = book['Master summary']
    assert master.cell(5, 3).value == sum(products.cell(r, 33).value for r in range(5, products.max_row + 1) if products.cell(r, 1).value)
