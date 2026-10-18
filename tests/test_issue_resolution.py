import io
import re
from decimal import Decimal

from fastapi.testclient import TestClient
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sqlalchemy import func, select

from app import ai_intake,intake_archive
from app.database import get_db
from app.main import app
from app.models import AllocationReview,AuditLog,ImportBatch,ImportReview,OrderItem,IntakeConfirmation
from app.validation import validate_system
from test_ai_intake import _session


def test_ai_spreadsheet_keeps_row_numbers_and_status_colors():
    workbook=Workbook();sheet=workbook.active
    for name,color in [('In order','FFFF00'),('Completed','92D050'),('Waiting','FFFFFF')]:
        sheet.append([name,'SKU',2])
        for cell in sheet[sheet.max_row]:cell.fill=PatternFill('solid',fgColor=color)
    data=io.BytesIO();workbook.save(data)
    result=ai_intake._spreadsheet_text('source.xlsx',data.getvalue())
    assert 'Row 1 [visual status hint: process' in result
    assert 'Row 2 [visual status hint: done' in result
    assert 'Row 3 [visual status hint: pending' in result


def test_source_review_links_and_explicit_resolution():
    db,order=_session()
    item=db.scalar(select(OrderItem));item.source_file='source.xlsx';item.source_sheet='Source';item.source_row=9
    batch=ImportBatch(filename='source.xlsx',sha256='1'*64);db.add(batch);db.flush()
    review=ImportReview(import_batch_id=batch.id,sheet='Source',row=9,raw_values='original evidence',reason='Check price')
    db.add(review);db.commit()
    app.dependency_overrides[get_db]=lambda:db
    try:
        with TestClient(app) as client:
            page=client.get(f'/source-reviews/import/{review.id}')
            assert page.status_code==200 and f'/worksheet/items/{item.id}?review=import-{review.id}' in page.text
            response=client.post(f'/source-reviews/import/{review.id}/resolve',data={'resolution':'Verified against the supplier invoice.'})
            assert response.status_code==200 and review.status=='Resolved'
            assert review.raw_values=='original evidence'
            assert db.scalar(select(AuditLog).where(AuditLog.entity_type=='ImportReview'))
    finally:app.dependency_overrides.clear();db.close()


def test_cannot_close_unallocated_source_review():
    db,order=_session()
    item=db.scalar(select(OrderItem));item.customer_id=None;item.allocation_status='Review'
    item.source_file='source.xlsx';item.source_sheet='Source';item.source_row=9
    review=AllocationReview(source_file='source.xlsx',sheet='Source',row=9,raw_values='original',reason='Unclear customer')
    db.add(review);db.commit()
    app.dependency_overrides[get_db]=lambda:db
    try:
        with TestClient(app) as client:
            response=client.post(f'/source-reviews/allocation/{review.id}/resolve',data={'resolution':'Needs further confirmation from the customer.'})
            assert response.status_code==422 and review.status=='Open'
            assert 'Allocate every matching' in response.text
            report=validate_system(db)
            issue=next(i for i in report.issues if i.code=='UNALLOCATED_PRODUCT')
            assert issue.action_url==f'/worksheet/items/{item.id}' and issue.suggestion
            assert all(i.action_url and i.suggestion for i in report.issues if i.severity!='pass')
    finally:app.dependency_overrides.clear();db.close()


def test_image_intake_archives_original_and_repeated_save_is_idempotent(monkeypatch):
    db,order=_session()
    seen=[]
    facts=ai_intake.IntakeFacts(customer_name='Maria Ala',brand=order.brand.name,product_name='Test coat',
        explicit_size='M',quantity=1,stated_price='200',price_type='wholesale',source_attachment_names=['order.png'])
    def analyze(message,attachments):
        seen.extend(attachments)
        return ai_intake.IntakeBundle(orders=[facts],source_summary='Test source')
    monkeypatch.setattr(ai_intake,'analyze_submission',analyze)
    monkeypatch.setattr(ai_intake,'search_product_images',lambda *_: (_ for _ in ()).throw(AssertionError('No decorative image search')))
    app.dependency_overrides[get_db]=lambda:db
    try:
        with TestClient(app) as client:
            image=io.BytesIO();Image.new('RGB',(20,20),'red').save(image,format='PNG')
            page=client.post('/ai-intake/analyze',data={'message':'Read this order','season_name':'SS26'},files={'attachments':('order.png',image.getvalue(),'image/png')})
            assert page.status_code==200 and seen[0][2]==image.getvalue()
            token=re.search(r'name="confirmation_token" value="([a-f0-9]+)"',page.text).group(1)
            bundle=re.search(r'name="source_bundle_token" value="([a-f0-9]+)"',page.text).group(1)
            assert (intake_archive.bundle_path(bundle)/'00-order.png').read_bytes()==image.getvalue()
            assert client.get(f'/intake-sources/{bundle}').status_code==200
            db.commit()
            data=dict(order_id=order.id,customer_name='Maria Ala',product_name='Test coat',size='M',quantity=1,stated_price='200',price_type='wholesale',worksheet_state='process',save_mode='batch',confirmation_token=token,source_bundle_token=bundle)
            first=client.post('/ai-intake/confirm',data=data)
            assert first.status_code==200 and first.json()['ok']
            second=client.post('/ai-intake/confirm',data=data)
            assert second.status_code==200 and second.json()['item_id']==first.json()['item_id']
            assert db.scalar(select(func.count(IntakeConfirmation.token)))==1
            assert db.scalar(select(func.count(OrderItem.id)))==2
            assert db.get(OrderItem,first.json()['item_id']).worksheet_status=='process'
            assert 'View saved AI source' in client.get(first.json()['redirect']).text
    finally:app.dependency_overrides.clear();db.close()
