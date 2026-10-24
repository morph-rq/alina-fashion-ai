from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from test_intake_workflow import env
from app import workbook_workspace as workspace, drafts
from app.models import IntakeSession


def workbook_bytes():
    book=Workbook();sheet=book.active;sheet.title='Supplier'
    sheet.append(['Article','Name','Size','Quantity','Price','Customer'])
    sheet.append(['001-A','Coat','38',2,100,'Anna'])
    sheet.append(['001-A','Coat','38',2,100,'Anna'])
    stream=BytesIO();book.save(stream);return stream.getvalue()


def upload(env,monkeypatch):
    client,factory,root=env
    monkeypatch.setattr(workspace,'WORKBOOK_DIR',root/'workbooks')
    response=client.post('/workbooks/upload',files={'file':('ORDER_SS27.xlsx',workbook_bytes())},follow_redirects=False)
    assert response.status_code==303
    return response.headers['location'].split('/')[-1]


def test_open_edit_export_and_conflict(env,monkeypatch):
    token=upload(env,monkeypatch);client,_,_=env
    page=client.get('/workbooks/'+token)
    assert page.status_code==200
    assert '001-A' in page.text and 'Save cell edits' in page.text
    response=client.post(f'/workbooks/{token}/cells',data={'sheet':0,'page':1,'revision':0,'cell:2:2':'Edited coat','cell:2:1':'=1+1'},follow_redirects=False)
    assert response.status_code==303
    conflict=client.post(f'/workbooks/{token}/cells',data={'sheet':0,'page':1,'revision':0,'cell:2:2':'Lost edit'})
    assert 'changed in another tab' in conflict.text
    result=client.get(f'/workbooks/{token}/download')
    book=load_workbook(BytesIO(result.content))
    assert book.active['B2'].value=='Edited coat'
    assert book.active['A2'].data_type=='s'
    assert book.active['A2'].value=='=1+1'
    original=load_workbook(workspace.folder(token)/'original.xlsx')
    assert original.active['B2'].value=='Coat'


def rule_form(revision=0):
    return {'sheet':0,'revision':revision,'season':'SS27','brand':'Supplier','start':2,'end':3,
            'map:article':'1','map:product_name':'2','map:explicit_size':'3','map:quantity':'4',
            'map:stated_price':'5','map:customer_name':'6','price_type':'wholesale','price_basis':'unit',
            'markup':'35','profile_name':'Supplier standard','notes':'Check sizes'}


def test_rules_preview_draft_and_idempotent_repeat(env,monkeypatch):
    token=upload(env,monkeypatch);client,factory,_=env
    assert client.post(f'/workbooks/{token}/rules',data=rule_form(),follow_redirects=False).status_code==303
    assert 'Supplier standard' in workspace.profiles()
    preview=client.post(f'/workbooks/{token}/preview',data={'sheet':0})
    assert preview.status_code==200 and '2 product rows in this preview' in preview.text
    response=client.post(f'/workbooks/{token}/draft',data={'sheet':0,'revision':1},follow_redirects=False)
    assert response.status_code==303
    again=client.post(f'/workbooks/{token}/draft',data={'sheet':0,'revision':1},follow_redirects=False)
    assert response.headers['location']==again.headers['location']
    with factory() as db:
        sessions=db.scalars(select(IntakeSession)).all()
        assert len(sessions)==1
        assert len(sessions[0].rows)==2  # identical supplier rows stay distinct
        for row in sessions[0].rows:
            assert drafts.value_of(row,'price')=='100.00'
            assert drafts.value_of(row,'order')['terms']['markup']=='0.35'


def test_invalid_quantity_and_profile_reuse(env,monkeypatch):
    token=upload(env,monkeypatch);client,_,_=env
    client.post(f'/workbooks/{token}/rules',data=rule_form())
    client.post(f'/workbooks/{token}/cells',data={'sheet':0,'page':1,'revision':1,'cell:2:4':'1.5'})
    response=client.post(f'/workbooks/{token}/draft',data={'sheet':0,'revision':2})
    assert 'quantity must be a whole number' in response.text
    second=upload(env,monkeypatch)
    response=client.post(f'/workbooks/{second}/rules',data={'sheet':0,'revision':0,'load_profile':'Supplier standard'},follow_redirects=False)
    assert response.status_code==303
    assert workspace.read_state(second)['sheets'][0]['rule']['markup']=='35'


def test_claude_reads_cells_and_applies_only_after_confirmation(env,monkeypatch):
    token=upload(env,monkeypatch);client,_,_=env
    from app import ai_intake
    def reply(schema,system,parts,**kwargs):
        assert '001-A' in parts[1]['text']
        assert 'D2' in parts[1]['text']
        return workspace.WorkbookAnswer(answer='D2 contains 2. I propose 3.',references=['D2'],changes=[workspace.CellProposal(row=2,column=4,value='3',reason='Requested quantity')])
    monkeypatch.setattr(ai_intake,'_generate',reply)
    response=client.post(f'/workbooks/{token}/ask',data={'sheet':0,'start':2,'end':3,'question':'Change quantity in D2 to 3'})
    assert response.status_code==200 and 'Apply these cell changes' in response.text
    sheet=workspace.read_state(token)['sheets'][0]
    assert sheet['rows'][1][3]==2
    proposal=sheet['assistant']['id']
    response=client.post(f'/workbooks/{token}/apply-claude',data={'sheet':0,'proposal':proposal})
    assert response.status_code==200
    assert workspace.read_state(token)['sheets'][0]['rows'][1][3]==3
    repeat=client.post(f'/workbooks/{token}/apply-claude',data={'sheet':0,'proposal':proposal})
    assert 'stale or already applied' in repeat.text


def test_claude_rejects_out_of_range_changes_and_stale_proposals(env,monkeypatch):
    token=upload(env,monkeypatch);client,_,_=env
    from app import ai_intake
    answer=workspace.WorkbookAnswer(answer='Proposed edit',references=['D2'],changes=[workspace.CellProposal(row=2,column=4,value='3',reason='Requested')])
    monkeypatch.setattr(ai_intake,'_generate',lambda *args,**kwargs:answer)
    result=client.post(f'/workbooks/{token}/ask',data={'sheet':0,'start':3,'end':3,'question':'Change D2'})
    assert 'outside the selected range' in result.text
    client.post(f'/workbooks/{token}/ask',data={'sheet':0,'start':2,'end':3,'question':'Change D2'})
    proposal=workspace.read_state(token)['sheets'][0]['assistant']['id']
    client.post(f'/workbooks/{token}/cells',data={'sheet':0,'page':1,'revision':0,'cell:2:4':'4'})
    result=client.post(f'/workbooks/{token}/apply-claude',data={'sheet':0,'proposal':proposal})
    assert 'stale or already applied' in result.text
    assert workspace.read_state(token)['sheets'][0]['rows'][1][3]==4


def test_saved_intake_opens_in_editor_without_reupload(env,monkeypatch):
    client,_,root=env
    from app import ai_intake
    monkeypatch.setattr(workspace,'WORKBOOK_DIR',root/'workbooks')
    monkeypatch.setattr(ai_intake,'ai_configured',lambda:False)
    intake=client.post('/intake',files={'attachments':('ORDER_SS27.xlsx',workbook_bytes())},follow_redirects=False)
    path=intake.headers['location']
    first=client.post(path+'/workbook',follow_redirects=False)
    second=client.post(path+'/workbook',follow_redirects=False)
    assert first.status_code==303
    assert first.headers['location']==second.headers['location']
    page=client.get(first.headers['location'])
    assert '001-A' in page.text
    assert 'Work with Claude' in page.text


def coloured_bytes():
    from openpyxl.styles import PatternFill,Font,Alignment
    book=load_workbook(BytesIO(workbook_bytes()));ws=book.active
    for row,colour in [(2,'FFF2CC'),(3,'C6E0B4')]:
        for cell in ws[row]:cell.fill=PatternFill('solid',fgColor=colour)
    ws['B2'].font=Font(bold=True,color='990000')
    ws['E2'].number_format='#,##0.00'
    ws.column_dimensions['B'].width=32
    stream=BytesIO();book.save(stream);return stream.getvalue()


def coloured_upload(env,monkeypatch):
    client,_,root=env
    monkeypatch.setattr(workspace,'WORKBOOK_DIR',root/'workbooks')
    result=client.post('/workbooks/upload',files={'file':('ORDER_SS27.xlsx',coloured_bytes())},follow_redirects=False)
    return result.headers['location'].split('/')[-1]


def test_source_colours_migrate_without_losing_cell_edits(env,monkeypatch):
    token=coloured_upload(env,monkeypatch);client,_,_=env
    state=workspace.read_state(token);sheet=state['sheets'][0]
    assert workspace.row_status(sheet,2)=='ready'
    assert workspace.row_status(sheet,3)=='shipped'
    assert sheet['styles']['2:2']['fill']=='#FFF2CC'
    assert sheet['styles']['2:2']['bold']
    assert sheet['column_dimensions']['2']['width']==32
    sheet['rows'][1][1]='User edit before colour migration'
    sheet.pop('format_version');sheet.pop('styles')
    workspace.save_json(workspace.folder(token)/'state.json',state)
    migrated=workspace.read_state(token)['sheets'][0]
    assert migrated['rows'][1][1]=='User edit before colour migration'
    assert migrated['styles']['2:2']['fill']=='#FFF2CC'
    page=client.get('/workbooks/'+token)
    assert page.status_code==200 and '#FFF2CC' in page.text
    assert 'Yellow: ready, awaiting shipment' in page.text


def test_autosave_status_markup_and_filtered_styled_export(env,monkeypatch):
    from app.workbook_formatting import row_pricing
    token=coloured_upload(env,monkeypatch);client,_,_=env
    client.post(f'/workbooks/{token}/rules',data=rule_form())
    response=client.post(f'/workbooks/{token}/cells',data={'sheet':0,'page':1,'revision':1,
        'meta:2:status':'cannot','meta:2:markup':'25','meta:2:note':'Factory refused',
        'meta:3:status':'ready','cell:3:5':'120'},headers={'X-Workbook-Autosave':'1'})
    assert response.status_code==200
    assert response.json()['rows']['2']['final_unit']=='125.00'
    sheet=workspace.read_state(token)['sheets'][0]
    assert row_pricing(sheet,3)['final_line']=='324.00'
    assert workspace.row_status(sheet,2)=='cannot'
    exported=client.get(f'/workbooks/{token}/download?sheets=0&rows=2&columns=1,2,5&extras=true')
    assert exported.status_code==200
    ws=load_workbook(BytesIO(exported.content)).active
    assert ws.max_row==2
    assert ws['B2'].value=='Coat'
    assert ws['B2'].fill.fgColor.rgb[-6:]=='FFC7CE'
    assert ws['B2'].font.bold
    assert ws['C2'].number_format=='#,##0.00'
    assert ws['D1'].value=='ALINA status'
    assert ws['F2'].value==125
    ready=client.get(f'/workbooks/{token}/download?sheets=0&status=ready')
    ready_sheet=load_workbook(BytesIO(ready.content)).active
    assert ready_sheet.max_row==2 and ready_sheet['E2'].value==120
    assert ready_sheet['B2'].fill.fgColor.rgb[-6:]=='FFF2CC'
    stale=client.post(f'/workbooks/{token}/cells',data={'sheet':0,'page':1,'revision':1},headers={'X-Workbook-Autosave':'1'})
    assert stale.status_code==409


def test_selection_filters_and_pricing_validation(env,monkeypatch):
    token=coloured_upload(env,monkeypatch);client,_,_=env
    client.post(f'/workbooks/{token}/rules',data=rule_form())
    state=workspace.read_state(token)
    state['sheets'][0]['rows'][2][5]='Bianca'
    workspace.save_json(workspace.folder(token)/'state.json',state)
    response=client.get(f'/workbooks/{token}/download?sheets=0&customer=Bianca&columns=1,6&extras=false')
    ws=load_workbook(BytesIO(response.content)).active
    assert ws.max_row==2 and ws.max_column==2 and ws['B2'].value=='Bianca'
    response=client.get(f'/workbooks/{token}/download?sheets=0&status=cannot')
    assert response.status_code==422
    invalid=client.post(f'/workbooks/{token}/cells',data={'sheet':0,'page':1,'revision':1,'meta:2:markup':'NaN'},headers={'X-Workbook-Autosave':'1'})
    assert invalid.status_code==422
    assert workspace.read_state(token)['sheets'][0]['revision']==1


def test_autosave_handles_wide_multipart_workbooks(env,monkeypatch):
    client,_,root=env
    monkeypatch.setattr(workspace,'WORKBOOK_DIR',root/'workbooks')
    book=Workbook()
    for r in range(1,41):book.active.append([f'R{r}C{c}' for c in range(1,31)])
    stream=BytesIO();book.save(stream)
    response=client.post('/workbooks/upload',files={'file':('WIDE_SS27.xlsx',stream.getvalue())},follow_redirects=False)
    token=response.headers['location'].split('/')[-1]
    values={'sheet':'0','page':'1','revision':'0',**{f'cell:{r}:{c}':f'R{r}C{c}' for r in range(1,41) for c in range(1,31)}}
    result=client.post(f'/workbooks/{token}/cells',files=[(key,(None,value)) for key,value in values.items()],headers={'X-Workbook-Autosave':'1'})
    assert result.status_code==200
    assert result.json()['revision']==1
