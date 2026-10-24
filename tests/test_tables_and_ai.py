"""Described tables (season/brand and custom), spreadsheet cell edits, and Settings › AI."""
from decimal import Decimal
from io import BytesIO
import json

import httpx
from openpyxl import load_workbook
from sqlalchemy import select

from test_workspace_features import workspace_session, workspace_client  # noqa: F401
from app import ai_intake, ai_settings
from app.config import settings
from app.models import AuditLog, BrandSeasonRule, CustomTable, CustomTableRow, Order, OrderItem, Season
from app.table_planner import parse_locally


def no_ai(monkeypatch):
    monkeypatch.setattr(ai_intake, 'ai_configured', lambda: False)


def test_season_description_creates_brand_tables_after_confirmation(workspace_session, workspace_client, monkeypatch):
    no_ai(monkeypatch)
    draft = workspace_client.post('/tables/plan', data={'description': 'New season FW27 with brands Lumen and Test Atelier, 30% markup, 40/60 prepayment'})
    assert draft.status_code == 200 and 'Not created yet' in draft.text and 'Lumen' in draft.text
    assert workspace_session.scalar(select(Season).where(Season.name == 'FW27')) is None   # draft only
    response = workspace_client.post('/tables/season', data={
        'season': 'fw27', 'brand_name': ['Lumen', 'Test Atelier', ''], 'markup': ['30', '30', ''],
        'initial': ['40', '', ''], 'transport': ['', '', '']}, follow_redirects=False)
    assert response.status_code == 303 and response.headers['location'].startswith('/seasons/FW27')
    season = workspace_session.scalar(select(Season).where(Season.name == 'FW27'))
    rules = {r.brand.name: r for r in workspace_session.scalars(select(BrandSeasonRule).where(BrandSeasonRule.season_id == season.id))}
    assert rules['Lumen'].markup_percentage == Decimal('0.30') and rules['Lumen'].initial_prepayment_percentage == Decimal('0.40')
    assert rules['Lumen'].ready_balance_percentage == Decimal('0.60') and rules['Test Atelier'].initial_prepayment_percentage == Decimal('0.50')
    assert len(workspace_session.scalars(select(Order).where(Order.season_id == season.id)).all()) == 2
    # Repeating it never duplicates tables or overwrites rules.
    workspace_client.post('/tables/season', data={'season': 'FW27', 'brand_name': ['Lumen'], 'markup': ['90'], 'initial': [''], 'transport': ['']})
    assert len(workspace_session.scalars(select(Order).where(Order.season_id == season.id)).all()) == 2
    assert rules['Lumen'].markup_percentage == Decimal('0.30')


def test_custom_table_described_edited_in_cells_and_exported(workspace_session, workspace_client, monkeypatch):
    no_ai(monkeypatch)
    plan = parse_locally('Samples to return: brand, item, qty, deadline\nLumen; blazer; 2; 2026-11-01')
    response = workspace_client.post('/tables/custom', data={
        'name': plan.title, 'column_name': [c.name for c in plan.columns], 'column_type': [c.type for c in plan.columns],
        'rows_json': json.dumps(plan.rows)}, follow_redirects=False)
    table = workspace_session.scalar(select(CustomTable))
    assert response.headers['location'] == f'/tables/{table.id}' and table.name == 'Samples to return'
    page = workspace_client.get(f'/tables/{table.id}')
    assert page.status_code == 200 and 'blazer' in page.text and 'data-edit-key' in page.text
    row = workspace_session.scalar(select(CustomTableRow))
    qty = next(c['key'] for c in json.loads(table.columns_json) if c['name'] == 'Qty')
    ok = workspace_client.post(f'/tables/{table.id}/cell', json={'row_id': row.id, 'key': qty, 'value': '3', 'revision': row.revision})
    assert ok.status_code == 200 and ok.json()['value'] == '3'
    stale = workspace_client.post(f'/tables/{table.id}/cell', json={'row_id': row.id, 'key': qty, 'value': '4', 'revision': row.revision - 1 if row.revision > 1 else 999})
    assert stale.status_code == 409
    bad = workspace_client.post(f'/tables/{table.id}/cell', json={'row_id': row.id, 'key': qty, 'value': 'many', 'revision': ok.json()['revision']})
    assert bad.status_code == 422
    added = workspace_client.post(f'/tables/{table.id}/rows/describe', data={'text': 'Vera; coat; 1; 2026-11-03\nAtelier; scarf; 4; 2026-11-05'}, follow_redirects=False)
    assert 'added=2' in added.headers['location']
    book = load_workbook(BytesIO(workspace_client.get(f'/tables/{table.id}/export.xlsx').content))
    sheet = book.active
    assert [c.value for c in sheet[1]] == ['Brand', 'Item', 'Qty', 'Deadline'] and sheet['C2'].value == 3 and sheet.max_row == 4
    assert workspace_session.scalar(select(AuditLog).where(AuditLog.entity_type == 'CustomTableRow')) is not None


def test_product_cells_save_through_item_rules(workspace_session, workspace_client):
    item = workspace_session.scalar(select(OrderItem))
    saved = workspace_client.post(f'/worksheet/items/{item.id}/cell', json={'field': 'size', 'value': 'L', 'revision': item.revision})
    assert saved.status_code == 200 and saved.json()['value'] == 'L'
    workspace_session.refresh(item)
    assert item.size == 'L'
    price = workspace_client.post(f'/worksheet/items/{item.id}/cell', json={'field': 'wholesale', 'value': '120,50', 'revision': item.revision})
    assert price.json()['value'] == '120.50'
    assert workspace_client.post(f'/worksheet/items/{item.id}/cell', json={'field': 'customer', 'value': 'NOPE', 'revision': price.json()['revision']}).status_code == 422
    assert workspace_client.post(f'/worksheet/items/{item.id}/cell', json={'field': 'quantity', 'value': '2', 'revision': 1}).status_code == 409
    assert 'data-edit-url="/worksheet/items/' in workspace_client.get('/worksheet').text


def test_ai_key_saved_in_settings_selects_sonnet_or_haiku(workspace_client, monkeypatch):
    for name in ('ai_api_key', 'anthropic_api_key', 'gemini_api_key', 'ai_model', 'ai_fallback_models', 'ai_provider'):
        monkeypatch.setattr(settings, name, '')
    assert 'Not connected' in workspace_client.get('/settings/ai').text
    assert 'error=' in workspace_client.post('/settings/ai', data={'api_key': 'AIza-gemini', 'model': 'claude-sonnet-5-5'}, follow_redirects=False).headers['location']
    seen = {}

    def fake_post(url, json=None, headers=None, timeout=None):
        seen.update(model=json['model'], key=headers['x-api-key'])
        return httpx.Response(200, json={'content': [{'type': 'text', 'text': 'OK'}]})
    monkeypatch.setattr(ai_settings.httpx, 'post', fake_post)
    response = workspace_client.post('/settings/ai', data={'action': 'test', 'api_key': 'sk-ant-test-key-1234', 'model': 'claude-haiku-5-5'}, follow_redirects=False)
    assert 'message=' in response.headers['location'] and seen == {'model': 'claude-haiku-5-5', 'key': 'sk-ant-test-key-1234'}
    assert settings.active_ai_provider == 'anthropic' and settings.active_ai_model == 'claude-haiku-5-5'
    assert settings.active_ai_fallback_models == 'claude-sonnet-5-5' and ai_intake.ai_configured()
    page = workspace_client.get('/settings/ai').text
    assert 'sk-ant-test-key-1234' not in page and '…1234' in page
    workspace_client.post('/settings/ai', data={'action': 'save', 'api_key': '', 'model': 'claude-sonnet-5-5'})
    assert settings.active_ai_model == 'claude-sonnet-5-5' and settings.active_ai_api_key == 'sk-ant-test-key-1234'

    def refused(url, json=None, headers=None, timeout=None):
        return httpx.Response(401, json={})
    monkeypatch.setattr(ai_settings.httpx, 'post', refused)
    assert 'error=' in workspace_client.post('/settings/ai', data={'action': 'test', 'api_key': 'sk-ant-wrong', 'model': 'claude-sonnet-5-5'}, follow_redirects=False).headers['location']
    assert settings.active_ai_api_key == 'sk-ant-test-key-1234'   # a failed test does not replace the working key
    workspace_client.post('/settings/ai', data={'action': 'remove'})
    assert settings.active_ai_api_key == ''


def test_tables_view_in_workspace_switch(workspace_client):
    page = workspace_client.get('/tables')
    assert page.status_code == 200 and 'Describe a table' in page.text and 'href="/tables"' in page.text
