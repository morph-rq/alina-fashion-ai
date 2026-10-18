import io
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine,select,func
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import excel_exchange as ex
from app.database import Base,get_db
from app.main import app
from app.models import Brand,Customer,Order,OrderItem,Season,PurchaseGroup,SellerCreditAccount,ImportBatch
from app.validation import validate_system


@pytest.fixture
def sample(tmp_path,monkeypatch):
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db=sessionmaker(bind=engine,expire_on_commit=False)()
    season=Season(name='SS26'); brand=Brand(name='Test brand'); customer=Customer(code='AA',name='Alina')
    db.add_all([season,brand,customer]); db.flush()
    order=Order(season_id=season.id,brand_id=brand.id)
    group=PurchaseGroup(name='Test group',season_id=season.id)
    db.add_all([order,group]); db.flush()
    item=OrderItem(order_id=order.id,customer_id=customer.id,purchase_group_id=group.id,name='Dress',size='38',quantity=2,wholesale_unit_price=Decimal('100'),markup_percentage=Decimal('.35'),worksheet_status='process',status='Cancelled')
    db.add(item);db.commit()
    monkeypatch.setattr(ex,'EXCHANGE_DIR',tmp_path/'exchange')
    yield db,item
    db.close();engine.dispose()


def edited(content,edit):
    wb=load_workbook(io.BytesIO(content));edit(wb)
    out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()


def apply(db,content):
    db.rollback()
    path=ex.EXCHANGE_DIR/'test-upload.xlsx';path.write_bytes(content)
    return ex.apply_upload(db,content,path)


def test_roundtrip_no_change_and_color_contract(sample):
    db,item=sample
    content=ex.export_editable(db)
    result=ex.inspect_upload(db,content)
    assert not result['errors'] and not result['rows'] and result['unchanged']==1
    wb=load_workbook(io.BytesIO(content))
    assert wb['Orders'].max_row==2 and wb['Orders']['A2'].value==item.id
    assert wb['Orders']['A2'].fill.fgColor.rgb.endswith('FFF2B3')
    assert not wb['Orders']._images
    assert len(wb['Orders'].data_validations.dataValidation)==1


def test_update_preserves_financial_and_source_state_and_duplicate_safe(sample):
    db,item=sample
    content=edited(ex.export_editable(db),lambda wb:wb['Orders'].__setitem__('H2','navy'))
    preview=ex.inspect_upload(db,content)
    assert preview['rows'][0]['changes']==[dict(field='Color',before='',after='navy')]
    assert item.color is None
    apply(db,content)
    assert item.color=='navy' and item.status=='Cancelled' and item.quantity==2
    with pytest.raises(ValueError,match='already been applied'):ex.inspect_upload(db,content)


def test_new_row_existing_customer_and_new_season(sample):
    db,item=sample
    def change(wb):
        wb['Orders'].append([None,'SS27','New brand','AA','Alina','SKU-7','Coat','black','M',1,250,35,None,'Not completed',None,'Source confirmed'])
    content=edited(ex.export_editable(db,blank=True),change)
    assert not ex.inspect_upload(db,content)['errors']
    apply(db,content)
    new=db.scalar(select(OrderItem).where(OrderItem.article=='SKU-7'))
    assert new.order.season.name=='SS27' and new.customer_id==item.customer_id
    assert new.worksheet_status=='pending' and db.scalar(select(func.count(Customer.id)))==1


def test_missing_rows_are_never_deleted(sample):
    db,item=sample
    content=edited(ex.export_editable(db),lambda wb:wb['Orders'].delete_rows(2))
    result=ex.inspect_upload(db,content)
    assert result['missing']==1 and not result['rows']
    apply(db,content)
    assert db.get(OrderItem,item.id) is not None


def test_resaving_a_used_download_cannot_duplicate_new_rows(sample):
    db,item=sample
    def change(wb):
        wb['Orders'].append([None,'SS26','Test brand','AA','Alina','NEW-1','Coat','black','M',1,200,35,None,'Not completed'])
    content=edited(ex.export_editable(db),change)
    apply(db,content)
    resaved=edited(content,lambda wb:wb['Start here'].__setitem__('A20','A harmless note'))
    assert content!=resaved
    with pytest.raises(ValueError,match='already been applied'):ex.inspect_upload(db,resaved)
    assert db.scalar(select(func.count(OrderItem.id)))==2


def test_stale_changes_conflict_but_unchanged_rows_do_not(sample):
    db,item=sample
    content=ex.export_editable(db)
    item.color='newer edit';db.commit()
    assert ex.inspect_upload(db,content)['unchanged']==1
    stale=edited(content,lambda wb:wb['Orders'].__setitem__('I2','40'))
    assert 'changed in the app' in ex.inspect_upload(db,stale)['errors'][0]['message']
    with pytest.raises(ValueError,match='conflicts'):apply(db,stale)
    db.rollback()
    assert item.color=='newer edit' and item.size=='38'


@pytest.mark.parametrize('cell,value,message',[
    ('J2',-1,'non-negative'),('K2','NaN','non-negative'),('I2','=1+1','formulas'),
    ('B2','SS28','season and brand'),('N2','green','Choose In process'),('O2',999,'Group ID'),
    ('E2','Wrong person','does not match'),('A2',999,'ID is not'),
])
def test_invalid_rows_never_apply(sample,cell,value,message):
    db,item=sample
    content=edited(ex.export_editable(db),lambda wb:wb['Orders'].__setitem__(cell,value))
    assert message in ex.inspect_upload(db,content)['errors'][0]['message']
    with pytest.raises(ValueError):apply(db,content)
    db.rollback()
    assert item.name=='Dress' and db.scalar(select(func.count(ImportBatch.id)))==0


def test_credit_ownership_cannot_be_moved_in_excel(sample):
    db,item=sample
    account=SellerCreditAccount(item_id=item.id,group_id=item.purchase_group_id,customer_id=item.customer_id)
    db.add(account);db.commit()
    def change(wb):
        wb['Orders']['D2']='BB';wb['Orders']['E2']='Other customer'
    content=edited(ex.export_editable(db),change)
    assert 'seller-credit history' in ex.inspect_upload(db,content)['errors'][0]['message']


def test_all_or_nothing_when_one_row_is_invalid(sample):
    db,item=sample
    def change(wb):
        wb['Orders']['H2']='valid change'
        wb['Orders'].append([None,'SS27','Brand',None,None,None,'Missing quantity'])
    content=edited(ex.export_editable(db),change)
    preview=ex.inspect_upload(db,content)
    assert len(preview['rows'])==1 and len(preview['errors'])==1
    with pytest.raises(ValueError):apply(db,content)
    db.rollback();assert item.color is None


def test_duplicate_existing_ids_blocked(sample):
    db,item=sample
    content=edited(ex.export_editable(db),lambda wb:wb['Orders'].append([c.value for c in wb['Orders'][2]]))
    assert 'Duplicate ID' in ex.inspect_upload(db,content)['errors'][0]['message']


def test_excel_pages_and_preview_leave_data_unchanged(sample):
    db,item=sample
    app.dependency_overrides[get_db]=lambda:db
    try:
        with TestClient(app) as client:
            assert client.get('/excel').status_code==200
            workbook=client.get('/excel/download').content
            modified=edited(workbook,lambda wb:wb['Orders'].__setitem__('G2','Updated dress'))
            page=client.post('/excel/preview',files={'file':('edited.xlsx',modified)})
            assert page.status_code==200 and 'Apply 1 reviewed changes' in page.text
            assert item.name=='Dress'
            assert '<img ' not in client.get('/worksheet').text
            assert client.get(f'/worksheet/items/{item.id}').status_code==200
    finally:app.dependency_overrides.clear()
