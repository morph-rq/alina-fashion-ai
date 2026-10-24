"""Source formatting and workflow colours for the local workbook editor."""
from __future__ import annotations
import colorsys
import io
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from xml.etree import ElementTree
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

STATUSES = {'ordered': ('Ordered / in progress', '#FFFFFF'),
            'ready': ('Ready · factory finished, awaiting shipment', '#FFF2CC'),
            'shipped': ('Done and shipped', '#C6E0B4'),
            'cannot': ('Cannot be done', '#FFC7CE'),
            'mixed': ('Mixed source colours · review', '#EEE8F4')}


def theme_palette(book):
    if not book.loaded_theme:
        return []
    ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    scheme = ElementTree.fromstring(book.loaded_theme).find('.//a:clrScheme', ns)
    if scheme is None:
        return []
    values = [next(iter(c)).get('lastClr') or next(iter(c)).get('val') for c in scheme]
    values[:4] = [values[1],values[0],values[3],values[2]]
    return values


def css_color(color, palette):
    if color is None:
        return None
    rgb = None
    if color.type == 'rgb': rgb = str(color.rgb)[-6:]
    elif color.type == 'indexed' and 0 <= color.indexed < len(COLOR_INDEX): rgb = COLOR_INDEX[color.indexed][-6:]
    elif color.type == 'theme' and 0 <= color.theme < len(palette): rgb = palette[color.theme]
    if not rgb or len(rgb)!=6:
        return None
    try: channels = [int(rgb[i:i+2],16) for i in (0,2,4)]
    except ValueError: return None
    # Excel tint is defined on the luminance component in HLS colour space.
    hue, light, saturation = colorsys.rgb_to_hls(*(c/255 for c in channels))
    tint = color.tint or 0
    light = light*(1+tint) if tint < 0 else light*(1-tint)+tint
    channels = [round(c*255) for c in colorsys.hls_to_rgb(hue,light,saturation)]
    return '#'+''.join(f'{c:02X}' for c in channels)


def colour_status(colour):
    if not colour: return None
    r,g,b = (int(colour[i:i+2],16)/255 for i in (1,3,5))
    hue,saturation,value = colorsys.rgb_to_hsv(r,g,b)
    if saturation < .16 or value < .3: return None
    if hue < .055 or hue > .95: return 'cannot'
    if .105 <= hue <= .19: return 'ready'
    if .20 <= hue <= .48: return 'shipped'
    return None


def read_formatting(path, bounds=None):
    book=load_workbook(path,read_only=True,data_only=True)
    palette=theme_palette(book); result={}
    try:
        for sheet in book.worksheets:
            if sheet.title=='_Exchange':continue
            if (sheet.max_row or 0)>50000 or (sheet.max_column or 0)>200:continue
            styles={}; cell_count=0
            limit = (bounds or {}).get(sheet.title, (sheet.max_row or 0, sheet.max_column or 0))
            for row in sheet.iter_rows(max_row=max(1,limit[0]),max_col=max(1,limit[1])):
                cell_count+=len(row)
                if cell_count>500000:break
                for cell in row:
                    if not getattr(cell,'has_style',False):continue
                    fill=css_color(cell.fill.fgColor,palette) if cell.fill.patternType=='solid' else None
                    font=cell.font
                    style={'fill':fill,'font':css_color(font.color,palette),'bold':bool(font.b),'italic':bool(font.i),
                           'number_format':cell.number_format,'align':cell.alignment.horizontal or 'general',
                           'wrap':bool(cell.alignment.wrap_text)}
                    styles[f'{cell.row}:{cell.column}']=style
            result[sheet.title]={'styles':styles,'format_version':2}
    finally:book.close()
    # Dimensions/merges are small XML metadata; no image decoding is needed.
    from zipfile import ZipFile
    import posixpath
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(path) as archive:
        workbook=ElementTree.fromstring(archive.read('xl/workbook.xml'))
        rels=ElementTree.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets={r.get('Id'):r.get('Target') for r in rels}
        for sheet in workbook.find('s:sheets',ns):
            name=sheet.get('name')
            if name not in result:continue
            rid=sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target=targets.get(rid,'')
            target=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
            xml=ElementTree.fromstring(archive.read(target))
            cols={}
            for col in xml.findall('s:cols/s:col',ns):
                for i in range(int(col.get('min')),min(int(col.get('max')),200)+1):
                    cols[str(i)]={'width':float(col.get('width','13')),'hidden':col.get('hidden')=='1'}
            result[name].update(column_dimensions=cols, merges=[x.get('ref') for x in xml.findall('s:mergeCells/s:mergeCell',ns)])
    return result


def row_status(sheet, number):
    override=sheet.get('row_meta',{}).get(str(number),{}).get('status')
    if override:return override
    row=sheet['rows'][number-1]
    if number < sheet['rule'].get('start',1):return 'ordered'
    if not any(decimal_value(v) is not None for v in row if v is not None):return 'ordered'
    # Prefer product descriptors; shipment and accounting columns can have independent fills.
    mapping=sheet['rule'].get('mapping',{})
    candidates=[int(mapping[k]) for k in ('article','product_name') if mapping.get(k)]
    styles=sheet.get('styles',{})
    found={colour_status(styles.get(f'{number}:{col}',{}).get('fill')) for col in candidates}
    found.discard(None)
    if not found:
        found={colour_status(styles.get(f'{number}:{col}',{}).get('fill')) for col,value in enumerate(row,1) if value is not None}
        found.discard(None)
    return next(iter(found)) if len(found)==1 else 'mixed' if found else 'ordered'


def decimal_value(value):
    if value is None or isinstance(value,bool):return None
    try:
        result=Decimal(str(value).replace(' ','').replace(',','.'))
        return result if result.is_finite() else None
    except InvalidOperation:return None


def row_pricing(sheet, number):
    rule=sheet['rule'];mapping=rule.get('mapping',{});row=sheet['rows'][number-1]
    def value(field):
        col=int(mapping.get(field) or 0)
        return row[col-1] if 0<col<=len(row) else None
    markup=decimal_value(sheet.get('row_meta',{}).get(str(number),{}).get('markup',rule.get('markup')))
    price=decimal_value(value('stated_price'));qty=decimal_value(value('quantity'))
    result={'markup':str(markup) if markup is not None else '', 'final_unit':None,'final_line':None}
    if not rule.get('start',1)<=number<=rule.get('end',len(sheet['rows'])):return result
    if not (value('article') or value('product_name')):return result
    if price is None or price<0 or rule.get('price_basis') not in ('unit','line') or rule.get('price_type') not in ('wholesale','final'):return result
    unit=price if rule['price_basis']=='unit' else price/qty if qty and qty>0 else None
    if unit is None:return result
    if rule['price_type']=='wholesale':
        if markup is None:return result
        unit*=1+markup/100
    unit=unit.quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
    result['final_unit']=str(unit)
    if qty and qty>0:result['final_line']=str((unit*qty).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
    return result


def filter_rows(sheet,status='',query='',customer=''):
    result=[]
    mapping=sheet['rule'].get('mapping',{})
    for number,row in enumerate(sheet['rows'],1):
        if status and row_status(sheet,number)!=status:continue
        if query and query.casefold() not in ' '.join(str(v) for v in row if v is not None).casefold():continue
        if customer:
            col=int(mapping.get('customer_name') or 0)
            name=row[col-1] if 0<col<=len(row) else sheet['rule'].get('customer','')
            if customer.casefold() not in str(name or '').casefold():continue
        result.append(number)
    return result


def build_export(state,indices,selected=None,status='',query='',columns=None,customer='',extras=False,headers=True):
    book=Workbook();book.remove(book.active)
    count=0
    for index in indices:
        sheet=state['sheets'][index]
        wanted=filter_rows(sheet,status,query,customer)
        if selected is not None:wanted=[n for n in wanted if n in selected]
        if (selected is not None or status or query or customer) and not wanted:continue
        prefix=list(range(1,sheet['rule'].get('start',1))) if headers else []
        wanted=sorted(set(wanted+prefix))
        cols=columns or list(range(1,sheet['columns']+1))
        if any(not 1<=c<=sheet['columns'] for c in cols):raise ValueError('An export column is outside the selected worksheet.')
        ws=book.create_sheet(sheet['name']);row_map={n:i+1 for i,n in enumerate(wanted)}
        labels=['ALINA status','ALINA markup %','ALINA final unit EUR','ALINA final line EUR','ALINA note']
        header=sheet['rule'].get('start',1)-1
        for number in wanted:
            row=sheet['rows'][number-1];out=row_map[number]
            status_value=row_status(sheet,number)
            explicit=sheet.get('row_meta',{}).get(str(number),{}).get('status')
            for target,col in enumerate(cols,1):
                value=row[col-1] if col<=len(row) else None
                cell=ws.cell(out,target,value)
                if isinstance(value,str):cell.data_type='s'
                style=sheet.get('styles',{}).get(f'{number}:{col}',{})
                fill=STATUSES[explicit][1] if explicit else style.get('fill')
                if fill:cell.fill=PatternFill('solid',fgColor=fill.lstrip('#'))
                cell.font=Font(name='Calibri',size=11,bold=style.get('bold',False),italic=style.get('italic',False),color=(style.get('font') or '#172033').lstrip('#'))
                cell.number_format=style.get('number_format','General')
                cell.alignment=Alignment(horizontal=style.get('align','general'),wrap_text=style.get('wrap',False))
            if extras:
                pricing=row_pricing(sheet,number)
                if number==header:
                    values=labels
                elif number>=sheet['rule'].get('start',1):
                    values=[STATUSES[status_value][0],float(pricing['markup']) if pricing['markup'] else None,
                            float(pricing['final_unit']) if pricing['final_unit'] else None,float(pricing['final_line']) if pricing['final_line'] else None,
                            sheet.get('row_meta',{}).get(str(number),{}).get('note','')]
                else:values=[]
                for offset,value in enumerate(values,len(cols)+1):
                    c=ws.cell(out,offset,value)
                    if isinstance(value,str):c.data_type='s'
                    c.fill=PatternFill('solid',fgColor=STATUSES[status_value][1].lstrip('#') if number!=header else 'E9EEF5')
                    if offset in (len(cols)+3,len(cols)+4):c.number_format='#,##0.00'
        if extras and header not in row_map:
            ws.insert_rows(1)
            for target,col in enumerate(cols,1):ws.cell(1,target,get_column_letter(col))
            for offset,label in enumerate(labels,len(cols)+1):ws.cell(1,offset,label)
        for target,col in enumerate(cols,1):
            ws.column_dimensions[get_column_letter(target)].width=min(60,max(9,sheet.get('column_dimensions',{}).get(str(col),{}).get('width',16)))
        for target in range(len(cols)+1,len(cols)+6):ws.column_dimensions[get_column_letter(target)].width=23
        ws.freeze_panes='A'+str(min(max(2,header+1),5))
        ws.sheet_view.showGridLines=True
        if wanted:ws.auto_filter.ref=ws.dimensions
        count+=len(wanted)
    if not book.worksheets:raise ValueError('No rows match your export selection.')
    stream=io.BytesIO();book.save(stream)
    return stream.getvalue()
