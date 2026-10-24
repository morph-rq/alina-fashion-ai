"""Exports for an exact selection: customer documents, internal (factory) reports, original-with-prices, ZIP.

Audiences are kept apart on purpose:
- ``customer``: one Excel/PDF per customer through ``customer_documents.generate`` — only that
  customer's selected lines, customer prices, no factory figures; the confidentiality scan runs on
  every file. When a customer's selection is their whole order, the full document (with payments
  and balance) is produced; otherwise a "selected positions" document with only its own totals.
- ``internal``: one report for the owner/factory with every selected allocation, customer codes,
  base and customer prices, and the factory figures per brand from the settlement engine.
- ``original``: the stored factory file, byte-for-byte except for two added columns
  ("Клиент · цена", "Клиент · сумма") on the rows of the selection — images, formulas,
  formatting and other sheets are untouched (``xlsx_stream.patch_workbook``). Internal only.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.calculations import ZERO, item_amounts, money
from app import customer_documents
from app.config import ROOT_DIR
from app.customer_documents import _image_file, _thumb
from app.factory_ledger import active_items, season_ledger
from app.models import Customer, GeneratedDocument, ImportBatch, OrderItem, Season
from app.selection import Selection, unit_price

STATUS_RU = {"Ordered": "Заказано", "Confirmed": "Подтверждено", "Production": "В производстве", "Ready": "Готово",
             "Assigned to shipment": "В отгрузке", "In transit": "В пути", "Arrived": "Прибыло", "Delivered": "Доставлено",
             "Cancelled": "Отменено"}


def slug(text: str) -> str:
    from app.factory_import import TRANSLIT
    latin = "".join(TRANSLIT.get(ch, ch) for ch in (text or "").upper())
    return re.sub(r"[^A-Za-z0-9]+", "_", latin).strip("_")[:60] or "export"


def customer_label(selection: Selection) -> str | None:
    """A neutral description of the selection for a customer document — never another customer's code."""
    f = selection.filters
    parts = []
    names = {"ready": "Готовые", "not_ready": "Не готовые", "cancelled": "Отменённые", "shipped": "Отгруженные"}
    if f.readiness:
        parts.append(names.get(f.readiness, STATUS_RU.get(f.readiness, f.readiness)))
    if f.categories:
        from app.selection import CATEGORIES
        parts.append(", ".join(CATEGORIES[c][0] for c in f.categories if c in CATEGORIES))
    if f.sizes:
        parts.append("размер " + ", ".join(f.sizes))
    if f.colors:
        parts.append("цвет " + ", ".join(f.colors))
    if f.item_ids and not parts:
        parts.append("выбранные позиции")
    if f.rows and not parts:
        parts.append("строки файла " + ", ".join(map(str, f.rows[:10])))
    if f.price_min or f.price_max:
        parts.append("по цене")
    return " · ".join(parts) or None


def plan(db: Session, selection: Selection) -> list[dict]:
    """One spec per customer in the selection: exact item ids, or None when it is the customer's whole order."""
    if selection.season is None:
        return []
    specs = []
    brands = sorted({i.order.brand.name for i in selection.items})
    for code in selection.customer_codes:
        customer = db.scalar(select(Customer).where(Customer.code == code))
        if customer is None or customer.is_owner:
            continue
        chosen = sorted(i.id for i in selection.items if i.customer_id == customer.id)
        scope = [i.id for o in selection.season.orders if not selection.filters.brands or o.brand.name in brands
                 for i in active_items(o) if i.customer_id == customer.id]
        whole = sorted(scope) == chosen
        specs.append({"customer": code, "season": selection.season.name,
                      "brands": brands if selection.filters.brands else None,
                      "item_ids": None if whole else chosen, "label": None if whole else customer_label(selection),
                      "items": len(chosen)})
    return specs


def make_customer_documents(db: Session, spec: dict, formats: tuple[str, ...], actor: str) -> list[GeneratedDocument]:
    customer = db.scalar(select(Customer).where(Customer.code == spec["customer"]))
    season = db.scalar(select(Season).where(Season.name == spec["season"]))
    return customer_documents.generate(db, customer, season, brands=spec.get("brands"), item_ids=spec.get("item_ids"),
                    label=spec.get("label"), formats=formats, actor=actor)


# --- Internal report ---------------------------------------------------------------------

def internal_snapshot(db: Session, selection: Selection) -> dict:
    rows, per_customer, per_brand = [], {}, {}
    for item in selection.items:
        amounts = item_amounts(item)
        code = item.customer.code if item.customer else "—"
        brand = item.order.brand.name
        rows.append({"id": item.id, "brand": brand, "customer": code, "name": item.name or "", "article": item.article or "",
                     "color": item.color or "", "size": item.size or "", "quantity": item.quantity,
                     "base": str(amounts.base), "markup": str(Decimal(item.markup_percentage or 0) * 100),
                     "unit": str(unit_price(item)), "total": str(amounts.final), "status": STATUS_RU.get(item.status, item.status),
                     "image": item.image_path, "source": f"{item.source_sheet or ''} {item.source_row or ''}".strip()})
        c = per_customer.setdefault(code, {"items": 0, "base": ZERO, "total": ZERO})
        c["items"] += 1
        c["base"] += amounts.base
        c["total"] += amounts.final
        per_brand.setdefault(brand, item.order_id)
    factory = []
    if selection.season is not None:
        ledger = season_ledger(db, selection.season)
        for view in ledger.orders:
            if view.order.brand.name in per_brand:
                factory.append({"brand": view.order.brand.name, "factory_base": str(view.factory_base),
                                "factory_total": str(view.factory_total), "deposit_rate": str(view.deposit_rate),
                                "deposit": str(view.deposit_required), "ready_balance": str(view.ready_factory_balance)})
    return {"title": selection.summary, "season": selection.season.name if selection.season else "",
            "rows": rows, "customers": {k: {kk: str(money(vv)) if isinstance(vv, Decimal) else vv for kk, vv in v.items()}
                                        for k, v in sorted(per_customer.items())},
            "factory": factory, "base_total": str(money(sum((Decimal(r["base"]) for r in rows), ZERO))),
            "customer_total": str(money(sum((Decimal(r["total"]) for r in rows), ZERO))),
            "note": "Суммы фабрики — по всему заказу бренда (условия фабрики считаются на заказ целиком)."}


def internal_excel(snap: dict, root: Path = ROOT_DIR) -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = "Позиции"
    ink = "1C2130"
    ws["A1"] = "ВНУТРЕННИЙ ОТЧЁТ — не для клиентов"
    ws["A1"].font = Font(bold=True, size=13, color="B42318")
    ws["A2"] = snap["title"]
    heads = ["Фото", "Бренд", "Клиент", "Наименование", "Артикул", "Цвет", "Размер", "Кол-во", "База", "Наценка %",
             "Цена клиента", "Сумма клиента", "Статус", "Источник"]
    for col, label in enumerate(heads, 1):
        cell = ws.cell(4, col, label)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=ink)
        cell.alignment = Alignment(wrap_text=True)
    money_fmt = '#,##0.00 "€"'
    r = 5
    for row in snap["rows"]:
        values = [None, row["brand"], row["customer"], row["name"], row["article"], row["color"], row["size"], row["quantity"],
                  float(row["base"]), float(row["markup"]) / 100, float(row["unit"]), None, row["status"], row["source"]]
        for col, value in enumerate(values, 1):
            ws.cell(r, col, value if value != "" else None)
        ws.cell(r, 12, f"=K{r}*H{r}")
        for col in (9, 11, 12):
            ws.cell(r, col).number_format = money_fmt
        ws.cell(r, 10).number_format = "0.##%"
        path = _image_file(row["image"], root)
        thumb = _thumb(path) if path else None
        if thumb is not None:
            picture = XLImage(thumb)
            ratio = picture.width / picture.height if picture.height else 1
            picture.height, picture.width = 60, int(60 * ratio)
            ws.add_image(picture, f"A{r}")
            ws.row_dimensions[r].height = 48
        r += 1
    ws.cell(r, 8, "Итого").font = Font(bold=True)
    ws.cell(r, 9, f"=SUM(I5:I{r - 1})").number_format = money_fmt
    ws.cell(r, 12, f"=SUM(L5:L{r - 1})").number_format = money_fmt
    for col, width in zip("ABCDEFGHIJKLMN", (12, 16, 9, 30, 18, 12, 8, 7, 12, 10, 13, 14, 14, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B5"
    summary = book.create_sheet("Сводка")
    summary.append(["Клиент", "Позиций", "База", "Сумма клиента"])
    for code, c in snap["customers"].items():
        summary.append([code, c["items"], float(c["base"]), float(c["total"])])
    summary.append([])
    summary.append(["Бренд", "База фабрики", "Фабрике всего", "Предоплата фабрике %", "Предоплата фабрике", "За готовые"])
    for f in snap["factory"]:
        summary.append([f["brand"], float(f["factory_base"]), float(f["factory_total"]), float(f["deposit_rate"]),
                        float(f["deposit"]), float(f["ready_balance"])])
    summary.append([])
    summary.append([snap["note"]])
    for row in summary.iter_rows():
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = money_fmt if cell.column != 4 or cell.value > 1 else "0%"
    for col, width in zip("ABCDEF", (22, 14, 16, 18, 18, 14)):
        summary.column_dimensions[col].width = width
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def internal_pdf(snap: dict, root: Path = ROOT_DIR) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from PIL import Image as PILImage
    from app.pdf_report import register_fonts
    regular, bold = register_fonts()
    out = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title="Внутренний отчёт", author="Alina Fashion")
    body = ParagraphStyle("b", fontName=regular, fontSize=7.5, leading=9.5)
    num = ParagraphStyle("n", parent=body, alignment=TA_RIGHT)
    head = ParagraphStyle("h", fontName=bold, fontSize=14, leading=18)
    warn = ParagraphStyle("w", parent=body, textColor=colors.HexColor("#B42318"), fontName=bold)
    eur = lambda v: f"€ {Decimal(v):,.2f}".replace(",", " ")  # noqa: E731
    story = [Paragraph("ВНУТРЕННИЙ ОТЧЁТ — не для клиентов", warn), Paragraph(escape(snap["title"]), head), Spacer(1, 3 * mm)]
    data = [[Paragraph(f"<b>{h}</b>", body) for h in ("Фото", "Бренд", "Клиент", "Изделие", "Артикул / цвет", "Размер", "Шт.",
                                                       "База", "%", "Цена", "Сумма", "Статус")]]
    for row in snap["rows"]:
        path = _image_file(row["image"], root)
        thumb = _thumb(path, (100, 125)) if path else None
        if thumb is not None:
            with PILImage.open(thumb) as probe:
                ratio = probe.width / probe.height if probe.height else 1
            thumb.seek(0)
            pic = Image(thumb, width=min(12 * mm, 15 * mm * ratio), height=min(15 * mm, 12 * mm / ratio))
        else:
            pic = Paragraph("—", body)
        data.append([pic, Paragraph(escape(row["brand"]), body), Paragraph(escape(row["customer"]), body),
                     Paragraph(escape(row["name"]), body), Paragraph(escape(" · ".join(x for x in (row["article"], row["color"]) if x)), body),
                     Paragraph(escape(str(row["size"])), body), Paragraph(str(row["quantity"]), num), Paragraph(eur(row["base"]), num),
                     Paragraph(f"{Decimal(row['markup']):g}%", num), Paragraph(eur(row["unit"]), num), Paragraph(eur(row["total"]), num),
                     Paragraph(escape(row["status"]), body)])
    data.append(["", "", "", "", "", "", Paragraph("<b>Итого</b>", body), Paragraph(f"<b>{eur(snap['base_total'])}</b>", num), "", "",
                 Paragraph(f"<b>{eur(snap['customer_total'])}</b>", num), ""])
    widths = [16, 26, 14, 58, 44, 14, 9, 22, 11, 22, 24, 22]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), .3, colors.HexColor("#DADFD6")),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDEFEA"))]))
    story.append(table)
    if snap["factory"]:
        story += [Spacer(1, 4 * mm), Paragraph("<b>Фабрика (по заказу бренда целиком)</b>", body)]
        rows = [[Paragraph(f"<b>{h}</b>", body) for h in ("Бренд", "База", "Фабрике всего", "Предоплата", "За готовые")]]
        rows += [[Paragraph(escape(f["brand"]), body), Paragraph(eur(f["factory_base"]), num), Paragraph(eur(f["factory_total"]), num),
                  Paragraph(eur(f["deposit"]), num), Paragraph(eur(f["ready_balance"]), num)] for f in snap["factory"]]
        story.append(Table(rows, colWidths=[50 * mm, 30 * mm, 30 * mm, 30 * mm, 30 * mm], hAlign="LEFT"))
    doc.build(story)
    return out.getvalue()


def make_internal(db: Session, selection: Selection, formats: tuple[str, ...], actor: str, root: Path | None = None) -> list[GeneratedDocument]:
    root = root or ROOT_DIR
    snap = internal_snapshot(db, selection)
    if not snap["rows"]:
        raise ValueError("В выборке нет позиций.")
    folder = root / "data" / "documents" / (selection.season.name if selection.season else "all") / "_internal"
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc)
    builders = {"xlsx": internal_excel, "pdf": internal_pdf}
    made = []
    for kind in ("xlsx", "pdf"):
        if kind not in formats:
            continue
        data = builders[kind](snap, root)
        path = folder / f"internal_{slug(selection.summary.removeprefix('Выбрано: '))}_{stamp:%Y%m%d_%H%M%S}.{kind}"
        path.write_bytes(data)
        doc = GeneratedDocument(kind="internal", file_format=kind, customer_id=None,
                                season_id=selection.season.id if selection.season else None,
                                brand_filter=selection.summary[:200], status="preview", version=1,
                                stored_path=str(path.relative_to(root)), sha256=hashlib.sha256(data).hexdigest(),
                                data_fingerprint=hashlib.sha256(json.dumps(snap, sort_keys=True).encode()).hexdigest(),
                                snapshot_json=json.dumps({"filters": selection.filters.to_dict(), "items": [r["id"] for r in snap["rows"]]}))
        db.add(doc)
        made.append(doc)
    db.commit()
    return made


# --- Original file with customer prices ---------------------------------------------------

def original_with_prices(db: Session, selection: Selection, actor: str, root: Path | None = None) -> list[GeneratedDocument]:
    root = root or ROOT_DIR
    from app.xlsx_stream import patch_workbook
    by_file: dict[str, list[OrderItem]] = {}
    for item in selection.items:
        if item.source_file and item.source_sheet and item.source_row:
            by_file.setdefault(item.source_file, []).append(item)
    if not by_file:
        raise ValueError("У выбранных позиций нет исходного файла фабрики.")
    made = []
    for filename, items in by_file.items():
        batch = db.scalar(select(ImportBatch).where(ImportBatch.filename == filename, ImportBatch.stored_path.is_not(None))
                          .order_by(ImportBatch.id.desc()))
        if batch is None:
            continue
        source = Path(batch.stored_path)
        source = source if source.is_absolute() else root / source
        if not source.exists() or source.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        book = load_workbook(source, read_only=True)
        widths = {}
        for sheet in {i.source_sheet for i in items}:
            if sheet in book.sheetnames:
                widths[sheet] = book[sheet].max_column or 1
        book.close()
        edits: dict[str, dict[tuple[int, int], object]] = {}
        rows: dict[tuple[str, int], list[OrderItem]] = {}
        for item in items:
            if item.source_sheet in widths:
                rows.setdefault((item.source_sheet, item.source_row), []).append(item)
        for (sheet, row), group in rows.items():
            col = widths[sheet] + 1
            cells = edits.setdefault(sheet, {})
            cells[(row, col)] = "; ".join(f"{i.customer.code if i.customer else '?'} {i.size or ''} {unit_price(i):.2f} €".replace("  ", " ")
                                          for i in group)
            cells[(row, col + 1)] = float(sum((item_amounts(i).final for i in group), ZERO))
        for sheet, cells in edits.items():
            first = min(r for r, _ in cells)
            cells.setdefault((max(first - 1, 1), widths[sheet] + 1), "Клиент · цена")
            cells.setdefault((max(first - 1, 1), widths[sheet] + 2), "Клиент · сумма")
        folder = root / "data" / "documents" / (selection.season.name if selection.season else "all") / "_internal"
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"{slug(Path(filename).stem)}_с_ценами_{datetime.now():%Y%m%d_%H%M%S}{source.suffix.lower()}"
        patch_workbook(source, target, edits)
        data = target.read_bytes()
        doc = GeneratedDocument(kind="original", file_format=source.suffix.lower().lstrip("."), customer_id=None,
                                season_id=selection.season.id if selection.season else None,
                                brand_filter=f"Оригинал с ценами: {filename}"[:200], status="preview", version=1,
                                stored_path=str(target.relative_to(root)), sha256=hashlib.sha256(data).hexdigest(),
                                data_fingerprint=hashlib.sha256(data).hexdigest(), snapshot_json=json.dumps({"items": [i.id for i in items]}))
        db.add(doc)
        made.append(doc)
    db.commit()
    if not made:
        raise ValueError("Исходный Excel фабрики не найден среди сохранённых оригиналов.")
    return made


def bundle(documents: list[GeneratedDocument], root: Path | None = None) -> bytes:
    root = root or ROOT_DIR
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for doc in documents:
            path = root / doc.stored_path
            if path.exists():
                archive.write(path, Path(doc.stored_path).name)
    return out.getvalue()
