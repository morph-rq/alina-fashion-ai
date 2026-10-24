"""Factory vs customer pricing, deposit reconciliation, factory import and the shared action layer."""
from __future__ import annotations

import re
import zipfile
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

import factory_fixture
from app import factory_import, ledger_actions as la
from app.database import Base
from app.factory_ledger import season_ledger
from app.models import (AllocationReview, AuditLog, Brand, Customer, CustomerPricingRule, FactoryAgreement,
                        FactoryRatePortion, Order, OrderItem, Payment, Season, SettlementSnapshot, Shipment)
from app.settlement import CustomerLine, Portion, factory_price, marked_up, settle
from app.services import country_is_transport_exempt
from app.calculations import transport_fee


# --- A. The interview's worked example ------------------------------------------------

def interview(aa=D("0.35"), ids=D("0.35"), ss=D("0.35"), payments=()):
    return settle([CustomerLine("AA", D(20), aa), CustomerLine("IDS", D(20), ids), CustomerLine("SS", D(30), ss),
                   CustomerLine("Alina", D(30))], [Portion(D("0.30"))], owner="Alina", payments=payments)


def test_a_exact_interview_scenario():
    s = interview()
    assert s.factory.base == D("100.00") and s.factory.total == D("130.00")
    assert s.factory_deposit_required == D("65.00")
    c = s.customers
    assert (c["AA"].total, c["IDS"].total, c["SS"].total) == (D("27.00"), D("27.00"), D("40.50"))
    assert (c["AA"].deposit_due, c["IDS"].deposit_due, c["SS"].deposit_due) == (D("13.50"), D("13.50"), D("20.25"))
    assert s.external_deposits_expected == D("47.25")
    assert s.owner_net_deposit_planned == D("17.75")        # not the spoken 17.50
    assert s.owner_nominal_cost == D("39.00") and s.owner_nominal_deposit == D("19.50")
    assert all(isinstance(v.total, D) for v in c.values())


def test_customer_rate_is_on_the_base_not_on_the_factory_price():
    assert marked_up(D(20), D("0.35")) == D("27.00")
    assert marked_up(marked_up(D(20), D("0.30")), D("0.35")) == D("35.10")  # the mistake we must not make
    assert interview().customers["AA"].total == D("27.00")


# --- B. Changing one customer's rate ------------------------------------------------

def test_b_changing_aa_rate_changes_only_aa():
    before, after = interview(), interview(aa=D("0.30"))
    assert after.factory.total == before.factory.total == D("130.00")
    assert after.factory_deposit_required == D("65.00")
    assert after.customers["AA"].total == D("26.00") and after.customers["AA"].deposit_due == D("13.00")
    assert after.customers["IDS"].total == before.customers["IDS"].total
    assert after.customers["SS"].total == before.customers["SS"].total
    assert after.external_deposits_expected == D("46.75") and after.owner_net_deposit_planned == D("18.25")


# --- C. Mixed factory percentage ------------------------------------------------------

def test_c_half_at_25_half_at_30():
    price = factory_price(D("1000"), [Portion(D("0.25"), share=D("0.5")), Portion(D("0.30"), share=D("0.5"))])
    assert price.total == D("1275.00")
    assert (price.total * D("0.5")).quantize(D("0.01")) == D("637.50")
    assert [p.base for p in price.portions] == [D("500.00"), D("500.00")]


def test_odd_cent_split_and_explicit_sections_and_roundup():
    price = factory_price(D("100.01"), [Portion(D("0.25"), share=D("0.5")), Portion(D("0.30"), share=D("0.5"))])
    assert sum(p.base for p in price.portions) == D("100.01")
    sections = factory_price(D("1000"), [Portion(D("0.23"), base=D("600"), label="RTW"), Portion(D("0.25"), base=D("400"), label="BAGS")])
    assert sections.total == D("1238.00")
    with pytest.raises(ValueError):
        factory_price(D("1000"), [Portion(D("0.23"), base=D("600"))])     # 400 not covered: never dropped silently
    assert marked_up(D("100.10"), D("0.35"), "euro_up") == D("136.10")    # ROUNDUP(35.035) = 36


# --- G. Payment reconciliation (pure) -------------------------------------------------

def test_g_planned_versus_confirmed_and_no_double_counting():
    s = interview(payments=[{"payer": "AA", "amount": D("13.50"), "to": "factory", "confirmed": True},
                            {"payer": "IDS", "amount": D("13.50"), "to": "factory", "confirmed": False},
                            {"payer": "Alina", "amount": D("17.75"), "to": "factory", "confirmed": True}])
    assert s.confirmed_factory_payments == D("31.25")                    # promised IDS money is not counted
    assert s.factory_deposit_outstanding == D("33.75")
    assert s.customers["IDS"].expected == D("13.50") and s.customers["IDS"].deposit_outstanding == D("13.50")
    assert s.owner_transfer_now == D("0.00")
    over = interview(payments=[{"payer": "Alina", "amount": D("200"), "to": "factory"}])
    assert over.factory_overpaid == D("70.00") and over.factory_balance_outstanding == D("0.00")
    forwarded = interview(payments=[{"payer": "SS", "amount": D("20.25"), "to": "owner"}])
    assert forwarded.confirmed_factory_payments == D("0.00") and forwarded.received_by_owner_to_forward == D("20.25")
    assert forwarded.customers["SS"].deposit_outstanding == D("0.00")


def test_overfunded_plan_is_credit_not_negative_transfer():
    s = settle([CustomerLine("AA", D(100), D("0.50")), CustomerLine("Alina", D(1))], [Portion(D("0.10"))], owner="Alina")
    assert s.owner_net_deposit_planned < 0 and s.overfunded > 0 and s.owner_transfer_now == D("0.00")


# --- Database: import, allocation, actions --------------------------------------------

@pytest.fixture()
def db(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'ledger.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    monkeypatch.setattr(factory_import, "ORIGINALS_DIR", tmp_path / "originals")
    monkeypatch.setattr(factory_import, "IMAGES_DIR", tmp_path / "data" / "images")
    monkeypatch.setattr(factory_import, "ROOT_DIR", tmp_path)
    with sessionmaker(bind=engine, expire_on_commit=False)() as session:
        for code, name in (("AA", None), ("IDS", None), ("SS", None), ("ALINA", "Alina")):
            session.add(Customer(code=code, name=name, is_owner=code == "ALINA"))
        session.commit()
        yield session


def imported(db, tmp_path, decisions=None):
    path = factory_fixture.build(tmp_path / "Dior cruise SS27 fixture.xlsx")
    analysis = factory_import.analyze_workbook(path, path.name, {"AA", "IDS", "SS", "ALINA"})
    original = factory_import.store_original(path, analysis.sha256, path.name)
    result = factory_import.commit_import(db, analysis, original, {"season": "SS27", "rates_confirmed": True, **(decisions or {})})
    return analysis, result


def test_e_import_reads_layout_images_rates_and_flags_ambiguity(db, tmp_path):
    analysis, result = imported(db, tmp_path)
    sheet = next(s for s in analysis.sheets if s.name == "Dior cruise")
    assert [s.kind for s in analysis.sheets] == ["summary", "order"]
    assert sheet.brand == "Dior cruise" and [x.rate for x in sheet.sections] == ["0.3"]
    assert D(sheet.computed_base) == D(factory_fixture.BASE_TOTAL) == D("8100")
    assert all(item.image_part for item in sheet.items)
    assert sheet.items[0].fill == "#C6E0B4"                       # supplier colour kept as evidence, not as status
    suit = next(i for i in sheet.items if i.article == "TEST-SUIT-02")
    assert [(u["customer"], u["size"]) for u in suit.units] == [("AA", "40"), ("ALINA", "38")]
    scarf = next(i for i in sheet.items if i.article == "TEST-SCF-05")
    assert scarf.confident and [u["customer"] for u in scarf.units] == ["IDS", "AA"]
    shirt = next(i for i in sheet.items if i.article == "TEST-SHT-06")
    assert not shirt.confident and "размер" in shirt.issue
    assert analysis.payments and analysis.payments[0]["amount"] == "1000.00"
    assert result["review"] == 1
    review = db.scalar(select(AllocationReview))
    assert review.row == shirt.row and review.status == "Open"
    unallocated = db.scalars(select(OrderItem).where(OrderItem.customer_id.is_(None))).all()
    assert len(unallocated) == 1 and unallocated[0].quantity == 2       # never silently assigned
    assert Path(factory_import.ORIGINALS_DIR).joinpath(f"{analysis.sha256}.xlsx").read_bytes() == \
        (tmp_path / "Dior cruise SS27 fixture.xlsx").read_bytes()
    stored = [i for i in db.scalars(select(OrderItem)) if i.image_path]
    assert len(stored) >= 6 and (tmp_path / stored[0].image_path).exists()
    again = factory_import.commit_import(db, analysis, Path(stored[0].image_path), {"season": "SS27"})
    assert again["duplicate"] is True


def test_d_one_product_two_customers_two_sizes(db, tmp_path):
    imported(db, tmp_path)
    suit = db.scalars(select(OrderItem).where(OrderItem.article == "TEST-SUIT-02")).all()
    assert sorted((i.customer.code, i.size, i.quantity) for i in suit) == [("AA", "40", 1), ("ALINA", "38", 1)]
    assert len({i.allocation_group_id for i in suit}) == 1                 # same product, two allocations


def test_settlement_from_imported_order_and_rule_changes(db, tmp_path):
    imported(db, tmp_path)
    season = db.scalar(select(Season))
    msg = la.commit_actions(db, [{"action": "set_customer_markup", "params": {"customer": "IDS", "markup": "35"}},
                                 {"action": "set_customer_markup", "params": {"customer": "AA", "markup": "30%"}},
                                 {"action": "set_customer_markup", "params": {"customer": "SS", "markup": 0.35}}], "test", "rates")
    assert "IDS" in msg[0]
    ledger = season_ledger(db, season)
    s = ledger.settlement
    # Base 8100 incl. the unresolved 600 shirt, counted as the owner's until it is resolved.
    assert s.factory.base == D("8100.00") and s.factory.total == D("10530.00") and s.factory_deposit_required == D("5265.00")
    ids_base = D(2000) + D(250)          # coat + one scarf
    aa_base = D(1500) + D(500) + D(250)
    assert s.customers["IDS"].total == marked_up(ids_base, D("0.35"))
    assert s.customers["AA"].total == marked_up(aa_base, D("0.30"))
    assert s.customers["SS"].total == marked_up(D(1000), D("0.35"))
    assert s.owner_net_deposit_planned == s.factory_deposit_required - s.external_deposits_expected
    assert any("без клиента" in w for w in ledger.warnings)
    # Changing AA leaves the factory side untouched.
    la.commit_actions(db, [{"action": "set_customer_markup", "params": {"customer": "AA", "markup": "25"}}], "test", "AA")
    s2 = season_ledger(db, season).settlement
    assert s2.factory.total == s.factory.total and s2.customers["IDS"].total == s.customers["IDS"].total
    assert s2.customers["AA"].total == marked_up(aa_base, D("0.25"))
    assert db.scalars(select(SettlementSnapshot)).all()
    assert db.scalar(select(AuditLog).where(AuditLog.entity_type == "CustomerPricingRule")) is not None


def test_manual_price_survives_rule_changes(db, tmp_path):
    imported(db, tmp_path)
    coat = db.scalar(select(OrderItem).where(OrderItem.article == "TEST-COAT-01"))
    la.commit_actions(db, [{"action": "set_manual_price", "params": {"item_id": coat.id, "unit_price": "1800"}}], "t", "price")
    la.commit_actions(db, [{"action": "set_customer_markup", "params": {"customer": "IDS", "markup": "40"}}], "t", "rate")
    db.refresh(coat)
    assert coat.selling_unit_price == D("1800.00") and coat.markup_source == "manual"
    ids = season_ledger(db, db.scalar(select(Season))).settlement.customers["IDS"]
    assert ids.total == D("1800.00") + marked_up(D(250), D("0.40"))


def test_resolve_ambiguous_allocation_then_partial_ready_and_cancel(db, tmp_path):
    imported(db, tmp_path)
    shirt = db.scalar(select(OrderItem).where(OrderItem.article == "TEST-SHT-06"))
    la.commit_actions(db, [{"action": "resolve_allocation", "params": {"item_id": shirt.id, "assignments": [
        {"customer": "IDS", "size": "40", "quantity": 1}, {"customer": "Alina", "size": "38", "quantity": 1}]}}], "t", "resolve")
    rows = db.scalars(select(OrderItem).where(OrderItem.article == "TEST-SHT-06")).all()
    assert sorted((r.customer.code, r.size) for r in rows) == [("ALINA", "38"), ("IDS", "40")]
    assert db.scalar(select(AllocationReview)).status == "Resolved"
    season = db.scalar(select(Season))
    before = season_ledger(db, season)
    assert before.ready_factory_balance == D("0.00") and not any("без клиента" in w for w in before.warnings)
    coat = db.scalar(select(OrderItem).where(OrderItem.article == "TEST-COAT-01"))
    la.commit_actions(db, [{"action": "set_status", "params": {"item_ids": [coat.id], "status": "Ready"}}], "t", "ready")
    ready = season_ledger(db, season)
    # Coat base 2000 → factory value 2600, half already in the deposit → 1300 now due for it.
    assert ready.ready_factory_balance == D("1300.00")
    assert ready.settlement.factory_deposit_required == before.settlement.factory_deposit_required
    dress = db.scalar(select(OrderItem).where(OrderItem.article == "TEST-DRS-04"))
    with pytest.raises(la.ActionError):
        la.commit_actions(db, [{"action": "set_status", "params": {"item_ids": [dress.id], "status": "Cancelled"}}], "t", "x")
    la.commit_actions(db, [{"action": "set_status", "params": {"item_ids": [dress.id], "status": "Cancelled",
                                                                 "reason": "клиент отказался"}}], "t", "cancel")
    after = season_ledger(db, season).settlement
    assert after.factory.base == before.settlement.factory.base - D(1000)
    assert "SS" not in after.customers
    assert after.customers["IDS"].total == before.settlement.customers["IDS"].total     # unrelated customer untouched


def test_g_payments_in_database_are_counted_once(db, tmp_path):
    imported(db, tmp_path)
    season = db.scalar(select(Season))
    la.commit_actions(db, [
        {"action": "record_payment", "params": {"payer": "AA", "amount": "5000", "to": "factory", "reference": "TX-1"}},
        {"action": "record_payment", "params": {"payer": "IDS", "amount": "1000", "to": "factory", "confirmed": False}},
        {"action": "record_payment", "params": {"payer": "Alina", "amount": "300", "reference": "OWN-1"}},
    ], "t", "payments")
    s = season_ledger(db, season).settlement
    assert s.confirmed_factory_payments == D("5300.00") and s.owner_paid_to_factory == D("300.00")
    assert s.customers["IDS"].expected == D("1000.00")
    with pytest.raises(la.ActionError, match="уже записана"):
        la.commit_actions(db, [{"action": "record_payment", "params": {"payer": "AA", "amount": "5000", "reference": "TX-1"}}], "t", "dup")
    with pytest.raises(la.ActionError, match="курс"):
        la.commit_actions(db, [{"action": "record_payment", "params": {"payer": "AA", "amount": "100", "currency": "CHF"}}], "t", "fx")
    p = db.scalar(select(Payment).where(Payment.status == "expected"))
    la.commit_actions(db, [{"action": "confirm_payment", "params": {"payment_id": p.id}}], "t", "confirm")
    assert season_ledger(db, season).settlement.confirmed_factory_payments == D("6300.00")
    order = db.scalar(select(Order))
    la.commit_actions(db, [{"action": "apply_credit", "params": {"order_id": order.id, "amount": "100", "source": "переплата SS26"}}], "t", "credit")
    assert season_ledger(db, season).settlement.credits_applied == D("100.00")


def test_preview_changes_nothing(db, tmp_path):
    imported(db, tmp_path)
    preview = la.preview_actions(db, [{"action": "set_customer_markup", "params": {"customer": "IDS", "markup": "35"}}])
    assert preview.ok and any(row["label"] == "Заказ IDS" for row in preview.diff)
    assert db.scalar(select(CustomerPricingRule)) is None
    bad = la.preview_actions(db, [{"action": "set_customer_markup", "params": {"customer": "NOBODY", "markup": "35"}}])
    assert not bad.ok and "не найден" in bad.error


def test_mixed_factory_rate_on_an_order(db, tmp_path):
    imported(db, tmp_path)
    order = db.scalar(select(Order))
    la.commit_actions(db, [{"action": "set_factory_rate", "params": {"order_id": order.id, "portions": [
        {"rate": "25", "share": "50"}, {"rate": "30", "share": "50"}]}}], "t", "mixed")
    s = season_ledger(db, db.scalar(select(Season))).settlement
    assert s.factory.total == D("10327.50")          # 4050 × 1.25 + 4050 × 1.30
    agreement = db.scalar(select(FactoryAgreement))
    assert agreement.confirmed and len(agreement.portions) == 2


# --- F. Ready, shipments, transport -------------------------------------------------

def test_f_transport_and_switzerland_exemption():
    class Rule:
        transport_exempt_country_codes = "CH"
    assert transport_fee(D("1000"), D("0.05")) == D("50.00")
    assert country_is_transport_exempt("CH", Rule()) and not country_is_transport_exempt("IT", Rule())


# --- Large workbook editing: patch the original ----------------------------------------

def test_patch_original_keeps_images_formulas_and_other_sheets(tmp_path):
    from app.xlsx_stream import image_anchors, patch_workbook, sheet_parts
    source = factory_fixture.build(tmp_path / "src.xlsx")
    target = tmp_path / "out.xlsx"
    patch_workbook(source, target, {"Dior cruise": {(5, 9): "AA 40, Алина 38 ✓", (4, 6): 2100, (40, 3): "новая строка"}})
    book = load_workbook(target)
    sheet = book["Dior cruise"]
    assert sheet["I5"].value == "AA 40, Алина 38 ✓" and sheet["F4"].value == 2100 and sheet["C40"].value == "новая строка"
    assert sheet["H4"].value == "=F4*G4" and len(sheet._images) == len(factory_fixture.ROWS)
    assert "B2:H2" in {str(r) for r in sheet.merged_cells.ranges} and sheet["I5"].comment is not None
    assert book["SS27"]["B4"].value == "='Dior cruise'!H11"
    with zipfile.ZipFile(target) as z:
        assert z.testzip() is None and len(image_anchors(z, dict(sheet_parts(z))["Dior cruise"])) == len(factory_fixture.ROWS)
        assert b'fullCalcOnLoad="1"' in z.read("xl/workbook.xml")


# --- D/E. Customer documents --------------------------------------------------------

def test_customer_documents_contain_only_that_customer_and_match(db, tmp_path):
    from app import customer_documents as cd
    imported(db, tmp_path)
    la.commit_actions(db, [{"action": "set_customer_markup", "params": {"customer": "IDS", "markup": "35"}},
                           {"action": "record_payment", "params": {"payer": "IDS", "amount": "500", "to": "factory"}}], "t", "x")
    season, ids = db.scalar(select(Season)), db.scalar(select(Customer).where(Customer.code == "IDS"))
    snapshot = cd.customer_snapshot(db, ids, season)
    rows = snapshot["orders"][0]["rows"]
    assert sorted(r["article"] for r in rows) == ["TEST-COAT-01", "TEST-SCF-05"]
    assert snapshot["orders"][0]["base"] == "2250.00" and snapshot["total"] == "3037.50"   # 2250 × 1.35
    assert snapshot["deposit"] == "1518.75" and snapshot["paid"] == "500.00" and snapshot["deposit_outstanding"] == "1018.75"
    docs = cd.generate(db, ids, season, root=tmp_path)
    xlsx = next(d for d in docs if d.file_format == "xlsx")
    pdf = next(d for d in docs if d.file_format == "pdf")
    book = load_workbook(tmp_path / xlsx.stored_path)
    brand = book["Dior cruise"]
    assert len(brand._images) == 2                                   # photos attached to the IDS rows
    values = [[c.value for c in row] for row in brand.iter_rows()]
    flat = " ".join(str(v) for row in values for v in row if v is not None)
    assert "TEST-COAT-01" in flat and "TEST-SUIT-02" not in flat
    assert 3037.5 in [v for row in values for v in row] and 787.5 in [v for row in values for v in row]
    text = cd._xlsx_text((tmp_path / xlsx.stored_path).read_bytes())
    for leaked in ("AA", "SS", "0.3", "30%", "TEST-SUIT-02", "Alina 38", "AA 40"):
        assert not re.search(rf"(?<!\w){re.escape(leaked)}(?!\w)", text), leaked
    pdf_text = cd._pdf_text((tmp_path / pdf.stored_path).read_bytes())
    assert "TEST-COAT-01" in pdf_text and "3 037.50" in pdf_text and "Предоплата" in pdf_text
    assert "AA" not in pdf_text.split() and "SS" not in pdf_text.split() and "TEST-SUIT-02" not in pdf_text
    assert not cd.is_outdated(db, xlsx)
    coat = db.scalar(select(OrderItem).where(OrderItem.article == "TEST-COAT-01"))
    la.commit_actions(db, [{"action": "update_item", "params": {"item_id": coat.id, "size": "42"}}], "t", "size")
    assert cd.is_outdated(db, xlsx)
    again = cd.generate(db, ids, season, root=tmp_path)
    assert again[0].version == 2 and (tmp_path / xlsx.stored_path).exists()     # v1 kept untouched


def test_confidentiality_check_blocks_leaks(db, tmp_path):
    from app import customer_documents as cd
    imported(db, tmp_path)
    ids = db.scalar(select(Customer).where(Customer.code == "IDS"))
    terms = cd.forbidden_terms(db, ids)
    snapshot = {"orders": [{"rows": [{"name": "Coat", "article": "X", "color": ""}]}]}
    assert cd.check_confidential("Coat for AA 40", terms, snapshot)
    assert cd.check_confidential("цена фабрики", terms, snapshot)
    assert not cd.check_confidential("Coat X 2250", terms, snapshot)
    owner = db.scalar(select(Customer).where(Customer.is_owner.is_(True)))
    with pytest.raises(ValueError):
        cd.customer_snapshot(db, owner, db.scalar(select(Season)))
