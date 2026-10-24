"""Upload check: everything that needs attention in one imported file, plus cost corrections.

Problems come from the import notes, allocation notes and protected changes recorded for the
file, and from the same product checks System Check runs. Colours follow the worksheet:
red = not done / open, yellow = in process / warning, green = done / resolved.
"""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import quote

from fastapi import Depends, Form, HTTPException, Request
from fastapi.responses import RedirectResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.audit import record_change
from app.calculations import item_amounts, money
from app.database import get_db
from app.models import (AllocationReview, Expense, ImportBatch, ImportChange, ImportReview, Order, OrderItem,
                        Season, Shipment)
from app.navigation import safe_next
from app.validation import item_issues
from app.worksheet import STATUS_LABELS, status_of

TONE_LABELS = {'error': 'Needs fixing', 'warning': 'Check', 'resolved': 'Resolved'}
# Excel fills match the worksheet export: red / yellow / green.
STATE_FILLS = {'pending': 'F8D7D3', 'process': 'FFF2B3', 'done': 'D9EEDB'}
TONE_FILLS = {'error': 'F8D7D3', 'warning': 'FFF2B3', 'resolved': 'D9EEDB'}
SHOW = {'all', 'problems', 'pending', 'process', 'done'}
EXPENSE_TYPES = ('exceptional', 'transport', 'bag', 'suitcase', 'fixed logistics')


def _name(value):
    return Path((value or '').replace('\\', '/')).name.casefold()


def batch_items(db, batch):
    """Products whose recorded source is this file (incremental re-imports keep their first file)."""
    wanted = _name(batch.filename)
    candidates = db.scalars(select(OrderItem).where(OrderItem.source_file.is_not(None)).options(
        selectinload(OrderItem.order).selectinload(Order.brand),
        selectinload(OrderItem.order).selectinload(Order.season).selectinload(Season.rules),
        selectinload(OrderItem.customer), selectinload(OrderItem.shipment_link),
    ).order_by(OrderItem.source_sheet, OrderItem.source_row, OrderItem.id)).all()
    return [item for item in candidates if _name(item.source_file) == wanted]


def collect(db, batch, show='all'):
    here = f'/imports/{batch.id}/check'
    problems, rows = [], []
    for review in db.scalars(select(ImportReview).where(ImportReview.import_batch_id == batch.id).order_by(ImportReview.id)):
        problems.append(SimpleNamespace(
            tone='error' if review.status == 'Open' else 'resolved', kind='Source note',
            where=review.sheet + (f' · row {review.row}' if review.row else ''),
            message=review.reason, item=None, fix_url=f'/source-reviews/import/{review.id}', fix_label='Review note'))
    for review in db.scalars(select(AllocationReview).where(AllocationReview.import_batch_id == batch.id).order_by(AllocationReview.id)):
        problems.append(SimpleNamespace(
            tone='error' if review.status == 'Open' else 'resolved', kind='Customer split',
            where=review.sheet + (f' · row {review.row}' if review.row else ''),
            message=f"{review.reason} (source: {review.raw_description or 'blank'})", item=None,
            fix_url=f'/source-reviews/allocation/{review.id}', fix_label='Split between customers'))
    for change in db.scalars(select(ImportChange).where(ImportChange.import_batch_id == batch.id, ImportChange.status == 'Review').order_by(ImportChange.id)):
        problems.append(SimpleNamespace(
            tone='error', kind='Protected change', where=change.sheet + (f' · row {change.row}' if change.row else ''),
            message=change.reason or f'{change.action} held for review', item=None,
            fix_url='/import#import-changes', fix_label='Compare before / after'))
    for item in batch_items(db, batch):
        found = item_issues(item, item.order)
        url = f'/worksheet/items/{item.id}?next={quote(here, safe="")}'
        where = (item.source_sheet or item.order.brand.name) + (f' · row {item.source_row}' if item.source_row else '')
        for issue in found:
            problems.append(SimpleNamespace(
                tone=issue.severity, kind='Product', where=where, message=issue.message, item=item,
                fix_url=url, fix_label='Open product'))
        rows.append(SimpleNamespace(item=item, state=status_of(item), value=item_amounts(item).final,
                                    issues=found, url=url,
                                    tone='error' if any(i.severity == 'error' for i in found) else 'warning' if found else None))
    order = {'error': 0, 'warning': 1, 'resolved': 2}
    problems.sort(key=lambda p: order[p.tone])
    counts = {tone: sum(p.tone == tone for p in problems) for tone in TONE_LABELS}
    states = {state: sum(r.state == state for r in rows) for state in STATUS_LABELS}
    if show == 'problems':
        rows = [r for r in rows if r.issues]
    elif show in STATUS_LABELS:
        rows = [r for r in rows if r.state == show]
    rows.sort(key=lambda r: (0 if r.tone == 'error' else 1 if r.tone == 'warning' else 2, r.item.source_sheet or '', r.item.source_row or 0))
    return SimpleNamespace(batch=batch, problems=problems, rows=rows, counts=counts, states=states,
                           here=here, show=show, total_value=money(sum((r.value for r in rows), Decimal(0))))


def check_workbook(data):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Problems'
    header = ['Colour', 'Type', 'Where', 'Problem', 'Product', 'Customer']
    sheet.append(header)
    for problem in data.problems:
        item = problem.item
        sheet.append([TONE_LABELS[problem.tone], problem.kind, problem.where, problem.message,
                      (item.name or item.article or f'#{item.id}') if item else '',
                      (item.customer.code if item.customer else 'Not allocated') if item else ''])
        for cell in sheet[sheet.max_row]:
            cell.fill = PatternFill('solid', fgColor=TONE_FILLS[problem.tone])
    products = book.create_sheet('Products')
    products.append(['Status', 'Sheet', 'Row', 'Article', 'Product', 'Customer', 'Brand', 'Season', 'Colour', 'Size',
                     'Qty', 'Wholesale unit', 'Markup %', 'Customer value', 'Problems'])
    for row in data.rows:
        i = row.item
        products.append([STATUS_LABELS[row.state], i.source_sheet, i.source_row, i.article, i.name,
                         i.customer.code if i.customer else 'Not allocated', i.order.brand.name, i.order.season.name,
                         i.color, i.size, i.quantity,
                         float(i.wholesale_unit_price) if i.wholesale_unit_price is not None else None,
                         float(i.markup_percentage * 100) if i.markup_percentage is not None else None,
                         float(row.value), '; '.join(issue.message for issue in row.issues)])
        for cell in products[products.max_row]:
            cell.fill = PatternFill('solid', fgColor=STATE_FILLS[row.state])
    for ws in (sheet, products):
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='243B35')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            width = max(len(str(c.value or '')) for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 9), 60)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def _money(value, label):
    try:
        amount = money(Decimal((value or '0').strip().replace(',', '.')))
    except InvalidOperation:
        raise HTTPException(422, f'Enter {label} as a number.') from None
    if amount < 0:
        raise HTTPException(422, f'{label.capitalize()} cannot be negative.')
    return amount


def _changed(db, entity, record, field, value, label):
    previous = getattr(record, field)
    if previous != value:
        setattr(record, field, value)
        record_change(db, entity_type=entity, entity_id=record.id, field=field, previous_value=previous,
                      new_value=value, source_user='web user', description=f'{label}: {field} {previous} → {value}')
        return 1
    return 0


def install_upload_check(app, templates, context):
    def get_batch(db, batch_id):
        batch = db.get(ImportBatch, batch_id)
        if batch is None:
            raise HTTPException(404, 'Import not found.')
        return batch

    @app.get('/imports/latest/check')
    def latest_check(db=Depends(get_db)):
        batch = db.scalar(select(ImportBatch).order_by(ImportBatch.id.desc()).limit(1))
        return RedirectResponse(f'/imports/{batch.id}/check' if batch else '/import', 303)

    @app.get('/imports/{batch_id}/check')
    def check_page(batch_id: int, request: Request, show: str = 'all', db=Depends(get_db)):
        data = collect(db, get_batch(db, batch_id), show if show in SHOW else 'all')
        batches = db.scalars(select(ImportBatch).order_by(ImportBatch.id.desc()).limit(20)).all()
        return templates.TemplateResponse(request, 'upload_check.html', context(
            request, check=data, batches=batches, status_labels=STATUS_LABELS, tone_labels=TONE_LABELS))

    @app.get('/imports/{batch_id}/check.xlsx')
    def check_export(batch_id: int, show: str = 'all', db=Depends(get_db)):
        data = collect(db, get_batch(db, batch_id), show if show in SHOW else 'all')
        filename = f'upload-check-{batch_id}.xlsx'
        return Response(check_workbook(data), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition': f'attachment; filename="{filename}"'})

    @app.post('/imports/{batch_id}/check/drive')
    def check_to_drive(batch_id: int, show: str = Form('all'), db=Depends(get_db)):
        from app.outbox import queue_named_export
        data = collect(db, get_batch(db, batch_id), show if show in SHOW else 'all')
        try:
            queue_named_export(db, f'upload-check-{Path(data.batch.filename).stem}.xlsx', check_workbook(data))
        except ValueError as exc:
            return RedirectResponse(f'/imports/{batch_id}/check?drive_error={quote(str(exc))}', 303)
        db.commit()
        return RedirectResponse(f'/imports/{batch_id}/check?drive=queued', 303)

    @app.post('/shipments/{shipment_id}/costs')
    def edit_shipment_costs(shipment_id: int, fixed_cost: str = Form('0'), suitcase_cost: str = Form('0'),
                            bag_cost: str = Form('0'), use_rule_transport: str = Form('on'),
                            transport_percentage: str = Form('0'), shipment_date: str = Form(''),
                            status: str = Form(''), notes: str = Form(''), next: str = Form(''), db=Depends(get_db)):
        from datetime import date
        shipment = db.get(Shipment, shipment_id)
        if shipment is None:
            raise HTTPException(404, 'Shipment not found.')
        try:
            when = date.fromisoformat(shipment_date) if shipment_date else None
        except ValueError:
            raise HTTPException(422, 'Enter a valid shipment date.') from None
        percent = _money(transport_percentage, 'transport percentage') / Decimal(100)
        if percent > 1:
            raise HTTPException(422, 'Transport percentage must be between 0% and 100%.')
        label = f'Shipment {shipment.name}'
        changes = sum([
            _changed(db, 'Shipment', shipment, 'fixed_cost', _money(fixed_cost, 'fixed cost'), label),
            _changed(db, 'Shipment', shipment, 'suitcase_cost', _money(suitcase_cost, 'suitcase cost'), label),
            _changed(db, 'Shipment', shipment, 'bag_cost', _money(bag_cost, 'bag cost'), label),
            _changed(db, 'Shipment', shipment, 'use_rule_transport', use_rule_transport == 'on', label),
            _changed(db, 'Shipment', shipment, 'transport_percentage', percent, label),
            _changed(db, 'Shipment', shipment, 'date', when, label),
            _changed(db, 'Shipment', shipment, 'status', status.strip() or shipment.status, label),
            _changed(db, 'Shipment', shipment, 'notes', notes.strip() or None, label),
        ])
        db.commit()
        target = safe_next(next, f"/logistics?season={quote(shipment.season.name) if shipment.season else ''}&focus=shipments")
        return RedirectResponse(target + ('&' if '?' in target else '?') + f'edited={changes}', 303)

    @app.post('/expenses/{expense_id}')
    def edit_expense(expense_id: int, amount: str = Form(...), expense_type: str = Form(...), note: str = Form(''),
                     next: str = Form(''), db=Depends(get_db)):
        expense = db.get(Expense, expense_id)
        if expense is None:
            raise HTTPException(404, 'Expense not found.')
        clean_type = expense_type.strip().lower()
        if clean_type not in EXPENSE_TYPES:
            raise HTTPException(422, 'Unsupported expense type.')
        label = f'Expense {expense.id}'
        changes = sum([
            _changed(db, 'Expense', expense, 'amount', _money(amount, 'expense amount'), label),
            _changed(db, 'Expense', expense, 'type', clean_type, label),
            _changed(db, 'Expense', expense, 'note', note.strip() or None, label),
        ])
        db.commit()
        target = safe_next(next, f"/logistics?season={quote(expense.season.name) if expense.season else ''}&focus=shipments")
        return RedirectResponse(target + ('&' if '?' in target else '?') + f'edited={changes}', 303)
