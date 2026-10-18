"""A deliberately small, reversible Excel editing contract, separate from audit exports."""
from __future__ import annotations

import hashlib
import io
import json
import re
import secrets
import zipfile
from decimal import Decimal, InvalidOperation

from fastapi import Depends, File, Form, Request, UploadFile
from fastapi.responses import RedirectResponse, Response
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select, text

from app.config import ROOT_DIR
from app.database import get_db
from app.models import Brand, BrandSeasonRule, Customer, ImportBatch, Order, OrderItem, PurchaseGroup, SellerCreditAccount
from app.models import Season
from app.worksheet import STATUS_LABELS, audit, status_of

EXCHANGE_DIR = ROOT_DIR / 'data' / 'excel-exchange'
HEADERS = ['ID', 'Season', 'Brand', 'Customer code', 'Customer name', 'Article / SKU', 'Product',
           'Color', 'Size', 'Quantity', 'Wholesale unit EUR', 'Markup %', 'Selling unit override EUR',
           'Worksheet status', 'Group ID', 'Note', 'Distributor (reference)', 'Seller (reference)',
           'Shipment (reference)', 'Source (reference)', 'Final unit EUR (calculated)', 'Product total EUR (calculated)']
EDITABLE = HEADERS[:16]
COLORS = {'In process': 'FFF2B3', 'Done': 'D9EEDB', 'Not completed': 'FFFFFF'}


def clean(value):
    return str(value).strip() if value is not None else ''


def number(value, label, *, integer=False, optional=False, maximum=Decimal('999999999.99')):
    if optional and not clean(value):
        return ''
    try:
        n = Decimal(clean(value).replace(' ', '').replace(',', '.'))
        precision = Decimal('1') if integer else Decimal('0.01')
        if not n.is_finite() or n < 0 or n > maximum or n != n.quantize(precision):
            raise ValueError()
        return str(int(n)) if integer else format(n.quantize(precision), 'f')
    except (ValueError, InvalidOperation):
        raise ValueError(f'{label}: enter a non-negative {"whole number" if integer else "number with up to 2 decimals"}.') from None


def normalize(values):
    values = dict(zip(EDITABLE, [clean(v) for v in values]))
    for key in ['ID', 'Group ID', 'Quantity']:
        values[key] = number(values[key], key, integer=True, optional=key != 'Quantity')
    for key in ['Wholesale unit EUR', 'Selling unit override EUR', 'Markup %']:
        values[key] = number(values[key], key, optional=key != 'Markup %',
                             maximum=Decimal('100') if key == 'Markup %' else Decimal('999999999.99'))
    return values


def item_values(item):
    group = item.purchase_group
    return [item.id, item.order.season.name, item.order.brand.name,
            item.customer.code if item.customer else '', item.customer.name if item.customer else '',
            item.article, item.name, item.color, item.size, item.quantity, item.wholesale_unit_price,
            item.markup_percentage * 100, item.selling_unit_price, STATUS_LABELS[status_of(item)],
            item.purchase_group_id, item.worksheet_note,
            group.distributor.name if group and group.distributor else '', group.seller if group else '',
            item.shipment.name if item.shipment else '', f'{item.source_sheet or "Manual"} / {item.source_row or ""}']


def token_path(token, suffix):
    if not re.fullmatch(r'[a-f0-9]{32}', token):
        raise ValueError('This preview is not valid. Upload your workbook again.')
    return EXCHANGE_DIR / f'{token}.{suffix}'


def export_editable(db, blank=False):
    EXCHANGE_DIR.mkdir(parents=True, exist_ok=True)
    token = secrets.token_hex(16)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Orders'
    ws.append(HEADERS)
    items = [] if blank else db.scalars(select(OrderItem).order_by(OrderItem.id)).all()
    baseline = {}
    for item in items:
        values = item_values(item)
        ws.append(values + [None, None])
        baseline[str(item.id)] = normalize(values[:16])
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill('solid', fgColor=COLORS[values[13]])
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Source text is never executable Excel input.
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='243B35')
        cell.font = Font(color='FFFFFF', bold=True)
    for index in range(2, ws.max_row + 1):
        ws[f'U{index}'] = f'=IF(M{index}<>"",M{index},IF(K{index}="","",ROUND(K{index}*(1+L{index}/100),2)))'
        ws[f'V{index}'] = f'=IF(U{index}="","",ROUND(U{index}*J{index},2))'
    widths = [9,12,18,17,25,23,32,18,12,12,20,12,24,20,12,38,23,23,23,28,24,24]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1,idx).column_letter].width = width
    ws.freeze_panes = 'G2'
    ws.auto_filter.ref = f'A1:V{max(2, ws.max_row)}'
    for col in ['K','L','M','U','V']:
        for row in range(2,ws.max_row+1):
            ws[f'{col}{row}'].number_format = '#,##0.00'
    validation = DataValidation(type='list', formula1='"In process,Done,Not completed"')
    validation.errorTitle = 'Choose a worksheet status'
    validation.error = 'Use the dropdown; color follows the status.'
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add('N2:N10001')
    for label,color in COLORS.items():
        ws.conditional_formatting.add('A2:V10001', FormulaRule(formula=[f'$N2="{label}"'], fill=PatternFill('solid',fgColor=color)))
    help_ws = wb.create_sheet('Start here')
    notes = [
        'ALINA — download, edit, upload',
        '1. Edit the Orders sheet. Keep its header row, ID column, and the hidden _Exchange sheet.',
        '2. Existing rows: change customer code, product details, prices, status, group ID or note.',
        '3. New rows: leave ID blank. Supply season, brand, product, quantity and markup %. Use an existing customer code or a new unique code + name.',
        '4. Upload at Excel files. Review every proposed change before Apply. Nothing is deleted when you remove a row.',
        'Yellow = In process; green = Done; white = Not completed. Use the status dropdown, not paint alone.',
        'Group ID must refer to an existing group in the same season. Use Groups to set seller, distributor and group costs.',
        'Existing Season and Brand are references. To change them, use the app. Customer names are reference-only for existing customer codes.',
        'Blank wholesale = unknown cost. Blank selling override = automatic markup. Markup is a percentage such as 35, not 0.35.',
        'Columns Q–T are references. Payments, seller credit, shipment history and group costs are never overwritten by this file.',
        'Columns U–V calculate the product selling unit and total for existing rows. Copy these formulas down for new rows; the app always recalculates totals itself. Group costs are separate.',
        'Colors do not record payments or change logistics status. Cancelled source rows remain excluded from accounting until reviewed in the app.',
        'Formulas in editable cells are not imported. Paste values. Download a fresh workbook after applying changes.',
        'Missing customer, size or price will be flagged with a link to resolve it. Unknown facts are never guessed.',
    ]
    for note in notes:
        help_ws.append([note])
        help_ws.cell(help_ws.max_row,1).alignment = Alignment(wrap_text=True,vertical='top')
        help_ws.row_dimensions[help_ws.max_row].height = 34
    help_ws.column_dimensions['A'].width = 115
    refs = wb.create_sheet('Groups reference')
    refs.append(['Group ID','Group','Season','Distributor','Seller'])
    for g in db.scalars(select(PurchaseGroup).order_by(PurchaseGroup.id)):
        refs.append([g.id,g.name,g.season.name,g.distributor.name if g.distributor else '',g.seller or ''])
    for column in ['A','B','C','D','E']:
        refs.column_dimensions[column].width = 28
    customers = wb.create_sheet('Customers reference')
    customers.append(['Customer code','Customer name'])
    for customer in db.scalars(select(Customer).order_by(Customer.code)):
        customers.append([customer.code, customer.name or ''])
    customers.column_dimensions['A'].width = 22
    customers.column_dimensions['B'].width = 38
    for sheet in [help_ws, refs, customers]:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type = 's'
    meta = wb.create_sheet('_Exchange')
    meta.append(['ALINA_EDITABLE_V1',token])
    meta.sheet_state = 'veryHidden'
    token_path(token,'json').write_text(json.dumps(baseline, ensure_ascii=False), encoding='utf-8')
    from app.storage import mark_dirty
    mark_dirty()  # The comparison baseline must be backed up even before an upload.
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def inspect_upload(db, content):
    digest = hashlib.sha256(content).hexdigest()
    if db.scalar(select(ImportBatch.id).where(ImportBatch.sha256 == digest)):
        raise ValueError('This exact file has already been applied. Download a fresh workbook before editing again.')
    if len(content) > 8 * 1024 * 1024:
        raise ValueError('Use an .xlsx file smaller than 8 MB. For source documents or images, use AI to table.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 40 * 1024 * 1024 or len(archive.infolist()) > 1500:
                raise ValueError('Workbook is too large. Use a smaller file.')
        wb = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
    except (zipfile.BadZipFile, KeyError, OSError):
        raise ValueError('Could not read this Excel file. Save as .xlsx, or use AI to table for another layout.') from None
    try:
        if '_Exchange' not in wb or 'Orders' not in wb or wb['_Exchange']['A1'].value != 'ALINA_EDITABLE_V1':
            raise ValueError('This is a different workbook layout. Use AI to table, or copy the rows into our blank template. Original SS26 files have a separate source importer.')
        try:
            exchange_token = clean(wb['_Exchange']['B1'].value)
            baseline = json.loads(token_path(exchange_token,'json').read_text(encoding='utf-8'))
        except FileNotFoundError:
            raise ValueError('This workbook belongs to another app copy or its baseline is missing. Download a fresh workbook here.') from None
        if db.scalar(select(ImportBatch.id).where(ImportBatch.summary_json.contains(f'"exchange_token": "{exchange_token}"'))):
            raise ValueError('This download has already been applied, even if it was renamed or saved again in Excel. Download a fresh workbook to continue without duplicate new rows.')
        ws = wb['Orders']
        if ws.max_row > 10001 or ws.max_column > 60:
            raise ValueError('Use at most 10,000 product rows and keep the original columns.')
        if [cell.value for cell in next(ws.iter_rows(max_row=1,max_col=len(HEADERS)))] != HEADERS:
            raise ValueError('The Orders headers changed. Download the template and keep the header row unchanged.')
        rows, errors, seen, unchanged = [], [], set(), 0
        new_customers = {}
        for row_number,cells in enumerate(ws.iter_rows(min_row=2,max_col=len(HEADERS)),2):
            raw = [c.value for c in cells[:16]]
            if not any(v is not None and str(v).strip() for v in raw):
                continue
            identity = clean(raw[0])
            try:
                if any(c.data_type == 'f' for c in cells[:16]):
                    raise ValueError('Paste values instead of formulas in editable columns.')
                data = normalize(raw)
                identity = data['ID']
                item = db.get(OrderItem,int(identity)) if identity else None
                original = baseline.get(identity)
                if identity:
                    if identity in seen:
                        raise ValueError('Duplicate ID. Keep one row per existing ID; clear ID only for a genuinely new product.')
                    seen.add(identity)
                    if not item or original is None:
                        raise ValueError('ID is not in this download. Download the latest workbook, or leave ID blank for a new product.')
                    if data == original:
                        unchanged += 1
                        continue
                    if normalize(item_values(item)[:16]) != original:
                        raise ValueError('This product changed in the app after download. Download a fresh file and copy your edits into it.')
                    if any(data[k] != original[k] for k in ['Season','Brand']):
                        raise ValueError('Existing season and brand cannot be moved through Excel. Edit the product in the app.')
                for key,maxlength in [('Season',30),('Brand',200),('Customer code',20),('Customer name',200),('Article / SKU',200),('Product',300),('Color',200),('Size',100),('Note',5000)]:
                    if len(data[key]) > maxlength:
                        raise ValueError(f'{key} is too long (maximum {maxlength} characters).')
                if not all(data[k] for k in ['Season','Brand','Product']):
                    raise ValueError('Fill Season, Brand and Product. Unknown size/customer/price may stay blank for later review.')
                if not identity and int(data['Quantity']) < 1:
                    raise ValueError('A new product needs quantity of at least 1.')
                if data['Worksheet status'] not in COLORS:
                    raise ValueError('Choose In process, Done, or Not completed in Worksheet status.')
                code, name = data['Customer code'], data['Customer name']
                customer = db.scalar(select(Customer).where(Customer.code == code)) if code else None
                if name and not code:
                    raise ValueError('Supply a customer code for this name, or clear both to leave allocation unresolved.')
                if customer and name != (customer.name or ''):
                    raise ValueError('Customer name does not match this code. Copy both from Customers reference; rename customers in the app.')
                if code and not customer:
                    if not name:
                        raise ValueError('A new customer code also needs a customer name.')
                    if code in new_customers and new_customers[code] != name:
                        raise ValueError('This new customer code has different names in the same file.')
                    new_customers[code] = name
                group = db.get(PurchaseGroup,int(data['Group ID'])) if data['Group ID'] else None
                if data['Group ID'] and (not group or group.season.name != data['Season']):
                    raise ValueError('Group ID must exist and belong to this season. Create or check it in Groups.')
                changing_owner = item and (item.customer_id != (customer.id if customer else None) or str(item.purchase_group_id or '') != data['Group ID'])
                if changing_owner and db.scalar(select(SellerCreditAccount.id).where(
                    (SellerCreditAccount.item_id == item.id) | (SellerCreditAccount.group_id == item.purchase_group_id))):
                    raise ValueError('Customer/group is locked by seller-credit history. Review credit in the app before moving funds or ownership.')
                changing_group = not item or str(item.purchase_group_id or '') != data['Group ID']
                if changing_group and ((group and group.shipment_id) or (item and (item.shipment_link or (item.purchase_group and item.purchase_group.shipment_id)))):
                    raise ValueError('Dispatched products/groups cannot be regrouped in Excel. Use Groups and Shipments to review history.')
                if changing_group and group and db.scalar(select(SellerCreditAccount.id).where(SellerCreditAccount.group_id == group.id)):
                    raise ValueError('The destination group has seller-credit history. Create a new group for these products.')
                changed = [dict(field=k,before=original[k],after=data[k]) for k in EDITABLE if original and data[k] != original[k]]
                warnings = []
                for key,label in [('Customer code','customer allocation'),('Size','size'),('Wholesale unit EUR','cost price'),('Group ID','shipment group')]:
                    if not data[key]: warnings.append(f'Confirm {label} after import')
                rows.append(dict(row=row_number,kind='Update' if item else 'New',data=data,changes=changed,warnings=warnings,
                                 url=f'/worksheet/items/{item.id}' if item else '/manual-entry'))
            except ValueError as exc:
                errors.append(dict(row=row_number,message=str(exc),url=f'/worksheet/items/{identity}' if identity.isdigit() and identity in baseline else '/excel'))
        return dict(rows=rows,errors=errors,unchanged=unchanged,missing=len(set(baseline)-seen),digest=digest,exchange_token=exchange_token)
    finally:
        wb.close()


def apply_upload(db, content, source_path):
    # Revalidate under the write lock: a preview never grants permission to overwrite later changes.
    db.execute(text('BEGIN IMMEDIATE'))
    result = inspect_upload(db,content)
    if result['errors']:
        raise ValueError('The file now has conflicts. Preview it again; no changes were saved.')
    for row in result['rows']:
        data = row['data']
        item = db.get(OrderItem,int(data['ID'])) if data['ID'] else None
        if not item:
            season = db.scalar(select(Season).where(Season.name == data['Season']))
            if not season:
                season = Season(name=data['Season']); db.add(season); db.flush()
            brand = db.scalar(select(Brand).where(Brand.name == data['Brand']))
            if not brand:
                brand = Brand(name=data['Brand'],default_markup_percentage=Decimal(data['Markup %'])/100); db.add(brand); db.flush()
            order = db.scalar(select(Order).where(Order.season_id == season.id,Order.brand_id == brand.id))
            if not db.scalar(select(BrandSeasonRule.id).where(BrandSeasonRule.season_id == season.id,BrandSeasonRule.brand_id == brand.id)):
                db.add(BrandSeasonRule(season_id=season.id,brand_id=brand.id,markup_percentage=Decimal(data['Markup %'])/100))
            if not order:
                order = Order(season_id=season.id,brand_id=brand.id); db.add(order); db.flush()
            item = OrderItem(order_id=order.id,source_file=str(source_path),source_sheet='Orders',source_row=row['row'],status='Ordered')
            db.add(item)
        customer = db.scalar(select(Customer).where(Customer.code == data['Customer code'])) if data['Customer code'] else None
        if data['Customer code'] and not customer:
            customer = Customer(code=data['Customer code'],name=data['Customer name']); db.add(customer); db.flush()
        item.customer_id = customer.id if customer else None
        item.allocation_status = 'Allocated' if customer else 'Review'
        for field,key in [('article','Article / SKU'),('name','Product'),('color','Color'),('size','Size'),('worksheet_note','Note')]:
            setattr(item,field,data[key] or None)
        item.quantity = int(data['Quantity'])
        item.wholesale_unit_price = Decimal(data['Wholesale unit EUR']) if data['Wholesale unit EUR'] else None
        item.selling_unit_price = Decimal(data['Selling unit override EUR']) if data['Selling unit override EUR'] else None
        item.markup_percentage = Decimal(data['Markup %'])/100
        item.worksheet_status = next(k for k,v in STATUS_LABELS.items() if v == data['Worksheet status'])
        item.purchase_group_id = int(data['Group ID']) if data['Group ID'] else None
        db.flush()
        audit(db,'OrderItem',item.id,'Editable Excel: '+json.dumps(row,ensure_ascii=False))
    db.add(ImportBatch(filename=source_path.name,sha256=result['digest'],stored_path=str(source_path),
                       summary_json=json.dumps(dict(editable_excel=True,exchange_token=result['exchange_token'],changed=len(result['rows']),unchanged=result['unchanged']))))
    db.commit()
    return result


def install_excel_exchange(app,templates,context):
    def page(request,db,**kwargs):
        return templates.TemplateResponse(request,'excel_exchange.html',context(request,**kwargs))

    @app.get('/excel')
    def hub(request: Request,db=Depends(get_db)):
        return page(request,db)

    @app.get('/excel/download')
    def download(blank: bool=False,db=Depends(get_db)):
        return Response(export_editable(db,blank),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition':f'attachment; filename="ALINA-{ "new-products" if blank else "editable-orders"}.xlsx"','Cache-Control':'no-store'})

    @app.post('/excel/preview')
    def preview(request: Request,file: UploadFile=File(...),db=Depends(get_db)):
        try:
            content = file.file.read(8*1024*1024+1)
            result = inspect_upload(db,content)
            token = secrets.token_hex(16)
            EXCHANGE_DIR.mkdir(parents=True,exist_ok=True)
            token_path(token,'xlsx').write_bytes(content)
            return page(request,db,result=result,token=token,filename=file.filename)
        except (ValueError,InvalidOperation) as exc:
            return page(request,db,error=str(exc))

    @app.post('/excel/apply')
    def apply(request: Request,token: str=Form(...),db=Depends(get_db)):
        try:
            path = token_path(token,'xlsx')
            result = apply_upload(db,path.read_bytes(),path)
            return RedirectResponse(f'/excel?saved={len(result["rows"])}',303)
        except (ValueError,FileNotFoundError) as exc:
            db.rollback()
            return page(request,db,error=str(exc) if isinstance(exc,ValueError) else 'Preview expired. Upload your file again.')
