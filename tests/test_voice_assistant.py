"""Voice assistant: transcript normalisation, speech text, exact selection, conversation flow, selective exports,
spoken confirmation rules and the ElevenLabs endpoints (mocked; nothing here calls ElevenLabs)."""
from __future__ import annotations

import io
import json
import zipfile
from decimal import Decimal as D

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from test_factory_settlement import db, imported  # noqa: F401
from app import agent, selection as sel, voice
from app.customer_documents import _pdf_text, _xlsx_text
from app.models import AgentProposal, CustomerPricingRule, GeneratedDocument, Order, OrderItem

CODES, BRANDS = ["AA", "IDS", "SS", "ALINA"], ["Dior cruise"]


def files(db, tmp_path, ids):
    return [(d, (tmp_path / d.stored_path).read_bytes()) for d in (db.get(GeneratedDocument, i) for i in ids)]


def doc_text(doc, data):
    return _xlsx_text(data) if doc.file_format == "xlsx" else _pdf_text(data)


# --- Speech recognition normalisation ------------------------------------------------------

@pytest.mark.parametrize("heard,expected", [
    ("Поставь ай ди эс тридцать пять процентов, а эй эй тридцать процентов", "Поставь IDS 35%, а AA 30%"),
    ("эс эс тридцать процентов", "SS 30%"),
    ("покажи диор круиз для айдиэс", "покажи Dior cruise для IDS"),
    ("вещи тридцать восьмого размера", "вещи 38 размера"),
    ("вещи сорокового размера", "вещи 40 размера"),
    ("IDS 35 процентов", "IDS 35%"),
    ("две тысячи пятьсот евро", "2500 евро"),
])
def test_spoken_forms_become_exact_codes_numbers_and_brands(heard, expected):
    assert voice.normalize_transcript(heard, CODES, BRANDS) == expected


def test_speech_text_is_short_plain_and_keeps_verified_amounts():
    reply = ("По заказу Dior cruise IDS должна внести предоплату €1 518.75.\n```json {\"x\": 1}```\n"
             "• Формула =SUM(H4:H6) см /board?season=SS27\n" + "Строка. " * 80)
    spoken = voice.to_speech(reply, CODES)
    assert "1 518 евро 75 центов" in spoken and "ай-ди-эс" in spoken
    assert not any(bad in spoken for bad in ("{", "json", "=SUM", "/board", "```"))
    assert len(spoken) < 360 and spoken.endswith("Подробности на экране.")
    assert voice.to_speech("Наценка 35%, у AA 30%", CODES) == "Наценка 35 процентов, у эй-эй 30 процентов."


def test_key_checks_explain_the_key_id_mistake():
    assert "ID ключа" in voice.key_problem("0123456789abcdef" * 4)
    assert voice.key_problem("abc") and voice.key_problem(voice.clean_key("-sk_test123")) is None


# --- Exact selection -------------------------------------------------------------------------

def test_selection_by_brand_customer_size_status_and_rows(db, tmp_path):
    imported(db, tmp_path)
    f, q, _ = sel.parse(db, "Покажи все вещи Dior cruise для IDS")
    s = sel.resolve(db, f)
    assert q is None and {i.customer.code for i in s.items} == {"IDS"} and len(s.items) == 2
    assert s.summary == "Выбрано: SS27 → Dior cruise → IDS → 2 позиции"
    f, _, _ = sel.parse(db, "Выдели все вещи 40 размера у AA")
    assert {(i.customer.code, i.size) for i in sel.resolve(db, f).items} == {("AA", "40")}
    f, _, _ = sel.parse(db, "Выдели все вещи 38 размера у Алины")
    assert {(i.customer.code, i.size) for i in sel.resolve(db, f).items} == {("ALINA", "38")}
    f, _, _ = sel.parse(db, "покажи платья SS")
    assert [i.name for i in sel.resolve(db, f).items] == ["Dress test"]
    assert sel.resolve(db, sel.Filters(readiness="ready")).items == []
    rows = sel.resolve(db, sel.Filters(sheet="Dior cruise", rows=[4])).items
    assert [i.article for i in rows] == ["TEST-COAT-01"]


def test_dior_versus_dior_cruise_and_unknown_customer(db, tmp_path):
    imported(db, tmp_path)
    season = db.scalars(select(Order)).first().season
    assert sel.match_brands("открой Dior", season, db) == (["Dior cruise"], None)     # the only Dior order
    from app.models import Brand
    db.add(Order(season_id=season.id, brand=Brand(name="Dior")))
    db.add(Order(season_id=season.id, brand=Brand(name="Dior beachwear")))
    db.commit()
    db.refresh(season)
    assert sel.match_brands("Dior cruise для IDS", season, db)[0] == ["Dior cruise"]
    assert sel.match_brands("покажи Dior", season, db)[0] == ["Dior"]
    _, question, _ = sel.parse(db, "Сделай Excel для Кристины")
    assert question and "не найден" in question["question"]


def test_same_product_name_never_pulls_in_other_customers(db, tmp_path):
    imported(db, tmp_path)
    f, _, _ = sel.parse(db, "покажи Scarf у IDS")
    items = sel.resolve(db, f).items
    assert {i.customer.code for i in items} == {"IDS"}          # the AA scarf of the same row stays out


# --- The conversation from the brief, offline -------------------------------------------------

def test_voice_conversation_select_markup_confirm_deposit_documents(db, tmp_path):
    imported(db, tmp_path)
    c = "voice1"
    first = agent.handle(db, c, "Открой диор круиз", channel="voice")
    assert first.panel["selection"]["summary"].startswith("Выбрано: SS27 → Dior cruise")
    ids = agent.handle(db, c, "Покажи ай ди эс", channel="voice")
    assert ids.panel["selection"]["customers"] == ["IDS"] and "Dior cruise → IDS → 2 позиции" in ids.reply
    markup = agent.handle(db, c, "Поставь ей тридцать пять процентов", channel="voice")
    assert markup.proposal and markup.proposal["status"] == "pending"
    assert db.scalar(select(CustomerPricingRule)) is None                          # nothing saved yet
    assert any(r["label"] == "IDS: предоплата" and D(r["after"]) == D("1518.75") for r in markup.proposal["diff"])
    saved = agent.handle(db, c, "Да, сохрани", channel="voice", screen={"proposal": {"id": markup.proposal["id"], "version": 1}})
    assert saved.reply.startswith("Сохранено") and db.scalar(select(CustomerPricingRule)).markup == D("0.35")
    deposit = agent.handle(db, c, "А сколько теперь предоплата?", channel="voice")
    assert "€1 518.75" in deposit.reply and "1 518.75" in deposit.speech
    assert "IDS должна внести предоплату в размере €1 518.75" in deposit.speech
    pdf = agent.handle(db, c, "Сделай PDF", channel="voice")
    assert [d["format"] for d in pdf.panel["documents"]] == ["pdf"] and pdf.panel["documents"][0]["customer"] == "IDS"
    xlsx = agent.handle(db, c, "И Excel тоже", channel="voice")
    assert [d["format"] for d in xlsx.panel["documents"]] == ["xlsx"] and xlsx.panel["documents"][0]["customer"] == "IDS"
    for doc, data in files(db, tmp_path, [d["id"] for d in pdf.panel["documents"] + xlsx.panel["documents"]]):
        text = doc_text(doc, data)
        assert "3 037.50" in text.replace(",", " ") or "3037.5" in text
        assert not any(code in text for code in ("AA", " SS ")) and "фабрик" not in text.casefold()
    aa = agent.handle(db, c, "Теперь сделай это для AA, но с 30%", channel="voice")
    assert aa.proposal and any(r["label"] == "Заказ AA" for r in aa.proposal["diff"])
    assert aa.proposal["documents"] == [{"customer": "AA", "season": "SS27", "brands": ["Dior cruise"]}]   # "это" = the same documents
    assert all(a["params"].get("brand") == "Dior cruise" for a in json.loads(
        db.scalar(select(AgentProposal).where(AgentProposal.public_id == aa.proposal["id"])).actions_json)["actions"])


def test_correction_rejection_and_undo(db, tmp_path):
    imported(db, tmp_path)
    c = "voice2"
    first = agent.handle(db, c, "IDS 35%, AA 30%")
    corrected = agent.handle(db, c, "Нет, у AA оставь 25", channel="voice")
    assert corrected.proposal and corrected.proposal["id"] != first.proposal["id"]
    plan = json.loads(db.scalar(select(AgentProposal).where(AgentProposal.public_id == corrected.proposal["id"])).actions_json)
    assert {a["params"]["customer"]: a["params"]["markup"] for a in plan["actions"]} == {"IDS": "35", "AA": "25"}
    assert db.scalar(select(AgentProposal).where(AgentProposal.public_id == first.proposal["id"])).status == "cancelled"
    agent.handle(db, c, "да", channel="voice")
    before_undo = {i.id: i.markup_percentage for i in db.scalars(select(OrderItem))}
    assert D("0.25") in before_undo.values()
    undo = agent.handle(db, c, "Отмени последнее изменение", channel="voice")
    assert undo.proposal and undo.proposal["status"] == "pending"
    agent.handle(db, c, "Да, сохрани", channel="voice")
    assert all(i.markup_percentage == D("0") for i in db.scalars(select(OrderItem)) if i.customer and not i.customer.is_owner)
    assert db.scalar(select(CustomerPricingRule)) is None


def test_a_yes_never_confirms_another_conversation_or_a_changed_or_risky_proposal(db, tmp_path):
    imported(db, tmp_path)
    pending = agent.handle(db, "conv-a", "IDS 35%")
    other = agent.handle(db, "conv-b", "Да, сохрани", channel="voice")
    assert not other.reply.startswith("Сохранено") and db.scalar(select(CustomerPricingRule)) is None
    stale = agent.handle(db, "conv-a", "да", channel="voice", screen={"proposal": {"id": pending.proposal["id"], "version": 0}})
    assert "обновилось" in stale.reply and db.scalar(select(CustomerPricingRule)) is None
    payment = agent.handle(db, "conv-c", "AA оплатила 500 евро, поступили фабрике")
    assert payment.proposal
    spoken = agent.handle(db, "conv-c", "Да, сохрани", channel="voice")
    assert "кнопк" in spoken.reply
    from app.models import Payment
    assert db.scalar(select(Payment)) is None


def test_cancel_item_by_voice_and_explain_colour(db, tmp_path):
    imported(db, tmp_path)
    c = "voice3"
    cancel = agent.handle(db, c, "Убери это платье у SS и пересчитай предоплату", channel="voice")
    assert cancel.proposal and "Dress test" in cancel.proposal["summary"]
    coat = db.scalar(select(OrderItem).where(OrderItem.article == "TEST-COAT-01"))
    why = agent.handle(db, c, "Почему здесь зелёное?", screen={"selected_item_ids": [coat.id]})
    assert "#C6E0B4" in why.reply and "зелёный" in why.reply and "Заказано" in why.reply


# --- Selective exports -----------------------------------------------------------------------

def test_exports_only_ids_dior_cruise_per_customer_isolation_and_three_rows(db, tmp_path):
    imported(db, tmp_path)
    ids = agent.handle(db, "x1", "Экспортируй только IDS из Dior cruise в PDF")
    docs = files(db, tmp_path, [d["id"] for d in ids.panel["documents"]])
    assert [(d.file_format, d.customer.code) for d, _ in docs] == [("pdf", "IDS")]
    assert "Coat test" in doc_text(*docs[0]) and "Suit test" not in doc_text(*docs[0])
    every = agent.handle(db, "x2", "Сделай отдельные PDF для каждого клиента")
    made = files(db, tmp_path, [d["id"] for d in every.panel["documents"]])
    assert sorted(d.customer.code for d, _ in made) == ["AA", "IDS", "SS"]          # never the owner
    for doc, data in made:
        others = {"AA", "IDS", "SS"} - {doc.customer.code}
        text = doc_text(doc, data)
        assert not any(f" {o} " in f" {text} " for o in others)
    allocated = list(db.scalars(select(OrderItem).where(OrderItem.customer_id.is_not(None)).order_by(OrderItem.id)))
    chosen = [i.id for i in allocated if i.customer.code == "AA"][:2] + [i.id for i in allocated if i.customer.code == "IDS"][:1]
    three = agent.handle(db, "x3", "Экспортируй только выделенные строки в Excel", screen={"selected_item_ids": chosen})
    assert "выделено 3" in three.reply
    rows = []
    for doc, data in files(db, tmp_path, [d["id"] for d in three.panel["documents"]]):
        assert json.loads(doc.snapshot_json)["subset"] is True
        sheet = load_workbook(io.BytesIO(data)).worksheets[1]
        rows += [sheet.cell(r, 3).value for r in range(4, sheet.max_row + 1) if str(sheet.cell(r, 3).value or "").startswith("TEST-")]
    assert len(rows) == 3


def test_ready_only_and_internal_report_and_original_with_prices(db, tmp_path):
    imported(db, tmp_path)
    aa = [i for i in db.scalars(select(OrderItem)) if i.customer and i.customer.code == "AA"]
    aa[0].status = "Ready"
    db.commit()
    ready = agent.handle(db, "x4", "Сделай Excel только по готовым вещам AA")
    doc, data = files(db, tmp_path, [ready.panel["documents"][0]["id"]])[0]
    snap = json.loads(doc.snapshot_json)
    assert snap["item_ids"] == [aa[0].id] and snap["label"] == "Готовые"
    both = agent.handle(db, "x5", "Сделай общий документ для фабрики и отдельные документы для каждого клиента")
    kinds = sorted({d["kind"] for d in both.panel["documents"]})
    assert kinds == ["customer", "internal"] and both.panel["bundle"]
    internal = next(d for d in both.panel["documents"] if d["kind"] == "internal" and d["format"] == "xlsx")
    text = doc_text(*files(db, tmp_path, [internal["id"]])[0])
    assert "ВНУТРЕННИЙ" in text and "IDS" in text and "AA" in text
    original = agent.handle(db, "x6", "Сделай точно такую же таблицу, как оригинал, но с моими ценами")
    doc, data = files(db, tmp_path, [original.panel["documents"][0]["id"]])[0]
    book = load_workbook(io.BytesIO(data))
    sheet = book["Dior cruise"]
    assert sheet["F4"].value == 2000 and sheet["H4"].value == "=F4*G4"          # original cells untouched
    assert "IDS" in str(sheet.cell(4, sheet.max_column - 1).value)
    assert len(sheet._images) == 6                                                # photos kept


# --- Voice HTTP endpoints (ElevenLabs mocked) ------------------------------------------------

@pytest.fixture()
def client(db, tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from app.database import get_db
    from app.main import app
    app.dependency_overrides[get_db] = lambda: (yield db)
    try:
        with TestClient(app) as c:
            yield c
    finally:
        app.dependency_overrides.clear()


def mock_transport(log):
    def handler(request: httpx.Request) -> httpx.Response:
        log.append((request.method, request.url.path, request.headers.get("xi-api-key")))
        if request.headers.get("xi-api-key") != "sk_good":
            return httpx.Response(401, json={"detail": {"status": "invalid_api_key", "message": "Invalid API key"}})
        if request.url.path == "/v1/single-use-token/realtime_scribe":
            return httpx.Response(200, json={"token": "sutkn_test"})
        if request.url.path.startswith("/v1/text-to-speech/"):
            body = json.loads(request.content)
            assert body["model_id"] in voice.TTS_MODELS and "{" not in body["text"]
            return httpx.Response(200, content=b"ID3" + b"\x00" * 2000, headers={"content-type": "audio/mpeg"})
        if request.url.path == "/v1/speech-to-text":
            return httpx.Response(200, json={"text": "покажи ай ди эс тридцать пять процентов"})
        return httpx.Response(403, json={"detail": {"status": "missing_permissions", "message": "missing the permission voices_read"}})
    return httpx.MockTransport(handler)


def test_voice_endpoints_without_key_explain_setup(client):
    assert client.get("/api/voice/config").json()["configured"] is False
    response = client.post("/api/voice/token")
    assert response.status_code == 409 and "/settings/voice" in response.json()["setup"]
    page = client.get("/settings/voice")
    assert page.status_code == 200 and "Не подключено" in page.text


def test_voice_endpoints_with_mocked_elevenlabs(client, monkeypatch):
    log = []
    transport = mock_transport(log)
    real = voice.ElevenLabs.__init__
    monkeypatch.setattr(voice.ElevenLabs, "__init__", lambda self, key, t=None, base=voice.API: real(self, key, transport, base))
    saved = client.post("/settings/voice", data={"action": "test", "api_key": "sk_bad", "voice_id": voice.DEFAULT_VOICE,
                                                 "model": "eleven_v4_turbo"}, follow_redirects=False)
    assert "error" in saved.headers["location"] and voice.current().configured is False
    saved = client.post("/settings/voice", data={"action": "test", "api_key": " -sk_good ", "voice_id": voice.DEFAULT_VOICE,
                                                 "model": "eleven_v4_turbo", "speak_replies": "1"}, follow_redirects=False)
    assert "message" in saved.headers["location"] and voice.current().api_key == "sk_good"
    page = client.get("/settings/voice")
    assert "sk_good" not in page.text and "…good" not in page.text and "Подключено" in page.text
    assert "sk_good" not in client.get("/api/voice/config").text
    token = client.post("/api/voice/token?hands_free=true").json()
    assert token["url"].startswith(voice.REALTIME_URL) and "token=sutkn_test" in token["url"] and "commit_strategy=vad" in token["url"]
    assert "sk_good" not in json.dumps(token)
    speech = client.post("/api/voice/speech", json={"text": "Предоплата €1 518.75 ```{\"a\":1}```"}).json()
    assert speech["text"] == "Предоплата 1 518 евро 75 центов."
    audio = client.get(f"/api/voice/speech/{speech['id']}")
    assert audio.status_code == 200 and audio.headers["content-type"] == "audio/mpeg" and audio.content.startswith(b"ID3")
    assert client.get(f"/api/voice/speech/{speech['id']}").status_code == 404           # single use
    heard = client.post("/api/voice/transcribe", files={"file": ("note.m4a", b"fake-audio", "audio/mp4")}).json()
    assert heard["text"] == "покажи IDS 35%" and heard["raw"].startswith("покажи ай ди эс")


def test_agent_message_api_accepts_voice_channel_and_returns_speech(client, db, tmp_path):
    response = client.post("/api/agent/message", json={"conversation": "a" * 16, "text": "Что изменилось за неделю?",
                                                       "channel": "voice", "raw": "что изменилось за неделю"})
    data = response.json()
    assert response.status_code == 200 and "reply" in data and "speech" in data
