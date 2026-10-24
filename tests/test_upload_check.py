"""Upload check page, colours, bulk marking, Excel export, cost edits and the no-rule System Check fix."""
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from test_workspace_features import workspace_session, workspace_client  # noqa: F401
from app import main
from app.models import (AllocationReview, AuditLog, Brand, Expense, ImportBatch, ImportReview, Order, OrderItem,
                        Season, Shipment)
from app.validation import validate_system


def upload(db):
    season = db.scalar(select(Season))
    order = db.scalar(select(Order))
    batch = ImportBatch(filename='SS26 master.xlsx', sha256='a' * 64)
    db.add(batch)
    db.flush()
    coat = db.scalar(select(OrderItem))
    coat.source_file, coat.size = 'data/imports/master/1/SS26 master.xlsx', 'M'
    unallocated = OrderItem(order_id=order.id, article='DRESS-2', name='Silk dress', quantity=2, size=None,
                            wholesale_unit_price=None, markup_percentage=Decimal('0.30'), allocation_status='Review',
                            source_file='SS26 master.xlsx', source_sheet='Atelier', source_row=7)
    other_file = OrderItem(order_id=order.id, article='OTHER', name='Other file', quantity=1, customer_id=None,
                           markup_percentage=Decimal('0.30'), source_file='another.xlsx')
    db.add_all([unallocated, other_file,
                AllocationReview(import_batch_id=batch.id, source_file='SS26 master.xlsx', sheet='Atelier', row=7,
                                 raw_description='MA/KI', raw_values='{}', reason='Two customers named'),
                ImportReview(import_batch_id=batch.id, sheet='Atelier', row=9, raw_values='{}', reason='Broken formula',
                             status='Resolved')])
    db.commit()
    return batch, unallocated, season


def test_check_lists_problems_for_this_file_only(workspace_session, workspace_client):
    batch, item, _ = upload(workspace_session)
    page = workspace_client.get(f'/imports/{batch.id}/check')
    assert page.status_code == 200
    text = page.text
    assert 'Silk dress' in text and 'Wool coat' in text and 'Other file' not in text
    assert 'Split between customers' in text and 'Product is not confidently allocated' in text
    assert 'Missing or zero price' in text and 'Resolved' in text
    problems = workspace_client.get(f'/imports/{batch.id}/check?show=problems').text
    assert 'Silk dress' in problems and 'Wool coat' not in problems.split('id="products"')[1]
    assert workspace_client.get('/imports/latest/check', follow_redirects=False).headers['location'] == f'/imports/{batch.id}/check'


def test_bulk_mark_returns_to_check_and_export_is_coloured(workspace_session, workspace_client):
    batch, item, _ = upload(workspace_session)
    here = f'/imports/{batch.id}/check?show=all'
    response = workspace_client.post('/worksheet/bulk', data={'item_ids': [item.id], 'state': 'done', 'next': here},
                                     follow_redirects=False)
    assert response.headers['location'].startswith(here) and 'updated=1' in response.headers['location']
    workspace_session.refresh(item)
    assert item.worksheet_status == 'done'
    book = load_workbook(BytesIO(workspace_client.get(f'/imports/{batch.id}/check.xlsx').content))
    problems, products = book['Problems'], book['Products']
    assert problems['A2'].value == 'Needs fixing' and problems['A2'].fill.fgColor.rgb.endswith('F8D7D3')
    colours = {row[4].value: row[0].fill.fgColor.rgb[-6:] for row in products.iter_rows(min_row=2)}
    assert colours['Silk dress'] == 'D9EEDB'          # marked done → green
    assert colours['Wool coat'] == 'F8D7D3'           # not completed → red
    item_page = workspace_client.get(f'/worksheet/items/{item.id}?next=/imports/{batch.id}/check')
    assert f'/imports/{batch.id}/check' in item_page.text


def test_bad_master_workbook_shows_reason_instead_of_server_error(workspace_session, workspace_client, tmp_path, monkeypatch):
    monkeypatch.setattr(main, 'ROOT_DIR', tmp_path)

    def broken(db, path):
        raise KeyError('Worksheet SS25 does not exist.')
    monkeypatch.setattr(main, 'import_workbook', broken)
    response = workspace_client.post('/import', files={'file': ('orders.xlsx', b'not really excel')}, follow_redirects=False)
    assert response.status_code == 303 and response.headers['location'].startswith('/import?error=')
    page = workspace_client.get(response.headers['location'])
    assert 'could not be imported' in page.text and 'SS25' in page.text
    assert not list((tmp_path / 'data' / 'imports' / 'master').glob('*/*'))


def test_upload_redirects_to_its_check(workspace_session, workspace_client, tmp_path, monkeypatch):
    import hashlib
    monkeypatch.setattr(main, 'ROOT_DIR', tmp_path)
    content = b'fake workbook bytes'

    def fake_import(db, path):
        db.add(ImportBatch(filename=path.name, sha256=hashlib.sha256(content).hexdigest()))
        return {'items': 0}
    monkeypatch.setattr(main, 'import_workbook', fake_import)
    response = workspace_client.post('/import', files={'file': ('orders.xlsx', content)}, follow_redirects=False)
    batch = workspace_session.scalar(select(ImportBatch))
    assert response.headers['location'] == f'/imports/{batch.id}/check?uploaded=1'


def test_shipment_costs_and_expenses_can_be_corrected_with_audit(workspace_session, workspace_client):
    db = workspace_session
    season = db.scalar(select(Season))
    shipment = Shipment(season_id=season.id, name='Truck 1', fixed_cost=Decimal('10'))
    db.add(shipment)
    db.flush()
    expense = Expense(season_id=season.id, shipment_id=shipment.id, type='bag', amount=Decimal('5'))
    db.add(expense)
    db.commit()
    assert 'Edit costs' in workspace_client.get('/logistics?season=AW26').text
    response = workspace_client.post(f'/shipments/{shipment.id}/costs', data={
        'fixed_cost': '12.50', 'suitcase_cost': '3', 'bag_cost': '0', 'use_rule_transport': 'off',
        'transport_percentage': '7', 'shipment_date': '2026-10-01', 'status': 'In transit', 'notes': ''},
        follow_redirects=False)
    assert response.status_code == 303 and 'edited=' in response.headers['location']
    db.refresh(shipment)
    assert shipment.fixed_cost == Decimal('12.50') and shipment.transport_percentage == Decimal('0.07')
    assert not shipment.use_rule_transport and shipment.status == 'In transit'
    workspace_client.post(f'/expenses/{expense.id}', data={'amount': '6.40', 'expense_type': 'suitcase', 'note': 'Extra case'})
    db.refresh(expense)
    assert expense.amount == Decimal('6.40') and expense.type == 'suitcase'
    fields = set(db.scalars(select(AuditLog.field).where(AuditLog.entity_type.in_(['Shipment', 'Expense']))))
    assert {'fixed_cost', 'transport_percentage', 'amount', 'type'} <= fields
    assert workspace_client.post(f'/expenses/{expense.id}', data={'amount': '-1', 'expense_type': 'bag'}).status_code == 422


def test_system_check_survives_order_without_brand_rule(workspace_session):
    db = workspace_session
    season = db.scalar(select(Season))
    brand = Brand(name='No Rule Brand')
    db.add(brand)
    db.flush()
    order = Order(season_id=season.id, brand_id=brand.id)
    db.add(order)
    db.flush()
    db.add(OrderItem(order_id=order.id, article='X', quantity=1, wholesale_unit_price=Decimal('10'),
                     markup_percentage=Decimal('0.30'), customer_id=db.scalar(select(OrderItem)).customer_id, size='M'))
    db.commit()
    codes = {issue.code for issue in validate_system(db).issues}
    assert 'MISSING_BRAND_RULE' in codes
