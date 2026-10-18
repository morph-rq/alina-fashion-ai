"""Private customer workbook: an explicit allowlist of customer-facing facts.

Ported from the main copy's ``excel_views.customer_workbook`` with three fixes:
selling-price overrides are shown as the customer price, refunds are labelled and
signed, and the workbook records when it was generated. Wholesale costs,
markup, internal notes, source paths, rules and other customers never appear.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import select

from app.calculations import PAYMENT_TYPE_LABELS, item_amounts, money, signed_payment_amount
from app.models import Customer, Season
from app.services import customer_statement


def customer_workbook(db, customer_id: int, season_id: int | None = None) -> bytes:
    from app.excel_export import DATE_FORMAT, MONEY_FORMAT, _finish_widths, _new_data_sheet
    customer = db.get(Customer, customer_id)
    if customer is None:
        raise ValueError("Customer not found")
    seasons = list(db.scalars(select(Season).where(Season.id == season_id) if season_id else select(Season).order_by(Season.id)))
    if season_id and not seasons:
        raise ValueError("Season not found")
    book = Workbook()
    book.remove(book.active)
    generated = datetime.now().astimezone()
    subtitle = f"{customer.name or customer.code} · statement generated {generated:%Y-%m-%d %H:%M}"
    statements = [customer_statement(db, customer_id, season.name) for season in seasons]
    statements = [s for s in statements if s.items or s.payments]
    summary_rows = [[s.season.name, float(s.ordered), float(s.initial_prepayment), float(s.ready_balance_due),
                     float(s.allocated_costs), float(s.account.total_charged), float(s.account.total_paid),
                     float(s.account.outstanding), float(s.account.overpayment)] for s in statements]
    summary = _new_data_sheet(book, "Summary", subtitle, ["Season", "Order value", "Initial due", "Ready due", "Delivery charges",
                                                          "Charged to date", "Paid (net of refunds)", "Balance due", "Credit"], summary_rows, 1)
    _finish_widths(summary, [14] + [20] * 8)
    for row in summary.iter_rows(min_row=5, min_col=2):
        for cell in row:
            cell.number_format = MONEY_FORMAT
    rows = []
    for statement in statements:
        for item in statement.items:
            amounts = item_amounts(item)
            unit = money(amounts.final / item.quantity) if item.quantity else amounts.final
            rows.append([statement.season.name, item.order.brand.name, item.article, item.name, item.color, item.size, item.quantity,
                         float(unit), float(amounts.final), item.status, item.ready_batch.name if item.ready_batch else None,
                         item.shipment.name if item.shipment else None, statement.expected_receipts.get(item.ready_batch_id),
                         item.requested_delivery_date])
    products = _new_data_sheet(book, "Products", subtitle, ["Season", "Brand", "Article / SKU", "Product", "Color", "Size", "Quantity",
                                                            "Unit price", "Total", "Status", "Ready batch", "Shipment",
                                                            "Expected delivery", "Requested delivery"], rows, 2)
    _finish_widths(products, [12, 23, 22, 32, 16, 12, 10, 17, 17, 21, 18, 18, 20, 20])
    for row in range(5, 5 + len(rows)):
        for col in (8, 9):
            products.cell(row, col).number_format = MONEY_FORMAT
        for col in (13, 14):
            products.cell(row, col).number_format = DATE_FORMAT
    payment_rows = [[s.season.name, p.date, PAYMENT_TYPE_LABELS.get(p.type, p.type), float(signed_payment_amount(p.amount, p.type))]
                    for s in statements for p in s.payments]
    payments = _new_data_sheet(book, "Payments", subtitle, ["Season", "Date", "Type", "Amount"], payment_rows, 3)
    _finish_widths(payments, [16, 20, 22, 20])
    for row in range(5, 5 + len(payment_rows)):
        payments.cell(row, 2).number_format = DATE_FORMAT
        payments.cell(row, 4).number_format = MONEY_FORMAT
    summary.cell(7 + len(summary_rows), 1, "Amounts and delivery estimates reflect the records at the time shown above. "
                 "Delivery estimates are rule-based, not carrier tracking.")
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
