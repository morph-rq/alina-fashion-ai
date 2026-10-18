from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app import ai_intake
from app.database import Base, get_db
from app.main import app, projected_profit
from app.models import Brand, BrandSeasonRule, Collection, Customer, Order, OrderItem, Season


@pytest.fixture()
def workspace_session():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as session:
        season = Season(name="AW26")
        brand = Brand(name="Test Atelier", default_markup_percentage=Decimal("0.30"))
        customer = Customer(code="MA", name="Maria")
        session.add_all([season, brand, customer])
        session.flush()
        rule = BrandSeasonRule(
            season_id=season.id,
            brand_id=brand.id,
            markup_percentage=Decimal("0.30"),
            initial_prepayment_percentage=Decimal("0.50"),
            ready_balance_percentage=Decimal("0.50"),
            transport_percentage=Decimal("0.05"),
            special_cost=Decimal("0"),
            transport_exempt_country_codes="CH",
            expected_receipt_days=14,
        )
        order = Order(season_id=season.id, brand_id=brand.id, status="Ordered")
        session.add_all([rule, order])
        session.flush()
        session.add(OrderItem(
            order_id=order.id,
            customer_id=customer.id,
            article="COAT-1",
            name="Wool coat",
            quantity=1,
            wholesale_unit_price=Decimal("100"),
            markup_percentage=Decimal("0.30"),
        ))
        session.commit()
        yield session


@pytest.fixture()
def workspace_client(workspace_session: Session):
    def override_db():
        yield workspace_session

    app.dependency_overrides[get_db] = override_db
    try:
        with TestClient(app) as client:
            yield client
    finally:
        app.dependency_overrides.clear()


def test_library_supports_collections_and_standalone_brands(workspace_client: TestClient, workspace_session: Session) -> None:
    page = workspace_client.get("/library")
    assert page.status_code == 200
    assert "Business library" in page.text
    assert "All stored documents" in page.text
    assert "Product references" in page.text  # Product images were deliberately removed from the library in this copy.

    created = workspace_client.post("/library/collections", data={"name": "Client edit", "description": "For Maria"})
    assert created.status_code == 200
    assert workspace_session.scalar(select(Collection).where(Collection.name == "Client edit")) is not None

    added_brand = workspace_client.post(
        "/library/brands",
        data={"name": "New House", "default_markup": "42", "website_url": "https://example.com", "notes": "New supplier"},
    )
    assert added_brand.status_code == 200
    brand = workspace_session.scalar(select(Brand).where(Brand.name == "New House"))
    assert brand is not None
    assert brand.default_markup_percentage == Decimal("0.4200")
    assert brand.website_url == "https://example.com"


def test_customer_profile_keeps_all_contact_details(workspace_client: TestClient, workspace_session: Session) -> None:
    customer = workspace_session.scalar(select(Customer).where(Customer.code == "MA"))
    response = workspace_client.post(
        f"/customers/{customer.id}/settings",
        data={
            "return_to": "detail", "name": "Maria Alex", "phone": "+39 123", "email": "maria@example.com",
            "country_code": "IT", "address": "Via Roma 1", "city": "Milan", "postal_code": "20100",
            "preferred_contact": "WhatsApp", "notes": "Call after 17:00", "alias": "Masha",
        },
    )
    assert response.status_code == 200
    workspace_session.refresh(customer)
    assert (customer.name, customer.email, customer.city, customer.preferred_contact) == (
        "Maria Alex", "maria@example.com", "Milan", "WhatsApp",
    )
    assert {alias.alias for alias in customer.aliases} == {"MASHA"}


def test_profit_is_deterministic_and_ai_rule_change_requires_confirmation(
    workspace_client: TestClient, workspace_session: Session, monkeypatch,
) -> None:
    profit = projected_profit(workspace_session)
    assert profit is not None
    assert profit["projected_revenue"] == Decimal("130.00")
    assert profit["projected_profit"] == Decimal("30.00")

    rule = workspace_session.scalar(select(BrandSeasonRule))
    monkeypatch.setattr(
        ai_intake,
        "business_assistant",
        lambda _message, _context: ai_intake.AssistantReply(
            answer="I can propose changing the markup. Review it first.",
            suggested_page="/rules",
            rule_proposal=ai_intake.AssistantRuleProposal(
                rule_id=rule.id, markup_percentage=40, initial_percentage=50, ready_percentage=50,
                explanation="Requested markup update",
            ),
        ),
    )
    answer = workspace_client.post("/api/assistant/help", json={"message": "Make markup 40%"})
    assert answer.status_code == 200
    assert answer.json()["profit"]["projected_profit"] == 30.0
    workspace_session.refresh(rule)
    assert rule.markup_percentage == Decimal("0.3000")

    rejected = workspace_client.post(
        "/assistant/apply-rule",
        data={"rule_id": rule.id, "markup_percentage": "40", "confirmation": "not-confirmed"},
    )
    assert rejected.status_code == 422
    workspace_session.refresh(rule)
    assert rule.markup_percentage == Decimal("0.3000")

    applied = workspace_client.post(
        "/assistant/apply-rule",
        data={
            "rule_id": rule.id, "markup_percentage": "40", "initial_percentage": "50", "ready_percentage": "50",
            "explanation": "Requested markup update", "confirmation": "APPLY_AI_RULE_PROPOSAL",
        },
    )
    assert applied.status_code == 200
    workspace_session.refresh(rule)
    assert rule.markup_percentage == Decimal("0.4000")


def test_every_page_uses_permanent_light_theme_and_ai_help(workspace_client: TestClient) -> None:
    page = workspace_client.get("/")
    assert page.status_code == 200
    assert 'data-theme="light"' in page.text
    assert "AI Help" in page.text
    assert "fashion-theme" not in page.text
    assert "theme-toggle" not in page.text
    assert '/static/refinement.css?v=20260827-2' in page.text

    mobile_css = workspace_client.get("/static/refinement.css")
    assert mobile_css.status_code == 200
    assert "calc(112px + env(safe-area-inset-bottom))" in mobile_css.text
    assert "scroll-padding-bottom" in mobile_css.text


def test_global_search_finds_records_and_excel_export_is_editable(workspace_client: TestClient) -> None:
    product_search = workspace_client.get("/api/search", params={"q": "COAT-1"})
    assert product_search.status_code == 200
    assert any(result["type"] == "Product" and result["title"] == "Wool coat" for result in product_search.json()["results"])

    customer_search = workspace_client.get("/api/search", params={"q": "Maria"})
    assert any(result["type"] == "Customer" for result in customer_search.json()["results"])

    exported = workspace_client.get("/exports/business.xlsx")
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats")
    workbook = load_workbook(BytesIO(exported.content), data_only=False)
    assert {"Summary", "Products", "Orders", "Customers", "Rules", "Collections"}.issubset(workbook.sheetnames)
    products = workbook["Products"]
    assert products["F5"].value == "COAT-1"
    assert products["M5"].value == '=IF(O5="Cancelled",0,ROUND(K5*(1+L5),2))'
    assert products["N5"].value == '=IF(O5="Cancelled",0,ROUND(K5*J5,2)+ROUND(ROUND(K5*J5,2)*L5,2))'
    assert products.freeze_panes == "A5"
    assert products.tables
