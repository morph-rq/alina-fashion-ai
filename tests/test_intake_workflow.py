"""Acceptance tests for the persistent intake workflow.

upload / import evidence -> draft rows -> focused questions -> corrections ->
validate -> save once -> export, with progress preserved between requests.
AI is replaced by a fake provider; no live service is called.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal as D
from io import BytesIO
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import ai_intake, source_store
from app.database import Base, get_db
from app.main import app
from app.models import (
    AuditLog, Brand, BrandSeasonRule, Customer, DraftQuestion, DraftRow, IntakeSession, OperationReceipt, Order, OrderItem,
    Payment, Season, SourceBlob, SourceItem,
)

def _png(color="navy"):
    from PIL import Image
    stream = BytesIO()
    Image.new("RGB", (40, 40), color).save(stream, "PNG")
    return stream.getvalue()


PNG = _png()


@pytest.fixture()
def env(tmp_path, monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(source_store, "ROOT_DIR", tmp_path)
    monkeypatch.setattr("app.image_manager.ROOT_DIR", tmp_path, raising=False)
    with factory() as db:
        season = Season(name="FW26")
        row_brand, zim = Brand(name="The Row"), Brand(name="Zimmermann")
        maria = Customer(code="MI", name="Maria Ivanova", country_code="IT")
        olga = Customer(code="OP", name="Olga Petrova", country_code="IT")
        db.add_all([season, row_brand, zim, maria, olga])
        db.flush()
        for brand, markup in ((row_brand, "0.30"), (zim, "0.25")):
            db.add(BrandSeasonRule(season_id=season.id, brand_id=brand.id, markup_percentage=D(markup),
                                   initial_prepayment_percentage=D("0.50"), ready_balance_percentage=D("0.50"),
                                   transport_percentage=D("0.05"), transport_exempt_country_codes="CH", expected_receipt_days=14))
            db.add(Order(season_id=season.id, brand_id=brand.id))
        db.commit()

    def override():
        with factory() as session:
            yield session

    app.dependency_overrides[get_db] = override
    client = TestClient(app)
    yield client, factory, tmp_path
    app.dependency_overrides.clear()


def fake_ai(monkeypatch, *bundles, check=None):
    calls = []
    queue = list(bundles)

    def analyze(message, attachments):
        if check:
            check(message, attachments)
        calls.append((message, [name for name, _mime, _content in attachments]))
        result = queue.pop(0) if len(queue) > 1 else queue[0]
        if isinstance(result, Exception):
            raise result
        return result

    monkeypatch.setattr(ai_intake, "ai_configured", lambda: True)
    monkeypatch.setattr(ai_intake, "analyze_submission", analyze)
    return calls


def facts(**values):
    return ai_intake.IntakeFacts(**values)


def bundle(*orders):
    return ai_intake.IntakeBundle(orders=list(orders), source_summary="test")


def start(client, factory, message="", files=()):
    with factory() as db:
        season_id = db.scalar(select(Season.id))
    response = client.post("/intake", data={"season_id": season_id, "message": message},
                           files=[("attachments", file) for file in files], follow_redirects=False)
    assert response.status_code == 303, response.text
    return response.headers["location"].rsplit("/", 1)[-1]


def session_of(factory, public_id):
    db = factory()
    return db, db.scalar(select(IntakeSession).where(IntakeSession.public_id == public_id))


def open_questions(factory, public_id):
    db, session = session_of(factory, public_id)
    return {q.field: q for q in session.questions if q.status == "open"}


def answer(client, public_id, question, choice=None, text=""):
    data = {"text": text}
    if choice is not None:
        data["choice"] = str(choice)
    response = client.post(f"/intake/{public_id}/questions/{question.id}", data=data, follow_redirects=False)
    assert response.status_code == 303 and "error=" not in response.headers["location"], response.headers["location"]


def save(client, factory, public_id, key="key-0123456789abcdef", then=""):
    db, session = session_of(factory, public_id)
    return client.post(f"/intake/{public_id}/commit", data={"operation_key": key, "expected_revision": session.revision, "then": then},
                       follow_redirects=False)


def test_upload_questions_correction_save_once_and_excel(env, monkeypatch):
    client, factory, root = env
    seen = {}

    def check(message, attachments):
        # Capture is durable before extraction starts.
        with factory() as db:
            blobs = db.scalars(select(SourceBlob)).all()
            seen["stored"] = all((root / blob.storage_path).is_file() for blob in blobs) and len(blobs) >= 2

    fake_ai(monkeypatch, bundle(facts(customer_name="Maria Ivanova", brand="The Row", product_name="Black coat", article="0451",
                                      color="black", stated_price="1200", source_attachment_names=["coat.png"],
                                      evidence_quotes=["Maria Ivanova, The Row, black coat, 38, 1200"])), check=check)
    public_id = start(client, factory, "Maria Ivanova, The Row, black coat, 1200", [("coat.png", PNG, "image/png")])
    assert seen["stored"]
    questions = open_questions(factory, public_id)
    # Customer and brand table were matched exactly, so they are not asked again.
    assert "customer" not in questions and "order" not in questions
    assert {"size", "price_meaning", "quantity"} <= set(questions)
    page = client.get(f"/intake/{public_id}")
    assert page.status_code == 200 and "Black coat" in page.text and "wholesale (factory) price" in page.text

    # The operator's natural answer fills several facts on the same row.
    client.post(f"/intake/{public_id}/messages", data={"message": "Wholesale, in euros. One piece, FW26. Size 38"})
    db, session = session_of(factory, public_id)
    assert len([row for row in session.rows if row.status == "draft"]) == 1
    assert session.state == "ready_for_review"
    client.post(f"/intake/{public_id}/messages", data={"message": "Actually size 40"})
    db, session = session_of(factory, public_id)
    rows = [row for row in session.rows if row.status == "draft"]
    assert len(rows) == 1
    size = __import__("json").loads(rows[0].fields_json)["size"]
    assert size["value"] == "40" and size["state"] == "confirmed" and size["history"][-1]["value"] == "38"

    first = save(client, factory, public_id, then="excel")
    assert first.status_code == 303 and "/receipts/" in first.headers["location"]
    replay = save(client, factory, public_id)  # Same key again: double click or lost response.
    assert replay.status_code == 303 and replay.headers["location"] == first.headers["location"].split("?")[0]
    with factory() as db:
        items = db.scalars(select(OrderItem)).all()
        assert len(items) == 1
        item = items[0]
        assert (item.article, item.size, item.quantity, item.wholesale_unit_price) == ("0451", "40", 1, D("1200.00"))
        assert item.customer.code == "MI" and item.order.brand.name == "The Row" and item.image_path
        assert db.scalar(select(func.count()).select_from(OperationReceipt)) == 1
        season_id = item.order.season_id
    # A new key after the draft is already saved still cannot create goods again.
    assert save(client, factory, public_id, key="another-key-0123456789").status_code == 303
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(OrderItem)) == 1
    receipt = client.get(first.headers["location"])
    assert receipt.status_code == 200 and "Season workbook" in receipt.text
    with factory() as db:
        stored = __import__("json").loads(db.scalar(select(OperationReceipt.result_json)))
        maria_id = str(db.scalar(select(Customer.id).where(Customer.code == "MI")))
    # The receipt's "after" balances include the product that was just saved (1200 + 30% markup).
    assert D(stored["before"][maria_id]["ordered"]) == 0 and D(stored["after"][maria_id]["ordered"]) == D("1560.00")
    for _attempt in range(2):  # Export retries never mutate business records.
        workbook = client.get(f"/exports/seasons/{season_id}.xlsx")
        assert workbook.status_code == 200
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(OrderItem)) == 1
    sheet = load_workbook(BytesIO(workbook.content))["Products"]
    values = [cell.value for cell in sheet[5]]
    assert values[5] == "0451" and sheet.cell(5, 6).data_type == "s" and values[8] == "40"


def test_several_customers_and_products_become_separate_rows_and_exports_stay_private(env, monkeypatch):
    client, factory, _root = env
    fake_ai(monkeypatch, bundle(
        facts(customer_name="Maria Ivanova", brand="The Row", product_name="Coat", explicit_size="38", quantity=1, quantity_explicit=True,
              stated_price="1000", price_type="wholesale", price_currency="EUR"),
        facts(customer_name="Olga Petrova", brand="Zimmermann", product_name="Dress", explicit_size="S", quantity=2, quantity_explicit=True,
              stated_price="300", price_type="wholesale", price_basis="unit"),
        facts(customer_name="Maria Ivanova", brand="The Row", product_name="Coat", explicit_size="38", quantity=1, quantity_explicit=True,
              stated_price="1000", price_type="wholesale"),  # Repeated evidence reinforces, never duplicates.
    ))
    public_id = start(client, factory, "two customers", [("list.png", PNG, "image/png")])
    db, session = session_of(factory, public_id)
    assert len([row for row in session.rows if row.status == "draft"]) == 2
    assert session.state == "ready_for_review", [q.text for q in session.questions if q.status == "open"]
    assert save(client, factory, public_id).status_code == 303
    with factory() as db:
        rows = {item.customer.code: item for item in db.scalars(select(OrderItem)).all()}
        assert rows["OP"].quantity == 2 and rows["OP"].wholesale_unit_price == D("300.00") and rows["OP"].markup_percentage == D("0.25")
        maria_id, olga_id = rows["MI"].customer_id, rows["OP"].customer_id
    private = client.get(f"/exports/customers/{maria_id}.xlsx?season=FW26")
    assert private.status_code == 200
    book = load_workbook(BytesIO(private.content))
    text = " ".join(str(cell.value) for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert "Olga" not in text and "Dress" not in text and "Coat" in text


def test_missing_size_is_asked_never_inferred_and_draft_survives_restart(env, monkeypatch):
    client, factory, _root = env
    fake_ai(monkeypatch, bundle(facts(customer_name="Maria Ivanova", brand="The Row", product_name="Coat", height_cm=170,
                                      quantity=1, quantity_explicit=True, stated_price="1000", price_type="wholesale")))
    public_id = start(client, factory, "Maria is 170 cm", [("note.txt", b"Maria 170cm coat 1000 wholesale", "text/plain")])
    questions = open_questions(factory, public_id)
    assert set(questions) == {"size"}
    blocked = save(client, factory, public_id)
    assert "error=" in blocked.headers["location"]
    # A new process sees the same draft and the same single question.
    app.dependency_overrides.clear()
    fresh = sessionmaker(bind=factory.kw["bind"], expire_on_commit=False)
    app.dependency_overrides[get_db] = lambda: (yield from [fresh()])
    page = client.get(f"/intake/{public_id}")
    assert "What size?" in page.text and "Coat" in page.text
    answer(client, public_id, open_questions(fresh, public_id)["size"], text="one size")
    assert save(client, fresh, public_id).status_code == 303
    with fresh() as db:
        assert db.scalar(select(OrderItem.size)) == "One size"


def test_conflicting_prices_keep_both_candidates_and_record_the_choice(env, monkeypatch):
    client, factory, _root = env
    common = dict(customer_name="Maria Ivanova", brand="The Row", product_name="Coat", explicit_size="38", quantity=1,
                  quantity_explicit=True, price_type="wholesale")
    fake_ai(monkeypatch, bundle(facts(stated_price="1000", **common)), bundle(facts(stated_price="1100", **common)))
    public_id = start(client, factory, "screenshot", [("a.png", PNG, "image/png")])
    client.post(f"/intake/{public_id}/messages", data={"message": ""}, files=[("attachments", ("b.png", _png("black"), "image/png"))])
    question = open_questions(factory, public_id)["price"]
    labels = [option["label"] for option in __import__("json").loads(question.options_json)]
    assert any("1000.00" in label for label in labels) and any("1100.00" in label for label in labels)
    answer(client, public_id, question, choice=1)
    db, session = session_of(factory, public_id)
    price = __import__("json").loads(session.rows[0].fields_json)["price"]
    assert price["value"] == "1100.00" and price["state"] == "confirmed" and price["history"]


def test_ai_unavailable_keeps_sources_and_manual_rows_still_save(env, monkeypatch):
    client, factory, _root = env
    fake_ai(monkeypatch, ValueError("AI order entry is temporarily unavailable. Try again in one minute."))
    public_id = start(client, factory, "", [("photo.png", PNG, "image/png")])
    db, session = session_of(factory, public_id)
    upload = next(source for source in session.sources if source.channel == "upload")
    assert upload.state == "failed" and upload.blob_id and "temporarily unavailable" in upload.state_reason
    assert any("could not analyse" in message.text for message in session.messages)
    client.post(f"/intake/{public_id}/rows", data={"kind": "new_item"})
    db, session = session_of(factory, public_id)
    row = session.rows[0]
    response = client.post(f"/intake/{public_id}/rows/{row.id}/fields", data={
        "expected_revision": row.revision, "field:customer": "Maria Ivanova", "field:order": "The Row", "field:name": "Bag",
        "field:size": "One size", "field:quantity": "1", "field:price": "500", "field:price_meaning": "wholesale"}, follow_redirects=False)
    assert "error=" not in response.headers["location"], response.headers["location"]
    assert save(client, factory, public_id).status_code == 303
    with factory() as db:
        assert db.scalar(select(OrderItem.name)) == "Bag"


def test_stale_review_is_refused_and_double_answer_is_harmless(env, monkeypatch):
    client, factory, _root = env
    fake_ai(monkeypatch, bundle(facts(customer_name="Maria Ivanova", brand="The Row", product_name="Coat", quantity=1, quantity_explicit=True,
                                      stated_price="1000", price_type="wholesale")))
    public_id = start(client, factory, "coat", [("a.png", PNG, "image/png")])
    question = open_questions(factory, public_id)["size"]
    answer(client, public_id, question, text="38")
    answer(client, public_id, question, text="40")  # Second submission of an answered question changes nothing.
    db, session = session_of(factory, public_id)
    assert __import__("json").loads(session.rows[0].fields_json)["size"]["value"] == "38"
    reviewed = session.revision
    row = session.rows[0]
    client.post(f"/intake/{public_id}/rows/{row.id}/fields", data={"field:color": "navy"})
    stale = client.post(f"/intake/{public_id}/commit", data={"operation_key": "stale-key-0123456789", "expected_revision": reviewed},
                        follow_redirects=False)
    assert "error=" in stale.headers["location"] and "changed after you reviewed" in __import__("urllib.parse").parse.unquote(stale.headers["location"])
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(OrderItem)) == 0


def test_correcting_saved_goods_targets_one_row_and_keeps_history(env, monkeypatch):
    client, factory, _root = env
    with factory() as db:
        order = db.scalar(select(Order).join(Brand).where(Brand.name == "The Row"))
        maria = db.scalar(select(Customer).where(Customer.code == "MI"))
        item = OrderItem(order_id=order.id, customer_id=maria.id, name="Black coat", size="38", quantity=1,
                         wholesale_unit_price=D("1000"), markup_percentage=D("0.30"))
        db.add(item)
        db.commit()
        item_id = item.id
    public_id = start(client, factory, "Change Maria's coat from size 38 to 40")
    db, session = session_of(factory, public_id)
    assert session.rows[0].kind == "update_item" and session.rows[0].target_item_id == item_id
    assert save(client, factory, public_id).status_code == 303
    with factory() as db:
        assert db.get(OrderItem, item_id).size == "40"
        assert db.scalar(select(func.count()).select_from(OrderItem)) == 1
        audit = db.scalar(select(AuditLog).where(AuditLog.field == "size"))
        assert audit.previous_value == "38" and "Change Maria" in audit.description


def test_payment_message_creates_a_payment_draft_with_duplicate_warning(env, monkeypatch):
    client, factory, _root = env
    with factory() as db:
        maria = db.scalar(select(Customer).where(Customer.code == "MI"))
        db.add(Payment(season_id=db.scalar(select(Season.id)), customer_id=maria.id, date=date(2026, 8, 1), amount=D("650"), type="payment"))
        db.commit()
    public_id = start(client, factory, "Record Maria Ivanova payment of 650 EUR on 2026-09-01")
    page = client.get(f"/intake/{public_id}")
    assert "Possible duplicate" in page.text
    assert save(client, factory, public_id).status_code == 303
    with factory() as db:
        payments = db.scalars(select(Payment).order_by(Payment.id)).all()
        assert len(payments) == 2 and payments[-1].date == date(2026, 9, 1) and payments[-1].amount == D("650.00")


def test_final_price_is_preserved_exactly(env, monkeypatch):
    client, factory, _root = env
    fake_ai(monkeypatch, bundle(facts(customer_name="Olga Petrova", brand="Zimmermann", product_name="Dress", explicit_size="S", quantity=1,
                                      quantity_explicit=True, stated_price="100", price_type="final")))
    public_id = start(client, factory, "final price 100", [("a.png", PNG, "image/png")])
    assert save(client, factory, public_id).status_code == 303
    from app.calculations import item_amounts
    with factory() as db:
        item = db.scalar(select(OrderItem))
        assert item_amounts(item).final == D("100.00")


def test_failed_storage_records_nothing_as_captured(env, monkeypatch):
    client, factory, _root = env

    def broken(content, root):
        raise source_store.CaptureError("disk full")

    monkeypatch.setattr(source_store, "_store_bytes", broken)
    with factory() as db:
        season_id = db.scalar(select(Season.id))
    response = client.post("/intake", data={"season_id": season_id}, files=[("attachments", ("a.png", PNG, "image/png"))])
    assert response.status_code == 422
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(SourceItem)) == 0
