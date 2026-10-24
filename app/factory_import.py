"""Deterministic reader for supplier/factory order workbooks (Excel), with operator review.

Flow: ``analyze_workbook`` (pure; nothing saved) → review screen (map new customer
initials, resolve ambiguous allocations, confirm factory percentages) → ``commit_import``.

Layouts handled (all seen in real SS24/SS27 files): one sheet per brand; header rows found
by their labels in English/Italian/French/Russian; size matrices (size labels across the
header, quantities in the cells) including per-row size grids (A/B/C letters); a single size
column with lists like "40, 42"; an unlabeled allocation column ("AA 40, Alina 38"); several
tables on one sheet; ``TOTAL`` → percentage row → grand total, with different percentages
per section (e.g. RTW 23 %, bags 25 %) and ``ROUNDUP``; Sped/Ready columns; photos anchored
to rows; damaged photos. A summary sheet ("ЗАКАЗЫ …") is recognised and not read as goods.

Money is never guessed: prices come from cells, totals are recomputed and compared with the
supplier's cached totals, and any difference is reported. Uploaded text is data, never
instructions.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.calculations import ZERO, money
from app.config import ROOT_DIR
from app import xlsx_stream

ORIGINALS_DIR = ROOT_DIR / "data" / "originals"
IMAGES_DIR = ROOT_DIR / "data" / "images" / "factory"
MAX_IMPORT_BYTES = 1024 * 1024 * 1024
MAX_SHEET_ROWS = 20000

HEADER_TOKENS = {
    "name": {"name", "style name", "description", "наименование", "название", "модель", "product", "item",
             "descrizione", "nom", "style", "type"},
    "article": {"article", "sku", "code", "style number", "article number", "артикул", "art", "codice",
                "style code", "reference", "ref", "model"},
    "color": {"color", "colour", "color description", "colore", "цвет", "couleur", "color name"},
    "color_code": {"color code", "colour code", "color number", "color id"},
    "size": {"size", "размер", "taglia", "taille", "t size", "sizes"},
    "qty": {"q ty", "qty", "quantity", "total qty", "кол во", "количество", "pcs", "qt", "units", "q.ty", "qta",
            "sell in qty"},
    "price": {"price", "price euro", "cost", "whl", "wholesale", "wholesale eur", "pvg", "pvg tbc", "цена",
              "prezzo", "prix", "unit price", "whs"},
    "retail": {"retail", "srp", "srp tbc", "rrp", "retail price"},
    "total": {"total", "total amt", "total value", "total ammount", "total amount", "total euro", "сумма",
              "amount", "totale", "итого"},
    "image": {"foto", "photo", "pic", "image", "picture", "pictures", "style image", "фото"},
    "look": {"look", "look number", "look sfilata", "look kering"},
    "category": {"category", "категория"},
    "composition": {"composition", "material", "fabrication", "состав"},
}
LABEL_TO_FIELD = {label: name for name, labels in HEADER_TOKENS.items() for label in labels}
SIZE_RE = re.compile(r"^(?:T\s?)?(?:\d{1,2}(?:[.,/]5|/)?|XXXS|XXS|XS|S|M|L|XL|XXL|XXXL|[2-5]XL|OS|UNI|TU)"
                     r"(?:\s*/\s*(?:XXS|XS|S|M|L|XL|XXL))?$", re.I)
TOTAL_WORDS = re.compile(r"\b(total|итого|всего|totale|sum)\b", re.I)
SHIP_HEADER = re.compile(r"(sped|spedizione|ready|отгруз|готов)", re.I)
NOT_CUSTOMER = {"ok", "no", "yes", "black", "white", "blue", "red", "pink", "beige", "grey", "gray", "green", "uni",
                "bags", "bag", "acc", "rtw", "look", "size", "total", "ready", "sped", "new", "sale", "changed",
                "comission", "commission", "the", "and", "with", "for", "milk", "silver", "cream", "olive", "khaki",
                "nero", "foto", "price", "qty", "sku", "eur", "euro", "tbc", "cancelled", "canceled", "cancel", "arrived",
                "delivered", "not", "pcs", "psc", "pc", "shipped", "sped", "spedito", "paid", "x", "or", "only",
                # garments and colours that customer reports (SS24 "Отчет_" files) write next to sizes
                "belt", "dress", "shirt", "skirt", "coat", "jacket", "pants", "pant", "trousers", "top", "knit", "sweater",
                "cardigan", "scarf", "shoes", "boots", "sneakers", "hat", "cap", "tshirt", "t-shirt", "jeans", "shorts", "blouse",
                "melange", "nougat", "snow", "navy", "camel", "brown", "ivory", "ecru", "sand", "taupe", "gold", "bordeaux",
                "burgundy", "orange", "yellow", "purple", "multi", "multicolor", "natural", "off", "dark", "light"}
STATUS_WORDS = {"CANCELLED": "Cancelled", "CANCELED": "Cancelled", "ОТМЕНА": "Cancelled", "DELIVERED": "Delivered",
                "ARRIVED": "Arrived", "READY": "Ready", "ГОТОВО": "Ready"}


def norm_label(value: Any) -> str:
    text = re.sub(r"[\s_]+", " ", re.sub(r"[^\w\s/]", " ", str(value or "").lower())).strip()
    return text


def parse_number(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(value))
            return number if number.is_finite() else None
        except InvalidOperation:
            return None
    text = str(value).strip().replace("€", "").replace("EUR", "").replace(" ", "").replace(" ", "")
    if not text or not re.fullmatch(r"-?[\d.,]+", text):
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
    elif "," in text:
        head, _, tail = text.rpartition(",")
        text = f"{head}.{tail}" if len(tail) != 3 else head + tail
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def size_label(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if isinstance(value, (int, float)) and not (0 <= float(value) <= 60):
        return None
    return text if SIZE_RE.match(text) else None


# --- Data structures ----------------------------------------------------------------

@dataclass
class Unit:
    customer: str | None
    size: str | None


@dataclass
class ItemRow:
    sheet: str
    row: int
    name: str | None
    article: str | None
    color: str | None
    quantity: int
    unit_price: str | None
    sizes: list[str]
    allocation_text: str | None
    units: list[dict]
    confident: bool
    issue: str | None
    section: str | None
    image_part: str | None
    status_hint: str | None
    shipment_hint: str | None
    fill: str | None
    look: str | None
    extra: dict = field(default_factory=dict)

    @property
    def base(self) -> Decimal:
        return money(Decimal(self.unit_price or 0) * self.quantity)


@dataclass
class SectionRate:
    sheet: str
    label: str
    rate: str | None
    rounding: str
    evidence: str
    source_total: str | None
    rows: list[int]
    share: str | None = None      # set when one block is split into portions (e.g. total/2 at 25 % and 30 %)


@dataclass
class SheetResult:
    name: str
    kind: str                 # order | summary | empty | unreadable
    brand: str
    items: list[ItemRow]
    sections: list[SectionRate]
    header_rows: list[int]
    columns: dict
    notes: list[str]
    source_grand_total: str | None
    computed_base: str = "0"


@dataclass
class Analysis:
    filename: str
    sha256: str
    size_bytes: int
    season: str | None
    sheets: list[SheetResult]
    tokens: dict[str, int]
    payments: list[dict]
    warnings: list[str]

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False, default=str)

    @staticmethod
    def from_json(text: str) -> "Analysis":
        data = json.loads(text)
        sheets = []
        for s in data["sheets"]:
            items = [ItemRow(**i) for i in s["items"]]
            sections = [SectionRate(**x) for x in s["sections"]]
            sheets.append(SheetResult(**{**s, "items": items, "sections": sections}))
        return Analysis(**{**data, "sheets": sheets})

    @property
    def items(self) -> list[ItemRow]:
        return [item for sheet in self.sheets for item in sheet.items]


# --- Allocation parsing -------------------------------------------------------------

ALLOC_TOKEN = re.compile(r"[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё.]*")
SIZE_TOKEN = re.compile(r"\d{1,2}(?:[.,]5)?(?:\s*\(\s*[A-Za-z]{1,3}\s*\))?|XXXS|XXS|XS|XXXL|XXL|XL|[2-5]XL|\b[SML]\b|OS|UNI", re.I)


HOMOGLYPHS = str.maketrans("АВСЕНКМОРТХаеосрх", "ABCEHKMOPTXaeocpx")
LATIN_TO_CYR = str.maketrans("ABCEHKMOPTXaeocpx", "АВСЕНКМОРТХаеосрх")
TRANSLIT = {"А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Е": "E", "Ё": "E", "Ж": "ZH", "З": "Z", "И": "I", "Й": "Y",
            "К": "K", "Л": "L", "М": "M", "Н": "N", "О": "O", "П": "P", "Р": "R", "С": "S", "Т": "T", "У": "U", "Ф": "F",
            "Х": "KH", "Ц": "TS", "Ч": "CH", "Ш": "SH", "Щ": "SCH", "Ъ": "", "Ы": "Y", "Ь": "", "Э": "E", "Ю": "YU", "Я": "YA"}


def canonical_token(word: str, known: set[str]) -> str:
    """Unify initials typed in mixed scripts: Cyrillic 'СС' = 'CC', 'Алина' = 'ALINA' when ALINA is known."""
    upper = word.upper()
    if upper in known:
        return upper
    latin = upper.translate(HOMOGLYPHS)
    if re.fullmatch(r"[A-Z.]+", latin):
        return latin
    translit = "".join(TRANSLIT.get(ch, ch) for ch in upper).replace("YA", "YA")
    for candidate in (translit, translit.replace("YA", "IA"), translit.rstrip("A") + "A"):
        if candidate in known:
            return candidate
    if re.search(r"[A-Z]", upper) and re.search(r"[А-ЯЁ]", upper):
        return upper.translate(LATIN_TO_CYR)
    return upper


def split_allocation(text: str | None, known: set[str]) -> list[tuple[str, list[str]]]:
    """'AA 40, Alina 38' -> [('AA', ['40']), ('ALINA', ['38'])]. Tokens are upper-cased."""
    if not text:
        return []
    raw = str(text).replace("\n", " ")
    raw = re.sub(r"(\d)\s*,\s*5\b", r"\1.5", raw)          # "39, 5" -> 39.5 (half size)
    raw = re.sub(r"(\d+)/(\d)\b", r"\1.\2", raw)             # "40/5" -> 40.5
    segments: list[tuple[str, list[str]]] = []
    position = 0
    pattern = re.compile(r"[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё.]*(?:\s+(?=[A-Za-zА-Яа-яЁё]))?")
    tokens = []
    for match in re.finditer(r"[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё.]*", raw):
        word = match.group(0).strip(".")
        upper = canonical_token(word, known)
        if len(upper) == 1 and upper not in known:
            continue  # a stray letter such as "44б"
        if SIZE_RE.match(word) and upper not in known and len(word) <= 3 and upper in {"XXS", "XS", "S", "M", "L", "XL", "XXL", "XXXL", "OS", "UNI"}:
            continue
        tokens.append((match.start(), match.end(), upper))
    merged = []
    for start, end, word in tokens:
        if merged and raw[merged[-1][1]:start].strip() == "" and f"{merged[-1][2]} {word}" in known:
            merged[-1] = (merged[-1][0], end, f"{merged[-1][2]} {word}")
        else:
            merged.append((start, end, word))
    for index, (start, end, word) in enumerate(merged):
        nxt = merged[index + 1][0] if index + 1 < len(merged) else len(raw)
        segment = raw[end:nxt]
        sizes = []
        for size in SIZE_TOKEN.findall(segment):
            size = re.sub(r"\s*\(.*\)", "", size).replace(",", ".").strip()
            sizes.append(size.upper())
        segments.append((word, sizes))
        position = nxt
    return [(word, sizes) for word, sizes in segments if word.lower() not in NOT_CUSTOMER and "." not in word.strip(".")]


def status_from_text(text: str | None) -> str | None:
    words = {w.upper() for w in re.findall(r"[A-Za-zА-Яа-яЁё]+", text or "")}
    if "NOT" in words and "ARRIVED" in words:
        return None
    return next((STATUS_WORDS[w] for w in ("CANCELLED", "CANCELED", "ОТМЕНА", "DELIVERED", "ARRIVED", "READY", "ГОТОВО") if w in words), None)


def norm_size(value: str | None) -> str | None:
    if value is None:
        return None
    text = str(value).strip().upper().replace(",", ".")
    text = re.sub(r"^T\s?", "", text)
    text = text.split("/")[0].strip() if re.match(r"^\d{2}\s*/\s*[A-Z]", text) else text
    text = re.sub(r"\.0$", "", text)
    return text or None


def allocate_units(segments: list[tuple[str, list[str]]], sizes: list[str | None], quantity: int) -> tuple[list[Unit], bool, str | None]:
    """Match parsed customer segments to the row's units. Returns (units, confident, issue)."""
    distinct = {norm_size(s) for s in sizes if s}
    explicit = [x for _, ss in segments for x in ss]
    if len(distinct) <= 1 and len(explicit) == quantity and all(ss for _, ss in segments):
        # One size cell ("38/2", "38.42") cannot describe several pieces; the initials list them exactly.
        return [Unit(c, norm_size(x)) for c, ss in segments for x in ss], True, None
    if len(distinct) == 1 and len(sizes) < quantity and not explicit:
        sizes = list(sizes) + [next(iter(distinct))] * (quantity - len(sizes))   # "UNI" for 2 pieces = both UNI
    sizes = [norm_size(s) for s in sizes] + [None] * max(0, quantity - len(sizes))
    sizes = sizes[:quantity] if quantity else sizes
    if quantity <= 0:
        return [], False, "Количество не указано"
    if not segments:
        return [Unit(None, s) for s in sizes], False, "Клиент не указан"
    if len(segments) == 1 and not segments[0][1]:
        return [Unit(segments[0][0], s) for s in sizes], True, None
    if len(segments) == 1 and segments[0][1] == [str(quantity)] and quantity <= 10 and not any(sizes):
        return [Unit(segments[0][0], None) for _ in range(quantity)], True, None   # "Alina 2" = two pieces
    sized = [(c, [norm_size(x) for x in ss]) for c, ss in segments if ss]
    unsized = [c for c, ss in segments if not ss]
    claimed = sum(len(ss) for _, ss in sized)
    known_sizes = [s for s in sizes if s]
    if sized and not unsized and claimed == quantity:
        if known_sizes and len(known_sizes) == quantity and Counter(known_sizes) != Counter(s for _, ss in sized for s in ss):
            pool = Counter(known_sizes)
            wanted = Counter(s for _, ss in sized for s in ss)
            return ([Unit(None, s) for s in sizes], False,
                    f"Размеры в описании ({', '.join(sorted(wanted.elements()))}) не совпадают с размерами строки ({', '.join(sorted(pool.elements()))})")
        return [Unit(c, s) for c, ss in sized for s in ss], True, None
    if sized and len(unsized) == 1 and claimed < quantity:
        remaining = Counter(known_sizes) - Counter(s for _, ss in sized for s in ss) if known_sizes else None
        units = [Unit(c, s) for c, ss in sized for s in ss]
        rest = list(remaining.elements()) if remaining is not None and sum(remaining.values()) == quantity - claimed else [None] * (quantity - claimed)
        units += [Unit(unsized[0], s) for s in rest]
        return units, True, None
    if not sized and len(unsized) == quantity:
        if len(set(known_sizes)) <= 1:
            return [Unit(c, sizes[i]) for i, c in enumerate(unsized)], True, None
        return ([Unit(None, s) for s in sizes], False,
                f"Неясно, какой размер у кого: {', '.join(unsized)} — размеры {', '.join(s for s in sizes if s)}")
    if not sized and len(unsized) > 1:
        return ([Unit(None, s) for s in sizes], False,
                f"{len(unsized)} клиента ({', '.join(unsized)}) на {quantity} шт. — неясно, кому сколько")
    return ([Unit(None, s) for s in sizes], False,
            f"Описание «{'; '.join(c + ' ' + ','.join(ss) for c, ss in segments)}» даёт {claimed} шт., а в строке {quantity}")


# --- Workbook reading ---------------------------------------------------------------

def _grid(path: Path, sheet: str, data_only: bool) -> dict[int, dict[int, Any]]:
    grid: dict[int, dict[int, Any]] = {}
    for number, cells in xlsx_stream.iter_sheet_rows(path, sheet, values=data_only):
        if number > MAX_SHEET_ROWS:
            break
        grid[number] = dict(cells)
    return grid


def _fills(path: Path, sheet: str, rows: set[int], cols: list[int]) -> dict[int, str]:
    from app.workbook_formatting import css_color, theme_palette
    result: dict[int, str] = {}
    if not rows or not cols:
        return result
    book = load_workbook(path, read_only=True)
    try:
        palette = theme_palette(book)
        worksheet = book[sheet]
        anywhere: dict[int, Counter] = {}
        for row in worksheet.iter_rows(min_row=min(rows), max_row=max(rows)):
            for cell in row:
                if getattr(cell, "row", None) not in rows or getattr(cell, "fill", None) is None or cell.fill.patternType != "solid":
                    continue
                colour = css_color(cell.fill.fgColor, palette)
                if not colour or colour in ("#FFFFFF", "#000000"):
                    continue
                if getattr(cell, "column", None) in cols:
                    result.setdefault(cell.row, colour)
                anywhere.setdefault(cell.row, Counter())[colour] += 1
        # A whole coloured row (six or more cells) is a mark on the item even when the name cell is plain.
        # Smaller blocks elsewhere are layout — e.g. Celine shades the size grid that applies — not a status.
        for number, counts in anywhere.items():
            colour, n = counts.most_common(1)[0]
            if n >= 6:
                result.setdefault(number, colour)
    except Exception:  # noqa: BLE001 - formatting is optional evidence
        return result
    finally:
        book.close()
    return result


def _header_score(cells: dict[int, Any]) -> tuple[int, dict[str, list[int]], list[int]]:
    mapping: dict[str, list[int]] = {}
    sizes = []
    for col, value in sorted(cells.items()):
        if isinstance(value, str) and value.startswith("="):
            continue
        label = norm_label(value)
        field_name = LABEL_TO_FIELD.get(label)
        if field_name is None and label:
            field_name = next((LABEL_TO_FIELD[l] for l in sorted(LABEL_TO_FIELD, key=len, reverse=True)
                               if " " in l and (label.startswith(l + " ") or label.endswith(" " + l))), None)
        if field_name:
            mapping.setdefault(field_name, []).append(col)
        elif size_label(value):
            sizes.append(col)
    runs = _size_runs(sizes)
    score = len(mapping) + (2 if runs else 0)
    return score, mapping, runs


def _pick(mapping: dict[str, list[int]], rows: list[dict[int, Any]]) -> dict[str, int]:
    """Where a label appears twice (e.g. 'Article' and 'Article number'), keep the fuller column."""
    result = {}
    for field_name, cols in mapping.items():
        result[field_name] = max(cols, key=lambda c: (sum(1 for r in rows if r.get(c) not in (None, "")), -cols.index(c)))
    return result


def _size_runs(cols: list[int]) -> list[int]:
    best: list[int] = []
    run: list[int] = []
    for col in sorted(cols):
        if run and col == run[-1] + 1:
            run.append(col)
        else:
            run = [col]
        if len(run) > len(best):
            best = list(run)
    return best if len(best) >= 3 else []


def _is_total_row(cells: dict[int, Any], formulas: dict[int, Any], mapping: dict[str, int]) -> bool:
    if any(isinstance(v, str) and TOTAL_WORDS.search(v) for v in cells.values()):
        return True
    name_cols = [mapping.get("name"), mapping.get("article")]
    has_name = any(c and cells.get(c) not in (None, "") for c in name_cols)
    return not has_name and any(isinstance(v, str) and v.upper().startswith(("=SUM", "=+SUM")) for v in formulas.values())


RATE_FORMULA = re.compile(r"^=\s*(ROUNDUP\s*\()?\s*\+?\s*([A-Z]{1,3}\d+|\d+(?:\.\d+)?)\s*\*\s*([A-Z]{1,3}\d+|\d*\.\d+)", re.I)


def _rate_in_row(number: int, cells: dict[int, Any], formulas: dict[int, Any]) -> tuple[Decimal, str, str] | None:
    """A percentage row: '=K15*J16' (J16 a fraction in this row), '=H10*0.3' or '=ROUNDUP(K5*0.35,0)'."""
    from openpyxl.utils import column_index_from_string
    for col, formula in sorted(formulas.items()):
        if not isinstance(formula, str):
            continue
        match = RATE_FORMULA.match(formula.replace("$", "").replace(" ", ""))
        if not match:
            continue
        rounding = "euro_up" if match.group(1) else "cent"
        for operand in (match.group(3), match.group(2)):
            ref = re.fullmatch(r"([A-Z]{1,3})(\d+)", operand.upper())
            if ref:
                if int(ref.group(2)) != number:
                    continue
                value = parse_number(cells.get(column_index_from_string(ref.group(1))))
            else:
                value = parse_number(operand)
            if value is not None and 0 < value < 1:
                return value, rounding, f"{get_column_letter(col)}{number}: {formula}"
    return None


PRODUCT_FORMULA = re.compile(r"^=\s*\+?\s*([A-Z]{1,3})(\d+)\s*\*\s*([A-Z]{1,3})(\d+)\s*$", re.I)


def _infer_columns(rows: list[dict[int, Any]], formula_rows: list[dict[int, Any]], mapping: dict[str, int],
                   size_cols: list[int], skip: set[int], labelled: set[int], known: set[str]) -> dict[str, int]:
    """Fill in unlabeled columns from the data: line-total formulas (=price*qty), codes, long text, initials."""
    from openpyxl.utils import column_index_from_string
    mapping = dict(mapping)
    # Line totals written as =H3*G3 identify the quantity and price columns exactly.
    votes: Counter = Counter()
    for cells, formulas in zip(rows, formula_rows):
        for col, formula in formulas.items():
            match = PRODUCT_FORMULA.match(str(formula).replace("$", "")) if isinstance(formula, str) else None
            if match and match.group(2) == match.group(4):
                votes[(col, column_index_from_string(match.group(1).upper()), column_index_from_string(match.group(3).upper()))] += 1
    if votes and not ({"qty", "price"} <= mapping.keys()):
        (total_col, left, right), _ = votes.most_common(1)[0]
        def small(col):
            vals = [parse_number(r.get(col)) for r in rows if parse_number(r.get(col)) is not None]
            return vals and sum(1 for v in vals if v <= 30) / len(vals) > .7
        qty_col, price_col = (left, right) if small(left) and not small(right) else (right, left)
        if not any(c in size_cols for c in (qty_col, price_col)) or "qty" not in mapping:
            mapping.setdefault("total", total_col)
            mapping.setdefault("qty", qty_col)
            mapping.setdefault("price", price_col)
    used = set(mapping.values()) | set(size_cols) | skip
    if "price" not in mapping or "qty" not in mapping or "size" not in mapping:
        numeric: dict[int, list[Decimal]] = {}
        sizes: Counter = Counter()
        for cells in rows:
            for col, value in cells.items():
                if col in used:
                    continue
                number = parse_number(value) if not isinstance(value, str) or not value.startswith("=") else None
                if number is not None:
                    numeric.setdefault(col, []).append(number)
                if size_label(value) or (isinstance(value, str) and re.fullmatch(r"\d{2}(?:[.,/]\d)?(?:\s*,\s*\d{2})*", value.strip())):
                    sizes[col] += 1
        total = max(1, len(rows))
        if "qty" not in mapping:
            best = max(((len(v), c) for c, v in numeric.items() if all(x == x.to_integral_value() and 0 < x <= 30 for x in v)
                        and len(v) / total > .6), default=None)
            if best:
                mapping["qty"] = best[1]
                used.add(best[1])
        if "price" not in mapping:
            candidates = [c for c, v in numeric.items() if c not in used and len(v) / total > .6
                          and sum(1 for x in v if x > 60) / len(v) > .7]
            qty_col = mapping.get("qty")

            def explains_total(col):
                # price × qty equals another column on every row → that column is the line total.
                if qty_col is None:
                    return False
                for other in numeric:
                    if other in (col, qty_col):
                        continue
                    pairs = [(parse_number(r.get(col)), parse_number(r.get(qty_col)), parse_number(r.get(other))) for r in rows]
                    pairs = [p for p in pairs if None not in p]
                    if pairs and len(pairs) / total > .6 and all(money(a * b) == money(c) for a, b, c in pairs):
                        mapping.setdefault("total", other)
                        return True
                return False
            exact = [c for c in candidates if explains_total(c)]
            best = exact[0] if exact else max(candidates, key=lambda c: sum(numeric[c]) / len(numeric[c]), default=None)
            if best is not None:
                mapping["price"] = best
                used.add(best)
        if "size" not in mapping and not size_cols:
            best = max(((n, c) for c, n in sizes.items() if c not in used and n / total > .5), default=None)
            if best:
                mapping["size"] = best[1]
                used.add(best[1])
    stats: dict[int, dict[str, float]] = {}
    values_by_col: dict[int, Counter] = {}
    for cells in rows:
        for col, value in cells.items():
            if col in used or not isinstance(value, str) or value.startswith(("=", "#")):
                continue   # formulas and Excel errors such as "#VALUE!" are not data
            text = value.strip()
            values_by_col.setdefault(col, Counter())[text.upper()] += 1
            entry = stats.setdefault(col, {"n": 0, "code": 0, "alloc": 0, "len": 0, "known": 0, "initials": 0})
            entry["n"] += 1
            entry["len"] += len(text)
            if re.fullmatch(r"[A-Z0-9][A-Z0-9_.\-/ ]{5,}", text.strip(), re.I) and re.search(r"\d", text) and re.search(r"[A-Za-z]", text) and " " not in text.strip():
                entry["code"] += 1
            words = re.findall(r"[A-Za-zА-Яа-яЁё]{2,}", text)
            if words and len(text) <= 70 and len(words) <= 6 and all(len(w) <= 12 for w in words) \
                    and not re.search(r"[|]", text) and sum(len(w) for w in words) <= 40:
                entry["alloc"] += 1
            if any(w.upper() in known for w in words):
                entry["known"] += 1
            # Initials look like "AB", "CC 42", "Ids 40, AL 38": short capitalised tokens, usually with sizes.
            tokens = re.findall(r"[A-Za-zА-Яа-яЁё]+", text)
            if tokens and all((len(t) <= 4 and t[:1].isupper()) or t.upper() in known for t in tokens) and \
                    any(len(t) <= 4 and (t.isupper() or t.upper() in known) for t in tokens):
                entry["initials"] += 1
    total = max(1, len(rows))
    if "article" not in mapping:
        best = max(((s["code"] / total, c) for c, s in stats.items()), default=(0, None))
        if best[0] >= 0.4:
            mapping["article"] = best[1]
            used.add(best[1])
    if "allocation" not in mapping:
        candidates = []
        for col, entry in stats.items():
            if col in used:
                continue
            distinct = values_by_col.get(col, Counter())
            constant = distinct and distinct.most_common(1)[0][1] / max(1, entry["n"]) > .8 and entry["known"] == 0 and len(rows) > 3
            if constant or (col in labelled and entry["known"] / total < .3):
                continue
            score = entry["alloc"] / total + entry["known"] / total + 2 * entry["initials"] / total
            if entry["alloc"] / total >= .5 or entry["known"] / total >= .3 or entry["initials"] / total >= .5:
                candidates.append((score, -col, col))
        if candidates:
            mapping["allocation"] = max(candidates)[2]
            used.add(mapping["allocation"])
    if "name" not in mapping:
        best = max(((s["len"] / max(1, s["n"]) * (s["n"] / total), c) for c, s in stats.items()
                    if c not in used and s["code"] / max(1, s["n"]) < .5), default=(0, None))
        if best[1] is not None and best[0] > 3:
            mapping["name"] = best[1]
    return mapping


def _own_line_formula(number: int, cells_v: dict, cells_f: dict, total_col: int | None) -> tuple[Decimal, int, str] | None:
    """(unit price, quantity, formula) from this row's "=A<n>*B<n>" whose cached value equals the product."""
    from openpyxl.utils import column_index_from_string
    candidates = sorted(cells_f.items(), key=lambda kv: kv[0] != total_col)
    for col, formula in candidates:
        match = PRODUCT_FORMULA.match(str(formula).replace("$", "")) if isinstance(formula, str) else None
        if not match or int(match.group(2)) != number or int(match.group(4)) != number:
            continue
        line = parse_number(cells_v.get(col))
        a = parse_number(cells_v.get(column_index_from_string(match.group(1).upper())))
        b = parse_number(cells_v.get(column_index_from_string(match.group(3).upper())))
        if not line or not a or not b or money(a * b) != money(line):
            continue
        whole = [x for x in (a, b) if x == x.to_integral_value() and 0 < x <= 1000]
        if not whole:
            continue
        qty = min(whole) if len(whole) == 2 else whole[0]
        price = b if qty is a else a
        if len(whole) == 2 and a == b:
            price = a
        return money(price), int(qty), str(formula)
    return None


def analyze_sheet(path: Path, sheet: str, anchors: list, known: set[str]) -> SheetResult:
    values = _grid(path, sheet, True)
    formulas = _grid(path, sheet, False)
    notes: list[str] = []
    brand = re.sub(r"\s+", " ", sheet).strip()
    if not formulas:
        return SheetResult(sheet, "empty", brand, [], [], [], {}, [], None)
    first_text = " ".join(str(v) for v in formulas.get(min(formulas), {}).values())
    refs_other = sum(1 for cells in list(formulas.values())[:40] for v in cells.values() if isinstance(v, str) and "!" in v and v.startswith("="))
    if re.search(r"заказы|orders summary", first_text, re.I) or refs_other >= 4:
        return SheetResult(sheet, "summary", brand, [], [], [], {}, ["Сводный лист — не товары."], None)
    numbers = sorted(formulas)
    title = None
    for number in numbers[:3]:
        texts = [v for v in formulas[number].values() if isinstance(v, str) and not v.startswith("=")]
        if len(formulas[number]) == 1 and texts and len(texts[0]) > 3:
            title = texts[0]
            break
    header: dict[str, int] = {}
    size_cols: list[int] = []
    grid_rows: dict[str, dict[int, str]] = {}
    header_rows: list[int] = []
    items: list[ItemRow] = []
    sections: list[SectionRate] = []
    pending_rows: list[int] = []
    last_total: tuple[str, Decimal | None, int] | None = None
    ship_cols: dict[int, str] = {}
    # Pass 1: find header rows (a sheet may hold several tables, and size-only header rows).
    header_at: dict[int, tuple[dict[str, list[int]], list[int]]] = {}
    for number in numbers:
        score, mapping, runs = _header_score(formulas[number])
        item_numbers = sum(1 for v in values.get(number, {}).values() if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 60)
        if item_numbers > 1:
            continue
        if score >= 3 and (mapping.keys() & {"price", "qty", "total", "name", "article"} or runs):
            header_at[number] = (mapping, runs)
        elif runs and len(runs) >= 4 and not mapping.keys() & {"name", "article"}:
            header_at[number] = (mapping, runs)   # size labels only (e.g. a shoe grid further down)
    if not header_at:
        # No header at all: start before the first row that looks like goods (text + price-like number).
        for number in numbers:
            cells = values.get(number, {})
            if any(isinstance(v, str) and len(v.strip()) > 2 for v in cells.values()) and \
                    any(isinstance(v, (int, float)) and not isinstance(v, bool) and v > 60 for v in cells.values()):
                header_at[number - 1] = ({}, [])
                notes.append("В листе нет заголовков — столбцы определены по данным и формулам строк.")
                break
    if not header_at:
        return SheetResult(sheet, "unreadable", brand, [], [], [], {}, ["Не найдены товары (название + цена)."], None)
    for number in sorted(set(numbers) | set(header_at)):
        cells_v = values.get(number, {})
        cells_f = formulas.get(number, {})
        if number in header_at:
            mapping_lists, runs = header_at[number]
            below = [n for n in numbers if n > number][:60]
            if header and not mapping_lists.keys() & {"name", "article", "price", "qty", "total"}:
                # A size-only row inside a table: new size labels; quantity/price may have moved too.
                size_cols = runs
                grid_rows = {"": {col: str(cells_f.get(col)).strip() for col in runs}}
                block = [n for n in below if not _is_total_row(values.get(n, {}), formulas.get(n, {}), header)][:8]
                moved = _infer_columns([values.get(n, {}) for n in block], [formulas.get(n, {}) for n in block],
                                       {}, size_cols, set(ship_cols), set(), known)
                for key in ("qty", "price", "total"):
                    if key in moved:
                        header[key] = moved[key]
                continue
            header = _pick(mapping_lists, [values.get(n, {}) for n in below])
            size_cols = runs or []
            if not size_cols:
                for above in range(max(1, number - 6), number):
                    _, _, above_runs = _header_score(formulas.get(above, {}))
                    if len(above_runs) > len(size_cols):
                        size_cols = above_runs
            header_rows.append(number)
            label_row = formulas.get(number, {}) if number in formulas else formulas.get(size_cols and min(
                (n for n in range(max(1, number - 6), number) if set(size_cols) <= set(formulas.get(n, {}))), default=number) or number, {})
            labels = {col: str(label_row.get(col)).strip() for col in size_cols if label_row.get(col) not in (None, "")}
            grid_rows = {"": labels}
            for above in range(max(1, number - 6), number + 1):
                row_cells = formulas.get(above, {})
                for col, value in row_cells.items():
                    if isinstance(value, str) and re.fullmatch(r"[A-Da-d]\s*", value) and size_cols and col < size_cols[0]:
                        grid_rows[value.strip().upper()] = {c: str(row_cells.get(c)).strip() for c in size_cols if row_cells.get(c) not in (None, "")}
            ship_cols = {col: str(v).strip() for col, v in cells_f.items() if isinstance(v, str) and SHIP_HEADER.search(v)
                         and col not in header.values()}
            labelled = {col for col, v in cells_f.items() if isinstance(v, str) and v.strip() and not size_label(v)}
            goods = [n for n in below if any(isinstance(v, str) and len(v.strip()) > 1 for v in values.get(n, {}).values())
                     and any(isinstance(v, (int, float)) and not isinstance(v, bool) and v > 60 for v in values.get(n, {}).values())
                     and not any(isinstance(v, str) and TOTAL_WORDS.search(v) for v in values.get(n, {}).values())
                     and not _rate_in_row(n, values.get(n, {}), formulas.get(n, {}))] or below
            header = _infer_columns([values.get(n, {}) for n in goods], [formulas.get(n, {}) for n in goods],
                                    header, size_cols, set(ship_cols), labelled, known)
            continue
        if not header:
            continue
        rate = _rate_in_row(number, cells_v, cells_f)
        is_total = _is_total_row(cells_v, cells_f, header)
        name = clean_text(cells_v.get(header.get("name"))) if header.get("name") else None
        article = clean_text(cells_v.get(header.get("article"))) if header.get("article") else None
        alloc_here = clean_text(cells_v.get(header.get("allocation"))) if header.get("allocation") else None
        if rate and not name and not article and not alloc_here:
            if not pending_rows and sections:
                # Another rate for the same goods: the block is split into equal portions (e.g. "=Q32/2").
                previous = sections[-1]
                group = [x for x in sections if x.rows == previous.rows]
                share = str(Decimal(1) / (len(group) + 1))
                for x in group:
                    x.share = share
                sections.append(SectionRate(sheet, f"Часть {len(group) + 1}", str(rate[0]), rate[1], rate[2], previous.source_total,
                                            list(previous.rows), share))
                continue
            label = (last_total[0] if last_total else "") or f"Часть {len(sections) + 1}"
            sections.append(SectionRate(sheet, label, str(rate[0]), rate[1], rate[2],
                                        str(last_total[1]) if last_total and last_total[1] is not None else None, list(pending_rows)))
            pending_rows = []
            continue
        if is_total and not name and not article and not alloc_here:
            label_text = " ".join(str(v) for v in cells_v.values() if isinstance(v, str))
            label = re.sub(TOTAL_WORDS, "", label_text).strip(" :") or ""
            total_col = header.get("total")
            cached = parse_number(cells_v.get(total_col)) if total_col else None
            last_total = (label, cached, number)
            continue
        price_probe = parse_number(cells_v.get(header.get("price"))) if header.get("price") else None
        if not name and not article and not (alloc_here and price_probe):
            continue
        if isinstance(name, str) and name.upper() == name and len(name) <= 5 and not article and not cells_v.get(header.get("price") or -1):
            continue  # section caption like "ACC"
        price = parse_number(cells_v.get(header.get("price"))) if header.get("price") else None
        qty_value = parse_number(cells_v.get(header.get("qty"))) if header.get("qty") else None
        if qty_value is None:
            # This row's own line formula (e.g. "=R10*O10") says where its quantity is.
            from openpyxl.utils import column_index_from_string
            for formula in cells_f.values():
                match = PRODUCT_FORMULA.match(str(formula).replace("$", "")) if isinstance(formula, str) else None
                if match and int(match.group(2)) == number:
                    a, b = column_index_from_string(match.group(1).upper()), column_index_from_string(match.group(3).upper())
                    other = b if a == header.get("price") else a if b == header.get("price") else None
                    if other is not None:
                        qty_value = parse_number(cells_v.get(other))
                        break
        cancelled = any(isinstance(v, str) and re.fullmatch(r"\s*(cancell?ed|cancel|annullato|отмен\w*|cancelado)\s*", v, re.I)
                        for v in cells_v.values())
        grid_key = ""
        for col, value in cells_v.items():
            if isinstance(value, str) and re.fullmatch(r"[A-Da-d]\s*", value) and col < (size_cols[0] if size_cols else 0):
                grid_key = value.strip().upper()
        labels = grid_rows.get(grid_key) or grid_rows.get("", {})
        sizes: list[str] = []
        for col in size_cols:
            count = parse_number(cells_v.get(col))
            if count is not None and count > 0 and count == count.to_integral_value() and count < 100:
                sizes += [labels.get(col) or get_column_letter(col)] * int(count)
        if header.get("size"):
            raw_size = cells_v.get(header["size"])
            if raw_size not in (None, ""):
                listed = [s for s in re.split(r"[,;\s]+", str(raw_size).replace(".0", "")) if s and s not in ("/",)]
                if not sizes:
                    sizes = [s for s in listed if size_label(s.split("/")[0]) or SIZE_RE.match(s)] or [str(raw_size).strip()]
        quantity = int(qty_value) if qty_value is not None and qty_value == qty_value.to_integral_value() and qty_value > 0 else len(sizes)
        if quantity <= 0:
            if price is None:
                continue
            quantity = 1
            notes.append(f"Строка {number}: количество не указано — принято 1 шт.")
        if price is None and header.get("total"):
            total_value = parse_number(cells_v.get(header["total"]))
            if total_value is not None and quantity:
                price = money(total_value / quantity)
                notes.append(f"Строка {number}: цена рассчитана из суммы строки.")
        own = _own_line_formula(number, cells_v, cells_f, header.get("total"))
        if own is not None:
            # The factory's own "=price*qty" for this row is the ground truth when headers are misaligned
            # (e.g. a second table whose captions sit one column off).
            own_price, own_qty, formula = own
            if price is None or money(price) != money(own_price) or quantity != own_qty:
                if price is not None or qty_value is not None:
                    notes.append(f"Строка {number}: цена и количество взяты из формулы строки фабрики ({formula}), "
                                 f"а не из заголовков ({price if price is not None else '—'} × {quantity}).")
                price, quantity = own_price, own_qty
        size_conflict = len(sizes) != quantity and size_cols and sizes
        if size_conflict:
            notes.append(f"Строка {number}: в размерной сетке {len(sizes)} шт., в количестве {quantity}.")
        alloc_text = alloc_here
        if not alloc_text and name and split_allocation(name, known) and all(
                w in known or w.translate(HOMOGLYPHS) in known for w, _ in split_allocation(name, known)):
            alloc_text, name = name, None   # initials typed in the name column
        segments = split_allocation(alloc_text, known)
        units, confident, issue = allocate_units(segments, sizes if len(sizes) == quantity else (sizes if sizes and len(set(sizes)) == 1 else []), quantity)
        if not confident and len(sizes) == quantity:
            units = [Unit(u.customer, sizes[i] if i < len(sizes) else None) for i, u in enumerate(units)]
        if price is None:
            issue = (issue + "; " if issue else "") + "нет цены"
            confident = False
        if size_conflict:
            issue = (issue + "; " if issue else "") + f"размерная сетка {len(sizes)} шт. ≠ количество {quantity}"
            confident = False
        if not article and name:
            combined = re.match(r"^([A-Z0-9][A-Z0-9_.\-]{6,})\s+-\s+\d?(.+)$", name)
            if combined:
                article, name = combined.group(1), combined.group(2).strip()
        if not name and not article:
            name = f"Без названия (строка {number})"
        image = next((a.media for a in anchors if a.row <= number <= max(a.row, a.to_row - 1) or a.row == number), None)
        if image is None:
            image = next((a.media for a in anchors if abs(a.row - number) <= 0 and a.col <= 4), None)
        status_hint = shipment_hint = None
        for col, label in ship_cols.items():
            if cells_v.get(col) not in (None, "", 0):
                if re.search(r"ready|готов", label, re.I):
                    status_hint = "Ready"
                else:
                    shipment_hint = label
        status_hint = status_hint or status_from_text(alloc_text)
        if cancelled or status_hint == "Cancelled":
            status_hint = "Cancelled"
        if header.get("total") and cells_f.get(header["total"]) in (None, "") and price is not None and not cancelled:
            notes.append(f"Строка {number}: в файле фабрики у строки нет суммы — фабрика могла не включить её в итог.")
        items.append(ItemRow(sheet, number, name, (article or "").replace("\t", "").strip() or None,
                             clean_text(cells_v.get(header.get("color"))) if header.get("color") else None,
                             quantity, str(money(price)) if price is not None else None, sizes, alloc_text,
                             [asdict(u) for u in units], confident, issue, None, image, status_hint, shipment_hint, None,
                             clean_text(cells_v.get(header.get("look"))) if header.get("look") else None,
                             {"cell": f"{get_column_letter(header.get('name') or header.get('article') or 1)}{number}"}))
        pending_rows.append(number)
    # Sections: rates found after blocks of rows.
    if sections:
        if len({s.rate for s in sections}) > 1 and not all(x.share for x in sections):
            for index, section in enumerate(sections):
                label = section.label or f"Часть {index + 1}"
                section.label = label
                for item in items:
                    if item.row in section.rows:
                        item.section = label
        else:
            notes.append(f"Процент фабрики: {Decimal(sections[0].rate) * 100:.2f}%".replace(".00", ""))
        uncovered = [i.row for i in items if not any(i.row in s.rows for s in sections)]
        if uncovered and len(sections) == 1:
            sections[0].rows += uncovered
        elif uncovered:
            notes.append(f"Строки {uncovered[:5]} после последнего процента — без процента фабрики.")
    else:
        notes.append("Процент фабрики в листе не найден — укажите его при проверке.")
    fills = _fills(path, sheet, {i.row for i in items}, [c for c in (header.get("name"), header.get("article")) if c])
    for item in items:
        item.fill = fills.get(item.row)
    # Supplier's own grand total (cached) for reconciliation.
    grand = None
    for number in sorted(formulas, reverse=True)[:6]:
        for col, formula in formulas[number].items():
            if isinstance(formula, str) and formula.startswith("=") and parse_number(values.get(number, {}).get(col)) is not None:
                grand = parse_number(values[number][col])
                break
        if grand is not None:
            break
    computed = money(sum((i.base for i in items if i.status_hint != "Cancelled"), ZERO))
    generic = re.fullmatch(r"(sheet|лист|page)\s*\d*", brand, re.I)
    return SheetResult(sheet, "order" if items else "empty", _brand_from_title(title) if title and generic else brand, items, sections, header_rows,
                       {k: get_column_letter(v) for k, v in header.items()}, notes,
                       str(grand) if grand is not None else None, str(computed))


def _brand_from_title(title: str) -> str:
    cleaned = re.sub(r"^\s*(order|заказ)\s+", "", title, flags=re.I)
    cleaned = re.sub(r"\b(SS|FW|AW)\s*\d{2,4}\b.*$", "", cleaned, flags=re.I)
    cleaned = re.sub(r"\b(deadline|alina)\b.*$", "", cleaned, flags=re.I).strip(" -·")
    return cleaned[:60] or title[:60]


def _payments_from_summary(path: Path, sheet: str) -> list[dict]:
    """Payment list under 'ОПЛАТЫ …' on a summary sheet: (label, date, amount, payer code)."""
    rows = _grid(path, sheet, True)
    result, active = [], False
    for number in sorted(rows):
        cells = rows[number]
        text = " ".join(str(v) for v in cells.values() if isinstance(v, str))
        if re.search(r"оплат", text, re.I) and not active and len(cells) <= 2:
            active = True
            continue
        if active:
            if re.search(r"всего|итого|отгруз", text, re.I):
                break
            ordered = [cells[c] for c in sorted(cells)]
            amount = next((parse_number(v) for v in ordered if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0), None)
            if amount is None:
                continue
            when = next((v for v in ordered if isinstance(v, datetime)), None)
            texts = [v for v in ordered if isinstance(v, str)]
            payer = texts[-1].strip() if len(texts) >= 2 or (texts and len(texts[-1]) <= 20) else None
            result.append({"row": number, "label": texts[0].strip() if texts else "", "date": when.date().isoformat() if when else None,
                           "amount": str(money(amount)), "payer": payer})
    return result


def season_from(filename: str, sheets: list[str]) -> str | None:
    for text in [filename, *sheets]:
        match = re.search(r"(?<![A-Za-z0-9])(SS|FW|AW|PF|RS)[ _-]*(20)?(\d{2})(?!\d)", text, re.I)
        if match:
            return match.group(1).upper() + match.group(3)
    return None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def convert_legacy(path: Path) -> Path:
    """Convert .xls to .xlsx with LibreOffice when installed (macros are not run)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice") or (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice" if Path("/Applications/LibreOffice.app").exists() else None)
    if not soffice:
        raise ValueError("Файлы .xls (старый Excel) нужно сохранить как .xlsx в Excel — LibreOffice для конвертации не установлен.")
    out = path.parent / "converted"
    out.mkdir(exist_ok=True)
    subprocess.run([soffice, "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                   check=True, timeout=300, capture_output=True)
    converted = out / (path.stem + ".xlsx")
    if not converted.exists():
        raise ValueError("Не удалось конвертировать .xls в .xlsx.")
    return converted


def analyze_workbook(path: Path, filename: str, known_tokens: set[str] | None = None, season_hint: str | None = None) -> Analysis:
    path = Path(path)
    size = path.stat().st_size
    if size > MAX_IMPORT_BYTES:
        raise ValueError("Файл больше 1 ГБ.")
    if not zipfile.is_zipfile(path):
        raise ValueError("Это не файл Excel .xlsx/.xlsm (или он повреждён).")
    known = {t.upper() for t in (known_tokens or set())}
    sheets: list[SheetResult] = []
    payments: list[dict] = []
    warnings: list[str] = []
    with zipfile.ZipFile(path) as archive:
        parts = xlsx_stream.sheet_parts(archive)
        anchors = {name: xlsx_stream.image_anchors(archive, part) for name, part in parts}
    for name, _part in parts:
        try:
            result = analyze_sheet(path, name, anchors.get(name, []), known)
        except Exception as exc:  # noqa: BLE001 - one odd sheet must not stop the others
            result = SheetResult(name, "unreadable", name.strip(), [], [], [], {}, [f"Лист не прочитан: {type(exc).__name__}"], None)
        if result.kind == "summary":
            payments += _payments_from_summary(path, name)
        sheets.append(result)
    tokens: Counter = Counter()
    for sheet in sheets:
        for item in sheet.items:
            for unit in item.units:
                if unit["customer"]:
                    tokens[unit["customer"]] += 1
            for word, _ in split_allocation(item.allocation_text, known):
                tokens.setdefault(word, 0)
    for sheet in sheets:
        if sheet.kind == "order" and sheet.source_grand_total and sheet.sections:
            expected = Decimal(sheet.computed_base)
            from app.settlement import Portion, factory_price
            portions, explicit = [], Decimal(0)
            for x in sheet.sections:
                if x.share:
                    portions.append(Portion(Decimal(x.rate), share=Decimal(x.share), policy=x.rounding))
                elif len({y.rate for y in sheet.sections}) > 1:
                    base = money(sum((i.base for i in sheet.items if i.row in x.rows and i.status_hint != "Cancelled"), ZERO))
                    explicit += base
                    portions.append(Portion(Decimal(x.rate), base=base, policy=x.rounding))
                else:
                    portions.append(Portion(Decimal(x.rate), share=Decimal(1), policy=x.rounding))
                    break
            try:
                expected = factory_price(expected, portions).total
            except ValueError:
                expected = None
            if expected is not None:
                diff = money(Decimal(sheet.source_grand_total) - expected)
                if diff and abs(diff) > Decimal("0.01"):
                    text = f"{sheet.brand}: итог фабрики {sheet.source_grand_total}, по строкам получается {expected} (разница {diff})."
                    unsummed = [int(n) for n in re.findall(r"Строка (\d+): в файле фабрики у строки нет суммы", " ".join(sheet.notes))]
                    if unsummed:
                        base = money(sum((i.base for i in sheet.items if i.row in unsummed and i.status_hint != "Cancelled"), ZERO))
                        text += (f" У строк {', '.join(map(str, unsummed))} в файле нет суммы (база €{base}) — похоже, фабрика их не посчитала."
                                 " Уточните у фабрики, входят ли они в заказ.")
                    warnings.append(text)
    return Analysis(filename, sha256_file(path), size, season_hint or season_from(filename, [s.name for s in sheets]),
                    sheets, dict(tokens), payments, warnings)


# --- Commit -------------------------------------------------------------------------

def store_original(path: Path, sha: str, filename: str) -> Path:
    ORIGINALS_DIR.mkdir(parents=True, exist_ok=True)
    target = ORIGINALS_DIR / f"{sha}{Path(filename).suffix.lower() or '.xlsx'}"
    if not target.exists():
        temp = target.with_suffix(target.suffix + ".part")
        shutil.copyfile(path, temp)
        if sha256_file(temp) != sha:
            temp.unlink(missing_ok=True)
            raise ValueError("Копия файла не совпала с оригиналом — импорт остановлен.")
        temp.replace(target)
        target.chmod(0o444)
    return target


def commit_import(db, analysis: Analysis, original: Path, decisions: dict, actor: str = "operator") -> dict:
    """Create season/brand/order/items/agreements from a reviewed analysis.

    ``decisions``: {"season": "SS27", "token_map": {"AL": "ALINA", "IDS": "new"}, "owner": "ALINA",
    "rates": {sheet: [{"label","rate","section"}]}, "skip_sheets": [...], "allocations": {"sheet:row": [{customer,size,quantity}]},
    "brand_names": {sheet: name}}.
    Rows that remain ambiguous are created without a customer and listed for review.
    """
    from sqlalchemy import func, select
    from app.audit import record_change
    from app.models import (AllocationReview, Brand, Customer, CustomerAlias, FactoryAgreement, FactoryRatePortion,
                            ImportBatch, Order, OrderItem, Season)
    from app.ledger_actions import effective_rule
    existing = db.scalar(select(ImportBatch).where(ImportBatch.sha256 == analysis.sha256))
    if existing is not None:
        return {"batch_id": existing.id, "duplicate": True, "created": 0, "review": 0}
    season_name = (decisions.get("season") or analysis.season or "").strip().upper()
    if not season_name:
        raise ValueError("Укажите сезон (например, SS27).")
    season = db.scalar(select(Season).where(Season.name == season_name)) or Season(name=season_name)
    db.add(season)
    db.flush()
    batch = ImportBatch(filename=analysis.filename[:300], sha256=analysis.sha256, stored_path=str(original),
                        summary_json=json.dumps({"kind": "factory", "season": season_name, "decisions": decisions}, ensure_ascii=False))
    db.add(batch)
    db.flush()
    token_map = {k.upper(): v for k, v in (decisions.get("token_map") or {}).items()}
    customers: dict[str, Customer] = {}

    def customer_for(token: str | None) -> Customer | None:
        if not token:
            return None
        token = token.upper()
        target = token_map.get(token, token)
        if target in ("", "skip", None):
            return None
        if target == "new":
            target = token
        if target in customers:
            return customers[target]
        customer = db.scalar(select(Customer).where(func.upper(Customer.code) == target.upper()))
        if customer is None:
            alias = db.scalar(select(CustomerAlias).where(func.upper(CustomerAlias.alias) == target.upper()))
            customer = alias.customer if alias else None
        if customer is None:
            code = re.sub(r"[^A-Za-zА-Яа-яЁё0-9]", "", target)[:20] or "X"
            customer = Customer(code=code, name=target.title() if len(target) > 3 else None)
            db.add(customer)
            db.flush()
            record_change(db, entity_type="Customer", entity_id=customer.id, field="code", previous_value=None,
                          new_value=code, source_user=actor, description=f"Новый клиент из файла {analysis.filename}")
        if token != customer.code.upper() and not db.scalar(select(CustomerAlias).where(func.upper(CustomerAlias.alias) == token)):
            db.add(CustomerAlias(alias=token, customer_id=customer.id))
            db.flush()
        customers[target] = customer
        return customer

    if decisions.get("owner"):
        owner = customer_for(decisions["owner"])
        if owner is not None:
            for other in db.scalars(select(Customer).where(Customer.is_owner.is_(True))):
                other.is_owner = False
            owner.is_owner = True
    created = review_count = 0
    image_root = IMAGES_DIR / analysis.sha256[:16]
    archive = zipfile.ZipFile(original)
    skip = set(decisions.get("skip_sheets") or [])
    only_rows = set(decisions["only_rows"]) if decisions.get("only_rows") is not None else None
    row_units = decisions.get("row_units") or {}
    try:
        for sheet in analysis.sheets:
            if sheet.kind != "order" or sheet.name in skip:
                continue
            if only_rows is not None and not any(f"{sheet.name}:{i.row}" in only_rows for i in sheet.items):
                continue
            brand_name = (decisions.get("brand_names") or {}).get(sheet.name) or sheet.brand
            brand = db.scalar(select(Brand).where(func.lower(Brand.name) == brand_name.lower())) or Brand(name=brand_name)
            db.add(brand)
            db.flush()
            order = db.scalar(select(Order).where(Order.season_id == season.id, Order.brand_id == brand.id))
            existed = order is not None
            if order is None:
                order = Order(season_id=season.id, brand_id=brand.id, source_sheet=sheet.name)
                db.add(order)
                db.flush()
            rates = (decisions.get("rates") or {}).get(sheet.name)
            if rates is None:
                rates = [{"label": s.label, "rate": s.rate, "rounding": s.rounding, "section": s.label if len({x.rate for x in sheet.sections}) > 1 else None}
                         for s in sheet.sections if s.rate is not None]
            if existed and decisions.get("keep_existing_terms"):
                rates = None      # a revision never silently changes agreed factory terms
            if rates:
                agreement = db.scalar(select(FactoryAgreement).where(FactoryAgreement.order_id == order.id)) or FactoryAgreement(order_id=order.id)
                db.add(agreement)
                agreement.portions.clear()
                db.flush()
                for index, r in enumerate(rates):
                    agreement.portions.append(FactoryRatePortion(position=index, label=str(r.get("label") or "")[:100],
                                                                 rate=Decimal(str(r["rate"])), section=r.get("section"),
                                                                 share=None if r.get("section") else Decimal("1") / len([x for x in rates if not x.get("section")]),
                                                                 rounding=r.get("rounding") or "cent"))
                agreement.confirmed = bool(decisions.get("rates_confirmed"))
                agreement.source_evidence = "; ".join(s.evidence for s in sheet.sections)[:2000]
                agreement.source_total = Decimal(sheet.source_grand_total) if sheet.source_grand_total else None
                order.source_total = agreement.source_total
            for item in sheet.items:
                key = f"{sheet.name}:{item.row}"
                if only_rows is not None and key not in only_rows:
                    continue
                chosen = (decisions.get("allocations") or {}).get(key)
                units = []
                if key in row_units:          # a revision: only the allocations that are new
                    units = [{"customer": u.get("customer"), "size": u.get("size")} for u in row_units[key]]
                    resolved = bool(units) and all(u["customer"] for u in units)
                elif chosen:
                    for part in chosen:
                        units += [{"customer": part.get("customer"), "size": part.get("size")}] * int(part.get("quantity") or 1)
                    resolved = len(units) == item.quantity and all(u["customer"] for u in units)
                else:
                    units = item.units
                    resolved = item.confident and all(u.get("customer") for u in units)
                image_path = None
                if item.image_part:
                    data = xlsx_stream.read_media(archive, item.image_part)
                    if data:
                        image_root.mkdir(parents=True, exist_ok=True)
                        target = image_root / f"{re.sub(r'[^A-Za-z0-9]+', '_', sheet.name)}_{item.row}{Path(item.image_part).suffix.lower()}"
                        target.write_bytes(data)
                        image_path = str(target.relative_to(ROOT_DIR)).replace("\\", "/")
                groups: dict[tuple, int] = {}
                for unit in units:
                    owner = customer_for(unit.get("customer")) if resolved else None
                    groups[(owner.id if owner else None, unit.get("size"))] = groups.get((owner.id if owner else None, unit.get("size")), 0) + 1
                if not resolved:
                    groups = {(None, ", ".join(item.sizes) or None): len(units) if key in row_units else item.quantity}
                group_id = f"{analysis.sha256[:12]}:{sheet.name}:{item.row}"
                for (customer_id, size), quantity in groups.items():
                    row = OrderItem(order_id=order.id, customer_id=customer_id, article=(item.article or None) and item.article[:200],
                                    name=(item.name or item.article or "")[:300], color=(item.color or None) and item.color[:200],
                                    size=(size or None) and str(size)[:100], quantity=quantity,
                                    wholesale_unit_price=Decimal(item.unit_price) if item.unit_price else None,
                                    markup_percentage=Decimal("0"), markup_source="rule",
                                    status=item.status_hint or "Ordered", image_path=image_path,
                                    image_source_filename=item.image_part, image_match_method="factory anchor" if image_path else None,
                                    source_sheet=sheet.name[:200], source_file=analysis.filename[:500], source_row=item.row,
                                    source_cell=item.extra.get("cell"), source_allocation_description=item.allocation_text,
                                    source_fill=item.fill, allocation_status="Allocated" if customer_id else "Needs review",
                                    allocation_group_id=group_id, pricing_section=item.section,
                                    source_note=item.shipment_hint and f"Отгрузка в файле: {item.shipment_hint}")
                    db.add(row)
                    db.flush()
                    row.order = order
                    rule = effective_rule(db, row)
                    if rule is not None:
                        row.markup_percentage = rule.markup
                    created += 1
                if not resolved:
                    review_count += 1
                    db.add(AllocationReview(import_batch_id=batch.id, source_file=analysis.filename[:500], sheet=sheet.name[:200],
                                            row=item.row, raw_description=item.allocation_text,
                                            raw_values=json.dumps(asdict(item), ensure_ascii=False, default=str),
                                            reason=item.issue or "Распределение не определено"))
    finally:
        archive.close()
    record_change(db, entity_type="ImportBatch", entity_id=batch.id, field=None, previous_value=None,
                  new_value=analysis.filename, source_user=actor,
                  description=f"Импорт файла фабрики {analysis.filename}: {created} строк, на проверку {review_count}")
    db.flush()
    from app.factory_ledger import record_snapshot
    db.expire_all()
    record_snapshot(db, db.get(Season, season.id), f"Импорт {analysis.filename}")
    db.commit()
    return {"batch_id": batch.id, "duplicate": False, "created": created, "review": review_count, "season": season_name}
